from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from .preflight_xlsx import _style_report_sheets


def _safe_sheet_title(title: str) -> str:
    """Sheet titles allow at most 31 chars and forbid ``[]:*?/\\``."""
    text = "".join("_" if char in r"[]:*?/\\" else char for char in (title or "功能"))
    return text.strip()[:31] or "功能"


def write_feature_log_xlsx(
    path: Path,
    sheets: list[tuple[str, list[tuple[Any, ...]], list[tuple[Any, ...]]]],
) -> None:
    """Write one 运行日志 workbook; each sheet records one flow feature.

    运行日志统一命名为 ``流程名_运行日志_时间戳.xlsx``，无论“是否输出结果”
    如何都会生成。它代替旧版把检查类合并在一起的“审核前检查”报告。
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    styled: list[tuple[Any, list[list[Any]]]] = []
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(_safe_sheet_title(title))
        styled.append((sheet, [list(headers), *[list(row) for row in rows]]))
    if not styled:
        sheet = workbook.create_sheet("运行日志")
        styled.append((sheet, [["流程", "说明"]]))
    _style_report_sheets(workbook, tuple(styled))
    workbook.save(path)


class FeatureLog:
    """Collect one sheet per flow feature, then write them to a single workbook.

    - ``add_sheet`` 在流程中每个功能执行后调用，收集标题、表头和行。
    - ``write`` 在流程结束时生成运行日志 xlsx；异常时也能写一份部分日志用于排错。

    运行日志由全局“输出流程运行日志”设置控制。检查/核对类步骤填
    “是否输出结果=是”时，这份日志就是其可保留输出；修改、汇总模块
    则按步骤设置保留阶段副本或最终工作簿。
    """

    def __init__(self, flow_name: str, output_dir: Path) -> None:
        self.flow_name = flow_name
        self.output_dir = output_dir
        self.timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        self.sheets: list[tuple[str, list[tuple[Any, ...]], list[tuple[Any, ...]]]] = []
        self.path = output_dir / f"{flow_name}_运行日志_{self.timestamp}.xlsx"

    def add_sheet(
        self,
        title: str,
        headers: tuple[Any, ...] | list[Any],
        rows: list[tuple[Any, ...]] | list[list[Any]],
    ) -> None:
        self.sheets.append((title, list(headers), [tuple(row) for row in rows]))

    def add_template_health(self, template_path: Path, items: list[Any]) -> None:
        """模板体检诊断：规则数量、校验区域重叠、#REF!、外部链接等，保留原审核前检查内容。"""
        rows = [
            (
                item.category,
                item.level,
                item.status,
                item.sheet_name or "",
                item.location or "",
                item.message,
            )
            for item in items
        ]
        self.add_sheet("模板体检", ("类别", "级别", "状态", "工作表", "位置", "说明"), rows)

    def write(self) -> Path | None:
        if not self.sheets:
            return None
        write_feature_log_xlsx(self.path, self.sheets)
        return self.path
