from __future__ import annotations

from collections import defaultdict
from numbers import Number
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import range_boundaries

from .models import CopyRange, PreflightItem, SourceMatch, StructureCheck, TemplateDefinition

STRUCTURE_MATCH_THRESHOLD = 0.90


def template_structure_values(template_workbook: Any, definition: TemplateDefinition) -> list[tuple[CopyRange, list[list[Any]]]]:
    """Read template labels in bulk once; source workbooks never need COM."""
    values: list[tuple[CopyRange, list[list[Any]]]] = []
    for item in definition.structure_ranges:
        value = template_workbook.Worksheets(item.sheet_name).Range(item.address).Value2
        if value is None:
            matrix = []
        elif not isinstance(value, tuple):
            matrix = [[value]]
        elif value and not isinstance(value[0], tuple):
            matrix = [list(value)]
        else:
            matrix = [list(row) for row in value]
        values.append((item, matrix))
    return values


def validate_source_xlsx(
    source_path: Path,
    definition: TemplateDefinition,
    expected_ranges: list[tuple[CopyRange, list[list[Any]]]],
    *,
    external_sheet_names: tuple[str, ...] = (),
) -> SourceMatch:
    """Check fixed labels from xlsx XML without launching Microsoft Excel."""
    workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
    try:
        source_names = set(workbook.sheetnames)
        required_names = {item.sheet_name for item in definition.copy_ranges}
        missing = tuple(sorted(required_names - source_names))
        if missing:
            return SourceMatch(False, 0.0, 0, 0, missing, "缺少模板要求的工作表：" + "、".join(missing))
        if not expected_ranges:
            return SourceMatch(False, 0.0, 0, 0, (), "模板没有可用的“表结构区域”命名区域，请先维护模板")

        checked = matched = 0
        mismatches: list[str] = []
        checks: list[StructureCheck] = []
        range_specs: list[tuple[CopyRange, list[list[Any]], tuple[int, int, int, int]]] = []
        bounds_by_sheet: dict[str, list[tuple[int, int, int, int]]] = defaultdict(list)
        for item, expected_values in expected_ranges:
            if item.sheet_name not in source_names:
                return SourceMatch(False, 0.0, 0, 0, (item.sheet_name,), "缺少模板要求的工作表：" + item.sheet_name)
            bounds = range_boundaries(item.address)
            range_specs.append((item, expected_values, bounds))
            bounds_by_sheet[item.sheet_name].append(bounds)

        # Read each compact bounding batch once.  Sparse, distant areas stay in
        # separate batches so a few named cells cannot force a whole sheet read.
        source_values: dict[tuple[str, int, int], Any] = {}
        for sheet_name, bounds_list in bounds_by_sheet.items():
            for min_col, min_row, max_col, max_row in _compact_batches(bounds_list):
                rows = workbook[sheet_name].iter_rows(
                    min_row=min_row, max_row=max_row,
                    min_col=min_col, max_col=max_col, values_only=True,
                )
                for row_number, row in enumerate(rows, start=min_row):
                    for column_number, value in enumerate(row, start=min_col):
                        source_values[(sheet_name, row_number, column_number)] = value

        for item, expected_values, bounds in range_specs:
            min_col, min_row, max_col, max_row = bounds
            for row_index in range(max_row - min_row + 1):
                expected_row = expected_values[row_index] if row_index < len(expected_values) else []
                for column_index in range(max_col - min_col + 1):
                    expected_value = expected_row[column_index] if column_index < len(expected_row) else None
                    if _is_blank(expected_value):
                        continue
                    actual_value = source_values.get(
                        (item.sheet_name, min_row + row_index, min_col + column_index)
                    )
                    checked += 1
                    is_match = _structure_values_equal(expected_value, actual_value)
                    cell = f"{_column_letter(min_col + column_index)}{min_row + row_index}"
                    checks.append(
                        StructureCheck(
                            item.sheet_name, cell, _display_value(expected_value),
                            _display_value(actual_value), is_match
                        )
                    )
                    if is_match:
                        matched += 1
                    elif len(mismatches) < 8:
                        mismatches.append(
                            f"{item.sheet_name}!{cell}：应为“{_display_value(expected_value)}”，"
                            f"实际“{_display_value(actual_value)}”"
                        )

        if checked == 0:
            return SourceMatch(False, 0.0, 0, 0, (), "模板没有可用的“表结构区域”命名区域，请先维护模板")
        score = matched / checked
        enough = matched >= 3 if checked >= 5 else matched >= 1
        details = f"表结构区域精确匹配 {matched}/{checked}（{score:.0%}）"
        if mismatches:
            details += "；不一致示例：" + "；".join(mismatches)
        if score < STRUCTURE_MATCH_THRESHOLD or not enough:
            return SourceMatch(False, score, matched, checked, (), details, tuple(checks))
        conflicts = sorted(source_names.intersection(external_sheet_names))
        if conflicts:
            return SourceMatch(False, 0.0, matched, checked, (), "外部文件工作表与报送文件重名，不能覆盖原表：" + "、".join(conflicts))
        return SourceMatch(True, score, matched, checked, (), details, tuple(checks))
    finally:
        workbook.close()


def validate_required_sheets_xlsx(
    source_path: Path, definition: TemplateDefinition, *, external_sheet_names: tuple[str, ...] = ()
) -> SourceMatch:
    """Minimal compatibility check used when a flow omits table-structure comparison."""
    workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
    try:
        source_names = set(workbook.sheetnames)
        missing = tuple(sorted({item.sheet_name for item in definition.copy_ranges} - source_names))
        if missing:
            return SourceMatch(False, 0.0, 0, 0, missing, "缺少模板要求的工作表：" + "、".join(missing))
        conflicts = sorted(source_names.intersection(external_sheet_names))
        if conflicts:
            return SourceMatch(False, 0.0, 0, 0, (), "外部文件工作表与报送文件重名，不能覆盖原表：" + "、".join(conflicts))
        return SourceMatch(True, 1.0, 0, 0, (), "已核对必需工作表")
    finally:
        workbook.close()


def _style_report_sheets(workbook: Workbook, sheets: tuple[tuple[Any, list[list[Any]]], ...]) -> None:
    """Shared header styling / autofilter / column width for report workbooks."""
    for sheet, rows in sheets:
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def write_preflight_report_xlsx(path: Path, template_path: Path, template_items: list[PreflightItem]) -> None:
    """Write the check report (模板体检 + 命名区域检查) without starting Excel.

    表结构比对结果已拆分为独立的 ``write_structure_report_xlsx``，不再混入本报告。
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    health_sheet = workbook.active
    health_sheet.title = "模板体检"
    health_rows = [("模板文件", "类别", "级别", "状态", "工作表", "位置", "说明")]
    health_rows.extend(
        (str(template_path), item.category, item.level, item.status,
         item.sheet_name or "—", item.location or "—", item.message)
        for item in template_items
    )
    _style_report_sheets(workbook, ((health_sheet, health_rows),))
    workbook.save(path)


def write_structure_report_xlsx(path: Path, source_matches: list[tuple[Path, SourceMatch]]) -> None:
    """Write the table-structure comparison result (核对类) without starting Excel."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    match_sheet = workbook.active
    match_sheet.title = "文件匹配"
    detail_sheet = workbook.create_sheet("结构比对明细")
    match_rows = [("报送文件", "匹配结果", "匹配率", "匹配标签数", "检查标签数", "缺少工作表", "说明")]
    match_rows.extend(
        (str(source), "通过" if result.matched else "不通过", f"{result.score:.0%}",
         result.matched_labels, result.checked_labels,
         "、".join(result.missing_sheets) or "无", result.details)
        for source, result in source_matches
    )
    if not source_matches:
        match_rows.append(("尚未检查报送文件", "—", "—", "—", "—", "—", "—"))
    detail_rows = [("报送文件", "工作表", "单元格", "模板值", "报送值", "比对结果")]
    for source, result in source_matches:
        for item in result.structure_checks:
            detail_rows.append(
                (
                    str(source), item.sheet_name, item.cell_address,
                    item.expected_value, item.actual_value,
                    "一致" if item.matched else "不一致",
                )
            )
    if len(detail_rows) == 1:
        detail_rows.append(("尚未执行表结构比对", "—", "—", "—", "—", "—"))
    _style_report_sheets(workbook, ((match_sheet, match_rows), (detail_sheet, detail_rows)))
    workbook.save(path)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _structure_values_equal(expected: Any, actual: Any) -> bool:
    if isinstance(expected, bool) or isinstance(actual, bool):
        return type(expected) is type(actual) and expected == actual
    if isinstance(expected, Number) and isinstance(actual, Number):
        return expected == actual
    return type(expected) is type(actual) and expected == actual


def _display_value(value: Any) -> Any:
    return "空" if value is None or value == "" else value


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _compact_batches(
    bounds_list: list[tuple[int, int, int, int]],
) -> list[tuple[int, int, int, int]]:
    batches: list[tuple[int, int, int, int]] = []
    for bounds in sorted(bounds_list, key=lambda item: (item[1], item[0])):
        best_index = None
        best_extra = None
        own_size = _bounds_size(bounds)
        for index, batch in enumerate(batches):
            merged = (
                min(batch[0], bounds[0]), min(batch[1], bounds[1]),
                max(batch[2], bounds[2]), max(batch[3], bounds[3]),
            )
            merged_size = _bounds_size(merged)
            extra = merged_size - _bounds_size(batch) - own_size
            if merged_size <= 10_000 and merged_size <= 4 * (_bounds_size(batch) + own_size):
                if best_extra is None or extra < best_extra:
                    best_index, best_extra = index, extra
        if best_index is None:
            batches.append(bounds)
        else:
            batch = batches[best_index]
            batches[best_index] = (
                min(batch[0], bounds[0]), min(batch[1], bounds[1]),
                max(batch[2], bounds[2]), max(batch[3], bounds[3]),
            )
    return batches


def _bounds_size(bounds: tuple[int, int, int, int]) -> int:
    return (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1)
