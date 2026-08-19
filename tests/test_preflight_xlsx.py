from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from src.base_audit.models import CopyRange, TemplateDefinition
from src.base_audit.preflight_xlsx import validate_source_xlsx, write_structure_report_xlsx


class PreflightXlsxTests(unittest.TestCase):
    def _validate(self, values, ranges):
        folder = TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        source = Path(folder.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "贷款明细"
        for address, value in values.items():
            sheet[address] = value
        workbook.save(source)
        workbook.close()
        definition = TemplateDefinition(
            [], [CopyRange("贷款明细", "A1")], False,
            [item for item, _ in ranges],
        )
        return validate_source_xlsx(source, definition, ranges)

    def test_blank_template_cells_are_skipped_but_numbers_and_codes_are_checked(self):
        area = CopyRange("贷款明细", "A1:D1")
        result = self._validate(
            {"A1": 1001, "B1": "001", "C1": "任意值", "D1": ""},
            [(area, [[1001.0, "001", "   ", None]])],
        )
        self.assertTrue(result.matched)
        self.assertEqual((2, 2), (result.matched_labels, result.checked_labels))

    def test_numeric_and_text_values_with_same_rendering_do_not_match(self):
        area = CopyRange("贷款明细", "A1")
        result = self._validate({"A1": "1001"}, [(area, [[1001]])])
        self.assertFalse(result.matched)
        self.assertEqual(0, result.matched_labels)

    def test_multiple_and_single_cell_areas_are_compared(self):
        first = CopyRange("贷款明细", "A1:B1")
        second = CopyRange("贷款明细", "D5")
        result = self._validate(
            {"A1": "甲", "B1": "乙", "D5": 9},
            [(first, [["甲", "乙"]]), (second, [[9]])],
        )
        self.assertTrue(result.matched)
        self.assertEqual(3, len(result.structure_checks))

    def test_merged_cell_only_compares_nonblank_template_anchor(self):
        folder = TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        source = Path(folder.name) / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "贷款明细"
        sheet.merge_cells("A1:B1")
        sheet["A1"] = "合并表头"
        workbook.save(source)
        workbook.close()
        area = CopyRange("贷款明细", "A1:B1")
        definition = TemplateDefinition([], [CopyRange("贷款明细", "A1")], False, [area])
        result = validate_source_xlsx(source, definition, [(area, [["合并表头", None]])])
        self.assertTrue(result.matched)
        self.assertEqual(1, result.checked_labels)

    def test_structure_details_keep_sheet_name_and_cell_address(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "机构甲.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "贷款明细"
            sheet["A1"], sheet["B1"] = "固定标题", "错误字段"
            workbook.save(source)
            workbook.close()

            structure = CopyRange("贷款明细", "A1:B1")
            definition = TemplateDefinition([], [CopyRange("贷款明细", "A1")], False, [structure])
            result = validate_source_xlsx(source, definition, [(structure, [["固定标题", "正确字段"]])])
            self.assertEqual(2, len(result.structure_checks))
            self.assertEqual("贷款明细", result.structure_checks[1].sheet_name)
            self.assertEqual("B1", result.structure_checks[1].cell_address)
            self.assertFalse(result.structure_checks[1].matched)

            report = root / "表结构比对.xlsx"
            write_structure_report_xlsx(report, [(source, result)])
            saved = load_workbook(report, read_only=True, data_only=True)
            detail = saved["结构比对明细"]
            self.assertEqual("工作表", detail["B1"].value)
            self.assertEqual("贷款明细", detail["B3"].value)
            self.assertEqual("B1", detail["C3"].value)
            saved.close()


if __name__ == "__main__":
    unittest.main()
