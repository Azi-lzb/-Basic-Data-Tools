from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from base_audit.history import (
    HISTORY_HEADERS,
    classify_current_issues,
    merge_history,
    organize_history_rule_numbers,
)
from base_audit.excel_com import ExcelSession
from base_audit.models import Issue


def issue(issue_id: str, period: str, *, triggered: bool = True) -> Issue:
    return Issue(
        issue_id=issue_id,
        period=period,
        batch_id="batch",
        audit_time="2026-01-01 00:00:00",
        triggered=triggered,
        status="新增",
        first_seen_period=period,
        previous_seen_period="",
        consecutive_count=1,
        org_code="001",
        org_name="甲银行",
        report_code="A01",
        sheet_name="基础数据",
        rule_id=issue_id,
        severity="错误",
        formula_cell="H5",
        target_cell="B5",
        target_value=1,
        formula_result="错误",
        message="测试",
        source_file="source.xlsx",
        audit_file="audit.xlsx",
    )


class HistoryTests(unittest.TestCase):
    def test_new_issue(self) -> None:
        current, resolved = classify_current_issues(
            "2026-02", [issue("NEW", "2026-02")], []
        )
        self.assertEqual(current[0].status, "新增")
        self.assertFalse(resolved)

    def test_continuous_issue(self) -> None:
        old = issue("SAME", "2026-01")
        old.consecutive_count = 2
        current, _ = classify_current_issues(
            "2026-02", [issue("SAME", "2026-02")], [old]
        )
        self.assertEqual(current[0].status, "连续出现")
        self.assertEqual(current[0].consecutive_count, 3)

    def test_recurring_issue(self) -> None:
        history = [issue("SAME", "2026-01"), issue("OTHER", "2026-02")]
        current, _ = classify_current_issues(
            "2026-03", [issue("SAME", "2026-03")], history
        )
        self.assertEqual(current[0].status, "再次出现")

    def test_resolved_issue(self) -> None:
        old = issue("OLD", "2026-01")
        _, resolved = classify_current_issues(
            "2026-02",
            [],
            [old],
            audited_org_codes={"001"},
            batch_id="new-batch",
        )
        self.assertEqual(resolved[0].status, "已整改")
        self.assertFalse(resolved[0].triggered)
        self.assertEqual(resolved[0].batch_id, "new-batch")

    def test_rerun_preserves_manual_feedback(self) -> None:
        old = issue("SAME", "2026-02")
        old.institution_feedback = "机构已说明"
        old.auditor_opinion = "下期继续关注"
        new = issue("SAME", "2026-02")
        merged = merge_history([old], [new])
        self.assertEqual(merged[0].institution_feedback, "机构已说明")
        self.assertEqual(merged[0].auditor_opinion, "下期继续关注")

    def test_multiple_history_rows_combine_manual_notes(self) -> None:
        first = issue("SAME", "2026-01")
        first.institution_feedback = "1. 机构说明甲"
        first.auditor_opinion = "1. 审核通过"
        second = issue("SAME", "2026-01")
        second.institution_feedback = "2. 机构说明乙"
        second.auditor_opinion = "2. 补充核实"
        current = issue("SAME", "2026-02")
        merge_history([first, second], [current])
        self.assertEqual(current.institution_feedback, "1. 机构说明甲\n2. 机构说明乙")
        self.assertEqual(current.auditor_opinion, "1. 审核通过\n2. 补充核实")

    def test_current_issue_result_row_uses_compact_columns(self) -> None:
        item = issue("SAME", "2026-02")
        item.comparison_value = "A2"
        item.reference_value = "B4"
        item.difference_value = "A2-B4"
        item.detail = "余额不应小于 0"
        self.assertEqual(
            (
                "source.xlsx", "基础数据", "B5", "错误",
                "A2", "B4", "A2-B4", "余额不应小于 0",
            ),
            ExcelSession._issue_result_row(item),
        )

    def test_history_and_current_result_use_the_same_columns(self) -> None:
        self.assertEqual(
            HISTORY_HEADERS,
            (
                "工作簿名", "工作表名", "定位单元格", "错误类型", "校验指标", "描述",
                "当前值", "对比值", "差值", "规则编号", "历史校验说明", "审核意见",
            ),
        )
        item = issue("规则｜1", "2026-02")
        item.institution_feedback = "1.历史说明"
        item.auditor_opinion = "1.审核通过"
        self.assertEqual(len(ExcelSession._issue_history_row(item)), len(HISTORY_HEADERS))

    def test_organize_history_rule_numbers_removes_legacy_serials(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "config.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "历史审核结果"
            sheet.append(HISTORY_HEADERS)
            prefix = "个人贷款｜贷款信息｜T26｜贷款余额"
            sheet.append(["", "", "", "", "", "", "", "", "问题甲", prefix + "｜1", "说明甲", "意见甲"])
            sheet.append(["", "", "", "", "", "", "", "", "问题乙", prefix + "｜1", "说明乙", "意见乙"])
            sheet.append(["", "", "", "", "", "", "", "", "问题丙", prefix + "｜3", "说明丙", "意见丙"])
            workbook.save(path)
            workbook.close()

            result = organize_history_rule_numbers(path)
            self.assertEqual((result.affected_groups, result.changed_rows), (1, 3))
            check = load_workbook(path, data_only=True)
            try:
                rows = list(check["历史审核结果"].iter_rows(min_row=2, values_only=True))
                self.assertEqual([row[9] for row in rows], [prefix, prefix, prefix])
                self.assertEqual([row[10] for row in rows], ["说明甲", "说明乙", "说明丙"])
            finally:
                check.close()


if __name__ == "__main__":
    unittest.main()
