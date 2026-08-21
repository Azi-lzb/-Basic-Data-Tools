from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from src.base_audit.merge_org import (
    _format_merged_sheet_sources,
    _organisation_from_source,
    _merge_one_org_openpyxl,
    _output_folder,
    _rebase_formula_to_local_sheets,
    _safe_sheet_name,
    _selected_sources,
    _source_period,
    _template_merge_output,
    _write_template_merge_report,
)


def _workbook(path: Path) -> None:
    book = Workbook()
    book.save(path)
    book.close()


def test_organisation_key_is_first_filename_segment() -> None:
    path = Path("TCL科技集团财务有限公司_金融基础数据-单位贷款_B00_2026-07-31_在线核查表.xlsx")
    assert _organisation_from_source(path) == "TCL科技集团财务有限公司"
    assert _source_period(path) == "2026-07-31"


def test_merged_sheet_log_groups_sheets_by_source_workbook() -> None:
    result = _format_merged_sheet_sources((
        "机构A_个人贷款.xlsx｜存量个人贷款信息",
        "机构A_个人贷款.xlsx｜个人客户基础信息",
        "机构A_单位贷款.xlsx｜存量单位贷款信息",
    ))
    assert result == (
        "机构A_个人贷款.xlsx（存量个人贷款信息、个人客户基础信息）\n"
        "机构A_单位贷款.xlsx（存量单位贷款信息）"
    )


def test_sheet_prefix_is_sanitized_and_deduplicated() -> None:
    existing: set[str] = set()
    assert _safe_sheet_name("单位贷款_明细/表", existing) == "单位贷款_明细_表"
    assert _safe_sheet_name("单位贷款_明细/表", existing) == "单位贷款_明细_表_2"
    assert len(_safe_sheet_name("A" * 40, existing)) == 31


def test_empty_selection_means_all_source_files() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        first = root / "机构A_表一_B00_2026-07-31_在线核查表.xlsx"
        second = root / "机构B_表二_B00_2026-07-31_在线核查表.xlsx"
        _workbook(first)
        _workbook(second)
        assert _selected_sources(root, [], recursive=True) == [first, second]


def test_output_folder_uses_requested_prefix_and_rejects_source_root() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        result = _output_folder(root, root / "执行结果", "组合联合核查表")
        assert result.parent == root / "执行结果"
        assert result.name.startswith("组合联合核查表_")
        try:
            _output_folder(root, root, "组合联合核查表")
        except ValueError as exc:
            assert "不能与源数据目录相同" in str(exc)
        else:
            raise AssertionError("源数据目录应被拒绝作为输出目录")


def test_template_merge_always_creates_a_separate_copy_next_to_base() -> None:
    with TemporaryDirectory() as folder:
        base = Path(folder) / "！金融基础数据-单位贷款202605.xlsx"
        personal = Path(folder) / "！金融基础数据-个人贷款202605.xlsx"
        _workbook(base)
        _workbook(personal)
        output, report = _template_merge_output(base, [personal])
        assert output.parent == base.parent / "联合模板"
        assert output != base
        assert output.name.startswith("！联合模板_单位贷款+个人贷款_")
        assert output.suffix == ".xlsx"
        assert report.name.startswith("！联合模板_单位贷款+个人贷款_检查报告_")


def test_template_merge_report_includes_base_template_sheets() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        base = root / "！个人贷款模板.xlsx"
        report = root / "联合模板检查报告.xlsx"
        _workbook(base)
        _write_template_merge_report(
            report,
            base_template=base,
            base_sheets=["集中系统数据", "参照表"],
            sources=[],
            copied=[],
            skipped=[],
            findings=[],
        )
        book = load_workbook(report, read_only=True, data_only=True)
        try:
            rows = list(book["工作表处理清单"].iter_rows(min_row=2, values_only=True))
            assert (base.name, "集中系统数据", "保留为底稿") in rows
            assert (base.name, "参照表", "保留为底稿") in rows
        finally:
            book.close()


def test_every_external_formula_is_rebased_to_a_local_sheet_reference() -> None:
    original = (
        "=SUMIFS('C:\\模板\\[个人贷款.xlsx]个人贷款发生额信息'!$E:$E,"
        "[1]个人贷款发生额信息!$D:$D,$B5)"
    )
    result = _rebase_formula_to_local_sheets(original)
    assert result == "=SUMIFS('个人贷款发生额信息'!$E:$E,'个人贷款发生额信息'!$D:$D,$B5)"


def test_openpyxl_merge_keeps_all_sheets_and_formulas() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        first = root / "机构A_个人贷款_B00_2026-07-31_在线核查表.xlsx"
        second = root / "机构A_单位贷款_B00_2026-07-31_在线核查表.xlsx"
        for path, sheet_name, hidden in ((first, "个人明细", False), (second, "单位明细", True)):
            book = Workbook()
            sheet = book.active
            sheet.title = sheet_name
            sheet["A1"] = "指标"
            sheet["B1"] = "金额"
            sheet["B2"] = "=1+1"
            sheet["A1"].font = Font(bold=True, color="FF0000")
            sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
            sheet.column_dimensions["A"].width = 22
            sheet.row_dimensions[1].height = 24
            sheet.freeze_panes = "A2"
            sheet.merge_cells("A4:B4")
            if hidden:
                book.create_sheet("封面")
                sheet.sheet_state = "hidden"
            book.save(path)
            book.close()
        output = root / "机构A_合并_2026-07-31.xlsx"
        merged_sheets: list[str] = []
        assert _merge_one_org_openpyxl([first, second], output, merged_sheets) == 3
        assert f"{first.name}｜个人明细" in merged_sheets
        assert f"{second.name}｜单位明细" in merged_sheets
        merged = load_workbook(output, read_only=False, data_only=False)
        try:
            assert len(merged.worksheets) == 3
            assert merged["个人明细"]["B2"].value == "=1+1"
            assert "A4:B4" in {str(item) for item in merged["个人明细"].merged_cells.ranges}
            assert merged["个人明细"]["A1"].font.bold is True
            assert merged["个人明细"]["A1"].fill.fgColor.rgb == "00FFFF00"
            assert merged["个人明细"].column_dimensions["A"].width == 22
            assert merged["个人明细"].row_dimensions[1].height == 24
            assert merged["个人明细"].freeze_panes == "A2"
            assert merged["单位明细"].sheet_state == "hidden"
        finally:
            merged.close()


def test_openpyxl_merge_skips_later_duplicate_sheet_names() -> None:
    with TemporaryDirectory() as folder:
        root = Path(folder)
        first = root / "机构A_个人贷款_B00_2026-07-31_在线核查表.xlsx"
        second = root / "机构A_单位贷款_B00_2026-07-31_在线核查表.xlsx"
        for path, value in ((first, "首份"), (second, "后份")):
            book = Workbook()
            book.active.title = "集中系统数据"
            book.active["A1"] = value
            book.save(path)
            book.close()
        output = root / "机构A_合并_2026-07-31.xlsx"
        assert _merge_one_org_openpyxl([first, second], output) == 1
        merged = load_workbook(output, read_only=True)
        try:
            assert merged.sheetnames == ["集中系统数据"]
            assert merged["集中系统数据"]["A1"].value == "首份"
        finally:
            merged.close()
