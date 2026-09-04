"""WPS OOXML conditional-format fallback (expression + cellIs).

The pywebview2 Windows edition reads a rendered conditional-format colour
through Excel/WPS ``DisplayFormat``.  WPS 12.0 returns ``None``, so the session
falls back to evaluating the submission's own OOXML rules.  These tests cover
``AND(C5>0,C5>1)`` relative-reference translation (C5/D5/E5/C26), destination-
relative descriptions, and the "don't guess" behaviour for unsupported rules.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from src.base_audit.conditional_format import evaluate_expression_formula
from src.base_audit.excel_com import ExcelSession
from src.base_audit.name_config import CONDITIONAL_FORMAT_EXTRACT_FUNCTION, FeatureMapping


def test_expression_and_relative_reference_conversion() -> None:
    wb = Workbook()
    ws = wb.active
    ws["C5"] = 2     # >1  -> True
    ws["D5"] = 0     # not >0 -> False (would be True if wrongly fixed on C5)
    ws["E5"] = 0.5   # >0 but not >1 -> False
    ws["C26"] = 3    # >1 -> True (proves the +21 row translation)
    formula = "AND(C5>0,C5>1)"

    results = {
        "C5": evaluate_expression_formula(formula, ws, 5, 3, 5, 3),
        "D5": evaluate_expression_formula(formula, ws, 5, 3, 5, 4),
        "E5": evaluate_expression_formula(formula, ws, 5, 3, 5, 5),
        "C26": evaluate_expression_formula(formula, ws, 5, 3, 26, 3),
    }
    wb.close()
    assert results == {"C5": True, "D5": False, "E5": False, "C26": True}


def _mapping() -> FeatureMapping:
    return FeatureMapping(
        "条件格式结果提取", CONDITIONAL_FORMAT_EXTRACT_FUNCTION,
        ("条件格式区域",), False, "", "",
    )


def _absolute(area: str) -> str:
    def part(text: str) -> str:
        letters = "".join(ch for ch in text if ch.isalpha())
        digits = "".join(ch for ch in text if ch.isdigit())
        return "${}${}".format(letters, digits)

    if ":" in area:
        start, end = area.split(":", 1)
        return "{}:{}".format(part(start), part(end))
    return part(area)


def _expression_workbook(path: Path, cells: dict[str, object], formula: str, area: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "报表"
    for address, value in cells.items():
        ws[address] = value
    fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    ws.conditional_formatting.add(area, FormulaRule(formula=[formula], fill=fill))
    wb.defined_names.add(DefinedName("条件格式区域", attr_text="'报表'!" + _absolute(area)))
    wb.save(path)
    wb.close()
    return path


def test_fallback_expression_relative_description_and_trigger(tmp_path: Path) -> None:
    path = _expression_workbook(
        tmp_path / "表达式.xlsx",
        {"C5": 2, "D5": 3, "E5": 4, "C26": 5},
        "AND(C5>0,C5>1)", "C5:E26",
    )
    issues, unsupported = ExcelSession()._extract_conditional_format_issues_from_rules(
        workbook_path=path, mapping=_mapping(), structure_ranges=[],
        period="2026-08", batch_id="t", audit_time="2026-08-01 00:00:00",
        org_code="", org_name="", source_file=path,
    )
    assert unsupported == 0
    by_cell = {item.target_cell: item.detail for item in issues}
    assert set(by_cell) == {"C5", "D5", "E5", "C26"}
    assert by_cell["C5"] == "条件格式规则：AND(C5>0,C5>1)"
    assert by_cell["D5"] == "条件格式规则：AND(D5>0,D5>1)"
    assert by_cell["E5"] == "条件格式规则：AND(E5>0,E5>1)"
    assert by_cell["C26"] == "条件格式规则：AND(C26>0,C26>1)"


def test_fallback_unsupported_rule_does_not_trigger(tmp_path: Path) -> None:
    path = _expression_workbook(
        tmp_path / "不支持.xlsx",
        {"C5": 1, "D5": 2},
        'AND(C5>0,INDIRECT("A1"))', "C5:D5",
    )
    issues, unsupported = ExcelSession()._extract_conditional_format_issues_from_rules(
        workbook_path=path, mapping=_mapping(), structure_ranges=[],
        period="2026-08", batch_id="t", audit_time="2026-08-01 00:00:00",
        org_code="", org_name="", source_file=path,
    )
    assert issues == []
    assert unsupported == 1


def test_fallback_scans_applies_to_without_named_range(tmp_path: Path) -> None:
    """Without a 条件格式区域 name, fall back to every sheet's applies-to."""
    path = tmp_path / "无命名区域.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "报表"
    ws["C5"] = 2
    ws["C6"] = 0
    fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    ws.conditional_formatting.add(
        "C5:C6", CellIsRule(operator="greaterThan", formula=["1"], fill=fill)
    )
    wb.save(path)
    wb.close()

    issues, unsupported = ExcelSession()._extract_conditional_format_issues_from_rules(
        workbook_path=path, mapping=_mapping(), structure_ranges=[],
        period="2026-08", batch_id="t", audit_time="2026-08-01 00:00:00",
        org_code="", org_name="", source_file=path,
    )
    assert unsupported == 0
    assert [item.target_cell for item in issues] == ["C5"]


def test_fallback_keeps_cellis_handling(tmp_path: Path) -> None:
    path = tmp_path / "cellIs.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "报表"
    ws["C5"] = 2
    ws["C6"] = 0
    fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    ws.conditional_formatting.add(
        "C5:C6", CellIsRule(operator="greaterThan", formula=["1"], fill=fill)
    )
    wb.defined_names.add(DefinedName("条件格式区域", attr_text="'报表'!$C$5:$C$6"))
    wb.save(path)
    wb.close()

    issues, unsupported = ExcelSession()._extract_conditional_format_issues_from_rules(
        workbook_path=path, mapping=_mapping(), structure_ranges=[],
        period="2026-08", batch_id="t", audit_time="2026-08-01 00:00:00",
        org_code="", org_name="", source_file=path,
    )
    assert unsupported == 0
    assert [item.target_cell for item in issues] == ["C5"]
    assert issues[0].detail == "条件格式规则：C5>1"
