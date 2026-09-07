"""基础数据季报表跨期比较（第一期：两期比较 + 大集中核对）。

对齐旧 VBA `基础数据季报表跨期比较.bas` 的核心口径：
- 输入为“当期目录 + 上期目录”，支持三种系统导出方式（自动识别）：
  1. 逐机构逐表：`机构代码#日期#01#表单#机构名.xls`，sheet=表单；
  2. 按报表划分：`banks#日期#01#表单.xls`，sheet 名=机构代码；
  3. 按机构划分：`reports#机构代码#日期#01#机构名.xls`，sheet=表单。
  数据工作表统一为：第 1 列指标编号、第 2 列指标名称、第 3 列值，第 4 行起。
- 指标/机构/警戒区间/校验规则配置放在独立工作簿 `报表采集系统_比较配置.xlsx`
  （用户维护、程序只读）。新格式为 4 表：指标参照、机构参照、警戒区间、
  校验规则（表头名驱动读取）；旧 5 表格式（特殊指标-自定义、复杂校验-自定义）
  自动识别并兼容。
- 大集中核对：读用户选择的大集中工作簿（集中系统数据 + 参照表），按
  机构代码→报表项目、指标代码→大集中指标名称映射，差异绝对值超过容差
  （默认 0.01 万元，即 100 元，可在设置中心调整）记“差异超过100元”，
  单边有值记“谨慎核实”。

纯 openpyxl/xlrd 实现，不启动 Excel/WPS/LibreOffice，Windows 与 UOS 通用。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook

# —— 口径常量（与 VBA 一致） ——
SOURCE_UNIT = "元"
DEST_UNIT = "万元"
UNIT_FACTOR = 1.0 / 10000.0          # 元 → 万元
CENTRAL_DIFF_TOLERANCE = 0.01        # 万元；差异绝对值超过它记“差异超过100元”
CENTRAL_UNIT_FACTOR = 10000.0        # 集中系统数据单位为亿元 → 万元
INDICATOR_NAME_CODE = "20201001"     # 按报表划分导出里取机构名称的指标编号
INDICATOR_CODE_CODE = "20201002"     # 同上，机构代码

PERIOD_SHEET_HEADERS = [
    "地区", "数据属性", "机构名称", "指标编码", "指标名称", "当期数", "上期数",
    "变动绝对值", "环比变动", "备注", "是否说明", "级别", "计算过程", "机构类别", "承接行",
    "社会信用代码", "数据日期", "币种", "频度",
]
CENTRAL_SHEET_HEADERS = [
    "地区", "数据属性", "机构名称", "指标编码", "指标名称", "基础数据值",
    "大集中值", "差异绝对值", "差异幅度", "是否说明",
]

CONFIG_WORKBOOK_NAME = "报表采集系统_比较配置.xlsx"
LEGACY_CONFIG_WORKBOOK_NAME = "跨期比较配置.xlsx"
CONFIG_INDICATOR_SHEET = "指标参照"
CONFIG_ORG_SHEET = "机构参照"
CONFIG_ALERT_SHEET = "警戒区间"
CONFIG_SPECIAL_SHEET = "特殊指标-自定义"
CONFIG_COMPLEX_SHEET = "复杂校验-自定义"
CONFIG_RULE_SHEET = "校验规则"       # 新格式统一规则表；存在即按新格式读取

# 统一校验规则的类型与级别（参考监管报表校验惯例：勾稽/阈值守恒用错误、
# 监管线与占比阈值用核实、趋势与提示类用提示）。
RULE_TYPE_EXPRESSION = "表达式"
RULE_TYPE_ACC_YEAR = "累计不降(当年)"
RULE_TYPE_ACC_HISTORY = "累计不降(历史)"
RULE_TYPES = (RULE_TYPE_EXPRESSION, RULE_TYPE_ACC_YEAR, RULE_TYPE_ACC_HISTORY)
RULE_LEVELS = ("错误", "核实", "提示")
DEFAULT_RULE_LEVEL = "提示"


class PeriodCompareError(RuntimeError):
    """跨期比较的明确业务失败（目录为空、文件名不合法、配置缺失等）。"""


@dataclass
class IndicatorDef:
    code: str
    name: str = ""
    data_type: str = ""            # 数据属性：余额/累发/个数/百分数…
    no_unit_convert: bool = False  # 是否不转换单位
    check_central: bool = False    # 是否与大集中核对
    central_name: str = ""         # 大集中报表查询指标名称


@dataclass
class OrgDef:
    name: str
    code: str = ""
    org_class: str = ""            # 机构类别
    bank_row: str = ""             # 承接行
    region: str = ""               # 地区/归属行
    report_item: str = ""          # 报表项目（大集中行标识）


@dataclass
class AlertRange:
    lower_bound: float    # 变幅下限（环比小数，0.3 = 30%）
    remark: str
    fill_row: bool = False


@dataclass
class SpecialRule:
    """特殊指标规则（对齐 VBA“特殊指标-自定义”）。

    第一期实现备注文本为“当年累计/历史累计指标比上期不应减少”的规则；
    其他 VBA 规则类型读取保留但跳过执行。
    """

    code: str
    name: str
    remark: str           # 备注文本，同时是规则类型标识与提示内容
    rmb_threshold: float | None = None   # 人民币阀值（亿元）
    usd_threshold: float | None = None
    data_type: str = ""
    explain: str = ""     # 说明/比较指标代码
    form: str = ""
    detailed: str = ""
    change_pct: float | None = None      # 绝对值变幅（%）
    disabled: bool = False

    @property
    def rule_kind(self) -> str:
        text = self.remark.replace("。", "").replace("：", ":").strip()
        if text.startswith("当年累计指标比上期不应减少"):
            return "当年累计不应减少"
        if text.startswith("历史累计指标比上期不应减少"):
            return "历史累计不应减少"
        return "未支持"


@dataclass
class ComplexRule:
    """复杂校验表达式规则（旧 5 表格式“复杂校验-自定义”，仅兼容保留）。

    表达式语法：`[机构,地区,类别,指标代码,数据属性,币种,频度,批次]` 为当期值，
    `{...}` 同结构为上期值，解析取第 4 段（下标 3）指标代码。新格式直接写
    `[指标代码]`（当期）或 `{指标代码}`（上期），也可写指标名称如
    `[一级资本净额]`。支持 Excel 风格 AND/OR/NOT（大小写均可）与 `<>`。
    """

    form: str
    desc: str
    rule: str
    invert: bool = False
    threshold: float | None = None
    disabled: bool = False
    note: str = ""


@dataclass
class CheckRule:
    """统一校验规则（新格式“校验规则”表的行）。

    type 决定 content 的解释：
    - ``表达式``：content 为条件表达式，真即命中；
    - ``累计不降(当年)`` / ``累计不降(历史)``：content 为指标代码清单
      （逗号/顿号/空白分隔），当期值低于上期值即命中；“当年”跨年跳过。
    level 为命中后的“级别”列取值（错误/核实/提示）。
    """

    rule_id: str = ""
    type: str = RULE_TYPE_EXPRESSION
    desc: str = ""
    content: str = ""
    level: str = DEFAULT_RULE_LEVEL
    disabled: bool = False
    note: str = ""


@dataclass
class PeriodConfig:
    indicators: dict[str, IndicatorDef] = field(default_factory=dict)
    orgs: dict[str, OrgDef] = field(default_factory=dict)      # key=机构名称
    orgs_by_code: dict[str, OrgDef] = field(default_factory=dict)
    alerts: list[AlertRange] = field(default_factory=list)
    specials: list[SpecialRule] = field(default_factory=list)      # 旧格式
    complex_rules: list[ComplexRule] = field(default_factory=list)  # 旧格式
    rules: list[CheckRule] = field(default_factory=list)           # 新格式
    # 加载期发现的问题（未知类型、非法级别、表达式语法错误、无法解析的
    # 指标名称引用等）；由调用方写入运行日志，不再静默忽略。
    warnings: list[str] = field(default_factory=list)

    def effective_rules(self) -> list[CheckRule]:
        """统一规则视图：新格式直接返回；旧格式把两张自定义表转换过来。"""
        if self.rules:
            return list(self.rules)
        converted: list[CheckRule] = []
        for item in self.specials:
            if item.rule_kind == "当年累计不应减少":
                rule_type = RULE_TYPE_ACC_YEAR
            elif item.rule_kind == "历史累计不应减少":
                rule_type = RULE_TYPE_ACC_HISTORY
            else:
                continue  # 未支持类型沿用旧行为：不执行
            converted.append(CheckRule(
                type=rule_type, desc=item.remark, content=item.code,
                level=DEFAULT_RULE_LEVEL, disabled=item.disabled,
            ))
        for item in self.complex_rules:
            if item.invert:
                continue  # 取反规则（VBA qfbz）从未执行过，保持跳过
            converted.append(CheckRule(
                type=RULE_TYPE_EXPRESSION, desc=item.desc, content=item.rule,
                level=DEFAULT_RULE_LEVEL, disabled=item.disabled, note=item.note,
            ))
        return converted

    def alert_remark(self, change_pct: float | None) -> str:
        """按 VBA AddbfMark 口径匹配警戒区间。

        配置中的“变幅下限”是小数（0.3 表示 30%）；环比变动以百分数传入
        （41.56 表示 +41.56%）。从大到小找第一个 ``环比% >= 下限×100`` 的
        档位取备注；恰好 -99% 记“缩小100倍及以上”，低于最小档记
        “近缩小100倍及以上”。
        """
        if change_pct is None or round(change_pct, 6) == 0:
            return ""
        matched: str | None = None
        for item in reversed(self.alerts):
            if round(change_pct - item.lower_bound * 100.0, 10) >= 0:
                matched = item.remark
                if round(change_pct - (-99.0), 10) == 0:
                    matched = "缩小100倍及以上"
                break
        if matched is None:
            return "近缩小100倍及以上"
        return matched


@dataclass
class IndicatorValue:
    org_name: str
    org_code: str
    code: str
    name: str
    value: Any            # float 或 str（文字型指标）
    date: str
    form: str             # 表单号 20201/20202/20203


# ---------------------------------------------------------------------------
# 配置读取
# ---------------------------------------------------------------------------

def _cell_bool(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip() in {"是", "true", "1", "y", "yes"}


def _optional_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value: Any) -> str:
    return str(value or "").strip()


# ---------------------------------------------------------------------------
# 新格式（4 表，表头名驱动）读取
# ---------------------------------------------------------------------------

# 表头别名表：新表头与旧表头名映射到同一字段，列顺序不再影响读取。
INDICATOR_HEADER_ALIASES = {
    "指标代码": "code", "指标名称": "name", "数据属性": "data_type",
    "不转换单位": "no_unit_convert", "是否不转换单位": "no_unit_convert",
    "大集中核对": "check_central", "是否与大集中核对": "check_central",
    "大集中指标名称": "central_name", "大集中报表查询指标名称": "central_name",
    "禁用": "disabled", "备注": "note",
}
ORG_HEADER_ALIASES = {
    "机构名称": "name", "社会信用代码": "code", "机构类别": "org_class",
    "承接行": "bank_row", "地区": "region", "归属行": "region",
    "报表项目": "report_item", "禁用": "disabled",
}
ALERT_HEADER_ALIASES = {
    "变幅下限": "lower", "变幅下限（小数，0.3=30%）": "lower",
    "备注": "remark", "整行填充": "fill_row", "是否整行填充": "fill_row",
}
RULE_HEADER_ALIASES = {
    "规则编号": "rule_id", "类型": "type", "描述": "desc",
    "规则内容": "content", "级别": "level", "禁用": "disabled", "备注": "note",
}


def _header_map(header_row: Iterable[Any], aliases: dict[str, str]) -> dict[str, int]:
    """表头文字 → 列下标；未知表头忽略，同一字段取先出现的列。"""
    result: dict[str, int] = {}
    for index, value in enumerate(header_row):
        field = aliases.get(str(value or "").strip())
        if field and field not in result:
            result[field] = index
    return result


def _cell(row: tuple[Any, ...], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def _sheet_rows(sheet, aliases: dict[str, str]) -> tuple[dict[str, int], list[tuple[Any, ...]]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}, []
    return _header_map(rows[0], aliases), rows[1:]


def _load_indicator_sheet(book, config: PeriodConfig) -> None:
    if CONFIG_INDICATOR_SHEET not in book.sheetnames:
        return
    mapping, rows = _sheet_rows(book[CONFIG_INDICATOR_SHEET], INDICATOR_HEADER_ALIASES)
    if "code" not in mapping:
        config.warnings.append(f"“{CONFIG_INDICATOR_SHEET}”缺少“指标代码”表头，该表未加载")
        return
    for row in rows:
        code = _norm_code(_cell(row, mapping, "code"))
        if not code or _cell_bool(_cell(row, mapping, "disabled")):
            continue
        config.indicators[code] = IndicatorDef(
            code=code,
            name=_text(_cell(row, mapping, "name")),
            data_type=_text(_cell(row, mapping, "data_type")),
            no_unit_convert=_cell_bool(_cell(row, mapping, "no_unit_convert")),
            check_central=_cell_bool(_cell(row, mapping, "check_central")),
            central_name=_text(_cell(row, mapping, "central_name")),
        )


def _load_org_sheet(book, config: PeriodConfig) -> None:
    if CONFIG_ORG_SHEET not in book.sheetnames:
        return
    mapping, rows = _sheet_rows(book[CONFIG_ORG_SHEET], ORG_HEADER_ALIASES)
    if "name" not in mapping:
        config.warnings.append(f"“{CONFIG_ORG_SHEET}”缺少“机构名称”表头，该表未加载")
        return
    for row in rows:
        name = _text(_cell(row, mapping, "name"))
        if not name or _cell_bool(_cell(row, mapping, "disabled")):
            continue
        org = OrgDef(
            name=name,
            code=_text(_cell(row, mapping, "code")),
            org_class=_text(_cell(row, mapping, "org_class")),
            bank_row=_text(_cell(row, mapping, "bank_row")),
            region=_text(_cell(row, mapping, "region")),
            report_item=_text(_cell(row, mapping, "report_item")),
        )
        config.orgs[name] = org
        if org.code:
            config.orgs_by_code[org.code] = org


def _load_alert_sheet(book, config: PeriodConfig) -> None:
    if CONFIG_ALERT_SHEET not in book.sheetnames:
        return
    mapping, rows = _sheet_rows(book[CONFIG_ALERT_SHEET], ALERT_HEADER_ALIASES)
    if "lower" not in mapping:
        config.warnings.append(f"“{CONFIG_ALERT_SHEET}”缺少“变幅下限”表头，该表未加载")
        return
    for row in rows:
        try:
            lower = float(_cell(row, mapping, "lower"))
        except (TypeError, ValueError):
            continue
        config.alerts.append(AlertRange(
            lower_bound=lower,
            remark=_text(_cell(row, mapping, "remark")),
            fill_row=_cell_bool(_cell(row, mapping, "fill_row")),
        ))


def _indicator_name_lookup(config: PeriodConfig) -> tuple[dict[str, str], list[str]]:
    """指标名称 → 代码；重名名称从查找表中剔除（不可按名称引用）。"""
    lookup: dict[str, str] = {}
    duplicated: set[str] = set()
    for code, indicator in config.indicators.items():
        if not indicator.name:
            continue
        if indicator.name in lookup:
            duplicated.add(indicator.name)
        else:
            lookup[indicator.name] = code
    for name in duplicated:
        lookup.pop(name, None)
    return lookup, sorted(duplicated)


def _load_rule_sheet(book, config: PeriodConfig) -> None:
    if CONFIG_RULE_SHEET not in book.sheetnames:
        return
    mapping, rows = _sheet_rows(book[CONFIG_RULE_SHEET], RULE_HEADER_ALIASES)
    if "content" not in mapping:
        config.warnings.append(f"“{CONFIG_RULE_SHEET}”缺少“规则内容”表头，该表未加载")
        return
    for row in rows:
        content = _text(_cell(row, mapping, "content"))
        if not content:
            continue
        rule_id = _text(_cell(row, mapping, "rule_id"))
        desc = _text(_cell(row, mapping, "desc"))
        label = f"规则 {rule_id}（{desc or content[:24]}）" if rule_id else f"规则（{desc or content[:24]}）"
        rule_type = _text(_cell(row, mapping, "type")) or RULE_TYPE_EXPRESSION
        if rule_type not in RULE_TYPES:
            config.warnings.append(f"{label}：未知类型“{rule_type}”（支持 {'、'.join(RULE_TYPES)}），已跳过")
            continue
        level = _text(_cell(row, mapping, "level")) or DEFAULT_RULE_LEVEL
        if level not in RULE_LEVELS:
            config.warnings.append(
                f"{label}：级别“{level}”无效（应为 {'/'.join(RULE_LEVELS)}），按“{DEFAULT_RULE_LEVEL}”处理"
            )
            level = DEFAULT_RULE_LEVEL
        config.rules.append(CheckRule(
            rule_id=rule_id,
            type=rule_type,
            desc=desc or content,
            content=content,
            level=level,
            disabled=_cell_bool(_cell(row, mapping, "disabled")),
            note=_text(_cell(row, mapping, "note")),
        ))
    _validate_rules(config)


def _validate_rules(config: PeriodConfig) -> None:
    """加载期校验：名称引用核对 + 表达式全 0 试编译，问题写入 warnings。"""
    lookup, duplicated = _indicator_name_lookup(config)

    def resolver(name: str) -> str:
        return lookup[name]

    for name in duplicated:
        config.warnings.append(f"指标名称“{name}”在指标参照中重复，校验规则中不能按名称引用它")
    for rule in config.rules:
        if rule.disabled or rule.type != RULE_TYPE_EXPRESSION:
            continue
        label = f"规则 {rule.rule_id}（{rule.desc}）" if rule.rule_id else f"规则（{rule.desc}）"
        codes, errors = _expression_codes(rule.content, resolver)
        for error in errors:
            config.warnings.append(f"{label}：{error}")
        for code in sorted(codes):
            if code not in config.indicators:
                config.warnings.append(f"{label}：引用的指标代码 {code} 不在指标参照中，请核实")
        _filled, render_errors = _render_expression(
            rule.content, lambda _code: 0.0, lambda _code: 0.0, resolver
        )
        if render_errors:
            continue  # 名称/占位符问题已在上面记录
        message = _compile_translated(_filled)
        if message:
            config.warnings.append(f"{label}：表达式语法错误（{message}），执行时将跳过")


def load_period_config(config_path: Path) -> PeriodConfig:
    if not config_path.is_file():
        raise PeriodCompareError(
            f"未找到比较配置：{config_path}。请在程序目录放置{CONFIG_WORKBOOK_NAME}"
            "（指标参照/机构参照/警戒区间/校验规则），参照说明维护后重试。"
        )
    config = PeriodConfig()
    book = load_workbook(config_path, read_only=True, data_only=True)
    try:
        if CONFIG_RULE_SHEET in book.sheetnames:
            _load_indicator_sheet(book, config)
            _load_org_sheet(book, config)
            _load_alert_sheet(book, config)
            _load_rule_sheet(book, config)
        else:
            _load_legacy_config(book, config)
    finally:
        book.close()
    config.alerts.sort(key=lambda item: item.lower_bound)
    return config


def _load_legacy_config(book, config: PeriodConfig) -> None:
    """旧 5 表格式：按列位置读取（用户手工绑定的旧配置簿继续可用）。"""
    if CONFIG_INDICATOR_SHEET in book.sheetnames:
        for row in book[CONFIG_INDICATOR_SHEET].iter_rows(min_row=2, values_only=True):
            code = _norm_code(row[0]) if row and row[0] is not None else ""
            if not code or (len(row) > 6 and _cell_bool(row[6])):
                continue
            config.indicators[code] = IndicatorDef(
                code=code,
                name=str(row[1] or "").strip() if len(row) > 1 else "",
                data_type=str(row[2] or "").strip() if len(row) > 2 else "",
                no_unit_convert=_cell_bool(row[3]) if len(row) > 3 else False,
                check_central=_cell_bool(row[4]) if len(row) > 4 else False,
                central_name=str(row[5] or "").strip() if len(row) > 5 else "",
            )
    if CONFIG_ORG_SHEET in book.sheetnames:
        for row in book[CONFIG_ORG_SHEET].iter_rows(min_row=2, values_only=True):
            name = str(row[0] or "").strip() if row and row[0] else ""
            if not name or (len(row) > 6 and _cell_bool(row[6])):
                continue
            org = OrgDef(
                name=name,
                code=str(row[1] or "").strip() if len(row) > 1 else "",
                org_class=str(row[2] or "").strip() if len(row) > 2 else "",
                bank_row=str(row[3] or "").strip() if len(row) > 3 else "",
                region=str(row[4] or "").strip() if len(row) > 4 else "",
                report_item=str(row[5] or "").strip() if len(row) > 5 else "",
            )
            config.orgs[name] = org
            if org.code:
                config.orgs_by_code[org.code] = org
    if CONFIG_ALERT_SHEET in book.sheetnames:
        for row in book[CONFIG_ALERT_SHEET].iter_rows(min_row=2, values_only=True):
            if not row or row[1] is None or row[1] == "":
                continue
            try:
                lower = float(row[1])
            except (TypeError, ValueError):
                continue
            config.alerts.append(
                AlertRange(
                    lower_bound=lower,
                    remark=str(row[2]).strip() if row[2] is not None else "",
                    fill_row=_cell_bool(row[3]) if len(row) > 3 else False,
                )
            )
    if CONFIG_SPECIAL_SHEET in book.sheetnames:
        for row in book[CONFIG_SPECIAL_SHEET].iter_rows(min_row=2, values_only=True):
            code = _norm_code(row[0]) if row and row[0] is not None else ""
            remark = str(row[2] or "").strip() if len(row) > 2 else ""
            if not code or not remark:
                continue
            config.specials.append(SpecialRule(
                code=code,
                name=str(row[1] or "").strip(),
                remark=remark,
                rmb_threshold=_optional_float(row[3]) if len(row) > 3 else None,
                usd_threshold=_optional_float(row[4]) if len(row) > 4 else None,
                data_type=str(row[5] or "").strip() if len(row) > 5 else "",
                explain=_norm_code(row[6]) if len(row) > 6 else "",
                form=_norm_code(row[7]) if len(row) > 7 else "",
                detailed=str(row[8] or "").strip() if len(row) > 8 else "",
                change_pct=_optional_float(row[9]) if len(row) > 9 else None,
                disabled=_cell_bool(row[10]) if len(row) > 10 else False,
            ))
    if CONFIG_COMPLEX_SHEET in book.sheetnames:
        for row in book[CONFIG_COMPLEX_SHEET].iter_rows(min_row=2, values_only=True):
            rule = str(row[2] or "").strip() if len(row) > 2 else ""
            desc = str(row[1] or "").strip() if len(row) > 1 else ""
            if not rule or not desc:
                continue
            config.complex_rules.append(ComplexRule(
                form=str(row[0] or "").strip(),
                desc=desc,
                rule=rule,
                invert=_cell_bool(row[3]) if len(row) > 3 else False,
                threshold=_optional_float(row[4]) if len(row) > 4 else None,
                disabled=_cell_bool(row[5]) if len(row) > 5 else False,
                note=str(row[6] or "").strip() if len(row) > 6 else "",
            ))


def ensure_default_config(config_path: Path) -> bool:
    """配置簿缺失时生成新版 4 表模板（含禁用状态的示例规则）；返回是否新建。

    旧文件名「跨期比较配置.xlsx」若与目标同目录存在，则直接改名沿用，
    用户自维护内容不丢失。
    """
    if config_path.is_file():
        return False
    legacy = config_path.with_name(LEGACY_CONFIG_WORKBOOK_NAME)
    if legacy.is_file():
        try:
            legacy.replace(config_path)
            return False
        except PermissionError:
            pass  # 旧文件被占用时退回生成默认配置，不影响执行
    book = Workbook()
    sheet = book.active
    sheet.title = CONFIG_INDICATOR_SHEET
    sheet.append(["指标代码", "指标名称", "数据属性", "不转换单位", "大集中核对", "大集中指标名称", "禁用", "备注"])
    org_sheet = book.create_sheet(CONFIG_ORG_SHEET)
    org_sheet.append(["机构名称", "社会信用代码", "机构类别", "承接行", "地区", "报表项目", "禁用"])
    alert_sheet = book.create_sheet(CONFIG_ALERT_SHEET)
    alert_sheet.append(["变幅下限（小数，0.3=30%）", "备注", "整行填充"])
    default_alerts = [
        (-0.96, "降幅(-96%,-90%],缩小10倍-100倍", "是"),
        (-0.9, "降幅(-90%,-80%],缩小5倍-10倍", ""),
        (-0.8, "降幅(-80%,-50%],缩小1倍-5倍", ""),
        (-0.5, "降幅(-50%,-30%],缩小1倍以内", ""),
        (-0.3, "", ""),
        (0.0, "", ""),
        (0.3, "增幅[30%,50%)", ""),
        (0.5, "增幅[50%,1倍)", ""),
        (1.0, "增幅[1倍,5倍)", ""),
        (5.0, "增幅[5倍,10倍)", ""),
        (10.0, "增幅[10倍,96倍)", "是"),
        (96.0, "近增幅100倍以上，请核实", "是"),
    ]
    for row in default_alerts:
        alert_sheet.append(row)
    rule_sheet = book.create_sheet(CONFIG_RULE_SHEET)
    rule_sheet.append(["规则编号", "类型", "描述", "规则内容", "级别", "禁用", "备注"])
    rule_sheet.append([
        "R001", RULE_TYPE_EXPRESSION, "示例：资产负债表不平衡",
        "[20202003] <> [20202004] + [20202005]", "错误", "是",
        "示例行：确认公式后把“禁用”清空即可启用；[代码]为当期值，{代码}为上期值，也可写指标名称",
    ])
    rule_sheet.append([
        "R002", RULE_TYPE_ACC_YEAR, "当年累计指标比上期不应减少。",
        "20203003", "提示", "是",
        "示例行：规则内容填指标代码清单，可用逗号/顿号分隔多个代码",
    ])
    config_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(config_path)
    return True


# ---------------------------------------------------------------------------
# 两期数据读取（三种导出方式归一化）
# ---------------------------------------------------------------------------

def _norm_code(value: Any) -> str:
    """指标编号统一为整数字符串（xlrd 会把 20201001 读成 float）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _read_sheet_rows(sheet_rows: Iterable[Iterable[Any]]) -> list[tuple[str, str, Any]]:
    """按第 4 行起读 (指标编号, 指标名称, 值)；前 3 行为标题/空行/表头。"""
    rows: list[tuple[str, str, Any]] = []
    for index, row in enumerate(sheet_rows):
        if index < 3:
            continue
        values = list(row) + [None, None, None]
        code = _norm_code(values[0])
        if not code:
            continue
        rows.append((code, str(values[1] or "").strip(), values[2]))
    return rows


def _form_from_sheet(sheet_name: str, rows: list[tuple[str, str, Any]], fallback: str) -> str:
    """从子工作表名称或指标编码取得稳定的报表标识。

    按机构划分的导出常把子工作表直接命名为报表名称，而按报表划分的
    导出则把子工作表命名为社会信用代码。后者不能把信用代码当作表单，
    因此优先取工作表名中的 20201 等表单号；没有时再从指标编码前五位
    推断，最后才保留有意义的工作表名称或文件名中的兜底表单号。
    """
    name = str(sheet_name or "").strip()
    matched = re.search(r"(?<!\d)(20\d{3})(?!\d)", name)
    if matched:
        return matched.group(1)
    # 中文报表名本身就是最清楚的子表身份；纯数字/字母的长串才可能是信用代码。
    if name and name.casefold() not in {"sheet", "sheet1", "sheet2", "sheet3", "封面"} and not re.fullmatch(r"[0-9A-Za-z]{8,18}", name):
        return name
    for code, _indicator_name, _value in rows:
        matched = re.match(r"^(20\d{3})\d{3,}$", str(code))
        if matched:
            return matched.group(1)
    if name and name.casefold() not in {"sheet", "sheet1", "sheet2", "sheet3", "封面"}:
        return name
    return str(fallback or "").strip()


def _open_read_only(path: Path):
    """按后缀分派 .xls（xlrd）/ .xlsx（openpyxl），返回 (sheet名→行迭代器) 工厂。"""
    suffix = path.suffix.casefold()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))

        def make_rows(sheet_name: str) -> list[list[Any]]:
            sh = book.sheet_by_name(sheet_name)
            return [ [sh.cell_value(r, c) if c < sh.ncols else None for c in range(3)]
                     for r in range(sh.nrows) ]

        return book.sheet_names(), make_rows, lambda: None
    book = load_workbook(path, read_only=True, data_only=True)

    def make_rows_xlsx(sheet_name: str) -> list[list[Any]]:
        return [list(row) + [None, None, None] for row in book[sheet_name].iter_rows(values_only=True)]

    return list(book.sheetnames), make_rows_xlsx, book.close


def parse_source_file(path: Path) -> list[IndicatorValue]:
    """把一个导出文件归一化为指标记录列表（自动识别三种导出方式）。"""
    stem = path.stem
    segments = stem.split("#")
    records: list[IndicatorValue] = []
    sheet_names, make_rows, close = _open_read_only(path)
    try:
        if len(segments) >= 4 and segments[0] in {"banks", "BANKS"}:
            # 按报表划分：banks#日期#01#表单.xls，sheet 名=机构代码
            date, form = segments[1], segments[3]
            for sheet_name in sheet_names:
                rows = _read_sheet_rows(make_rows(sheet_name))
                lookup = {code: value for code, _name, value in rows}
                org_code = _norm_code(sheet_name)
                org_name = str(lookup.get(INDICATOR_NAME_CODE, "") or "")
                for code, name, value in rows:
                    records.append(IndicatorValue(org_name, org_code, code, name, value, date, form))
        else:
            # 逐机构逐表 / 按机构划分：文件名给出机构（或表单），sheet 名含 2020 为表单
            if len(segments) >= 5:
                org_code, date, form_from_name, org_name = segments[0], segments[1], segments[3], segments[4]
            else:
                raise PeriodCompareError(f"无法识别的导出文件名：{path.name}")
            for sheet_name in sheet_names:
                rows = _read_sheet_rows(make_rows(sheet_name))
                form = _form_from_sheet(sheet_name, rows, form_from_name)
                name_lookup = {code: value for code, _n, value in rows}
                resolved_name = str(name_lookup.get(INDICATOR_NAME_CODE, "") or org_name)
                resolved_code = str(name_lookup.get(INDICATOR_CODE_CODE, "") or org_code)
                for code, name, value in rows:
                    records.append(IndicatorValue(resolved_name, resolved_code, code, name, value, date, form))
    finally:
        close()
    return records


def list_period_pairs(
    current_dir: Path | None,
    previous_dir: Path | None,
) -> list[dict[str, object]]:
    """按“工作簿 → 工作表 → 机构 + 表单”配对两期数据，供界面展示。

    三种导出方式的配对键：
    - 逐机构逐表（机构代码#日期#01#表单#机构名）：机构名 + 表单号；
    - 按机构划分（reports#代码#日期#01#机构名）：逐子工作表配对；
    - 按报表划分（banks#日期#01#表单）：逐社会信用代码子工作表配对。
    每行均保留工作簿名、工作表名与数据日期。只在单期出现的子表同样列出，
    另一侧标记缺失，便于发现漏报。
    """

    def collect(directory: Path | None) -> dict[tuple[str, str], dict[str, str]]:
        result: dict[tuple[str, str], dict[str, str]] = {}
        if directory is None or not directory.is_dir():
            return result
        for path in sorted(directory.iterdir()):
            if not path.is_file() or path.suffix.casefold() not in {".xls", ".xlsx"} or path.name.startswith("~$"):
                continue
            for item in _list_source_sheets(path):
                key = (item["orgKey"], item["form"])
                result[key] = item
        return result

    cur = collect(current_dir)
    pre = collect(previous_dir)
    rows: list[dict[str, object]] = []
    for key in sorted(set(cur) | set(pre), key=lambda k: (k[0], k[1])):
        cur_item = cur.get(key)
        pre_item = pre.get(key)
        rows.append({
            "org": key[0],
            "form": key[1],
            "orgName": (cur_item or pre_item or {}).get("orgName", key[0]),
            "curName": cur_item["name"] if cur_item else "",
            "curPath": cur_item["path"] if cur_item else "",
            "curSheet": cur_item["sheet"] if cur_item else "",
            "curDate": cur_item["date"] if cur_item else "",
            "preName": pre_item["name"] if pre_item else "",
            "prePath": pre_item["path"] if pre_item else "",
            "preSheet": pre_item["sheet"] if pre_item else "",
            "preDate": pre_item["date"] if pre_item else "",
            "matched": cur_item is not None and pre_item is not None,
            "side": "both" if cur_item and pre_item else "cur" if cur_item else "pre",
        })
    return rows


def _list_source_sheets(path: Path) -> list[dict[str, str]]:
    """读取一份导出工作簿的子表身份，供跨期配对清单使用。"""
    segments = path.stem.split("#")
    sheet_names, make_rows, close = _open_read_only(path)
    items: list[dict[str, str]] = []
    try:
        is_banks = len(segments) >= 4 and segments[0].casefold() == "banks"
        if is_banks:
            date, fallback_form = segments[1], segments[3]
            fallback_org_code, fallback_org_name = "", ""
        elif len(segments) >= 5:
            if segments[0].casefold() == "reports":
                fallback_org_code, date, fallback_form, fallback_org_name = (
                    segments[1], segments[2], "", segments[4]
                )
            else:
                fallback_org_code, date, fallback_form, fallback_org_name = (
                    segments[0], segments[1], segments[3], segments[4]
                )
        else:
            # 非标准命名也列出其工作表，避免在界面中静默漏掉文件。
            date, fallback_form, fallback_org_code, fallback_org_name = "", "", "", path.stem

        for sheet_name in sheet_names:
            rows = _read_sheet_rows(make_rows(sheet_name))
            if not rows:
                continue
            lookup = {code: value for code, _name, value in rows}
            org_code = _norm_code(lookup.get(INDICATOR_CODE_CODE)) or _norm_code(fallback_org_code)
            org_name = str(lookup.get(INDICATOR_NAME_CODE, "") or fallback_org_name or sheet_name).strip()
            if is_banks and not org_code:
                org_code = _norm_code(sheet_name)
            form = _form_from_sheet(sheet_name, rows, fallback_form)
            items.append({
                "orgKey": org_code or org_name,
                "orgName": org_name,
                "form": form,
                "name": path.name,
                "path": str(path),
                "sheet": str(sheet_name),
                "date": str(date),
            })
    finally:
        close()
    return items


def load_period_directory(
    directory: Path,
    *,
    label: str,
    on_step: Callable[[str], None] | None = None,
) -> dict[tuple[str, str], IndicatorValue]:
    if not directory.is_dir():
        raise PeriodCompareError(f"{label}目录不存在：{directory}")
    files = sorted(
        p for p in directory.iterdir()
        if p.is_file() and p.suffix.casefold() in {".xls", ".xlsx"} and not p.name.startswith("~$")
    )
    if not files:
        raise PeriodCompareError(f"{label}目录中没有 .xls/.xlsx 报送文件：{directory}")
    records: dict[tuple[str, str], IndicatorValue] = {}
    for path in files:
        if on_step:
            on_step(f"读取{label}文件：{path.name}")
        for record in parse_source_file(path):
            records[(record.org_name, record.code)] = record
    return records


# ---------------------------------------------------------------------------
# 两期比较
# ---------------------------------------------------------------------------

def _numeric(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _apply_unit(value: Any, indicator: IndicatorDef | None) -> Any:
    number = _numeric(value)
    if number is None:
        return str(value).strip() if value is not None else ""
    if indicator is not None and indicator.no_unit_convert:
        return number
    return number * UNIT_FACTOR


def _vba_has_value(record: IndicatorValue, indicator: IndicatorDef | None) -> bool:
    """VBA 读数口径：文字型指标有非空值才读；数值型指标非 0 才读。

    对应 VBA ``If (data_type="文字" And value<>"") Or (data_type<>"文字" And
    ycV(value)<>0)``；ycV 把非数值文本当 0。零值/空行不进入比较与输出。
    """
    if record is None:
        return False
    is_text = indicator is not None and indicator.data_type == "文字"
    if is_text:
        return str(record.value).strip() != ""
    return _numeric(record.value) not in (None, 0.0)


def compare_periods(
    current: dict[tuple[str, str], IndicatorValue],
    previous: dict[tuple[str, str], IndicatorValue],
    config: PeriodConfig,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    keys = list(current.keys()) + [k for k in previous.keys() if k not in current]
    for org_name, code in keys:
        cur = current.get((org_name, code))
        pre = previous.get((org_name, code))
        if cur is None and pre is None:
            continue
        indicator = config.indicators.get(code)
        # VBA 口径：两期读数时零值/空行不进入结果（文字型有值即读）。
        if not _vba_has_value(cur, indicator) and not _vba_has_value(pre, indicator):
            continue
        sample = cur or pre
        org = config.orgs.get(org_name)
        is_text = indicator is not None and indicator.data_type == "文字"
        cur_value = _apply_unit(cur.value, indicator) if cur else None
        pre_value = _apply_unit(pre.value, indicator) if pre else None
        if is_text:
            # VBA 按数据属性判断文字型：即使值恰好是数字也只比较字符串。
            cur_value = str(cur.value).strip() if cur is not None else None
            pre_value = str(pre.value).strip() if pre is not None else None
        if isinstance(cur_value, str) or isinstance(pre_value, str):
            difference = "文字变动" if (cur_value or "") != (pre_value or "") else ""
            change_pct: float | None = None
            remark = "文字变动" if difference else ""
        else:
            cur_num = cur_value if cur_value is not None else 0.0
            pre_num = pre_value if pre_value is not None else 0.0
            if cur is not None and pre is None and cur_num != 0:
                difference, change_pct, remark = "", None, "本期有，上期无"
            elif pre is not None and cur is None and pre_num != 0:
                difference, change_pct, remark = "", None, "本期无，上期有"
            else:
                difference = cur_num - pre_num
                if pre_num != 0:
                    # 与 VBA 一致：环比 = 增减额 / 上期 * 100（上期取原值）。
                    change_pct = difference / pre_num * 100.0
                    remark = config.alert_remark(change_pct)
                else:
                    change_pct = None
                    remark = ""
        rows.append({
            "地区": org.region if org else "",
            "数据属性": indicator.data_type if indicator else "",
            "机构名称": org_name,
            "指标编码": code,
            "指标名称": (cur.name if cur else pre.name if pre else ""),
            "当期数": cur_value,
            "上期数": pre_value,
            "变动绝对值": difference,
            "环比变动": change_pct,
            "备注": remark,
            "机构类别": org.org_class if org else "",
            "承接行": org.bank_row if org else "",
            "社会信用代码": (cur.org_code if cur else pre.org_code if pre else ""),
            "数据日期": (cur.date if cur else pre.date if pre else ""),
            "币种": "人民币",
            "是否说明": "",
            "级别": "",
            "计算过程": "",
            "频度": "季",
        })
    rows.sort(key=lambda row: (row["机构名称"], row["指标编码"]))
    return rows


# ---------------------------------------------------------------------------
# 大集中比较
# ---------------------------------------------------------------------------

def _load_central_matrix(central_path: Path) -> tuple[dict[tuple[str, str], float], set[str]]:
    """读集中系统数据 sheet，返回 {(行标签, 指标名): 值} 与 {指标名集合}。"""
    if not central_path.is_file():
        raise PeriodCompareError(f"大集中数据文件不存在：{central_path}")
    book = load_workbook(central_path, read_only=True, data_only=True)
    try:
        if "集中系统数据" not in book.sheetnames:
            raise PeriodCompareError("大集中文件缺少“集中系统数据”工作表")
        sheet = book["集中系统数据"]
        matrix: list[list[Any]] = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()
    if not matrix:
        raise PeriodCompareError("“集中系统数据”工作表为空")
    # 表头行：首个含非空文本的行（真实文件里表头在数据行之前）；
    # 行标签列：表头为空但其下各行是文本的那一列。
    header_index = next(
        (
            i for i, row in enumerate(matrix)
            if any(isinstance(v, str) and v.strip() for v in row)
        ),
        None,
    )
    if header_index is None:
        raise PeriodCompareError("“集中系统数据”中未找到表头行")
    header = matrix[header_index]
    label_col = next(
        (
            c for c in range(min(4, len(header)))
            if (header[c] is None or str(header[c]).strip() in {"", " "})
            and any(
                r > header_index and c < len(matrix[r]) and isinstance(matrix[r][c], str) and matrix[r][c].strip()
                for r in range(header_index + 1, min(header_index + 5, len(matrix)))
            )
        ),
        None,
    )
    if label_col is None:
        raise PeriodCompareError("“集中系统数据”中未找到机构行标签列")
    values: dict[tuple[str, str], float] = {}
    headers: set[str] = set()
    for c, title in enumerate(header):
        if c == label_col or title is None or str(title).strip() == "":
            continue
        headers.add(str(title).strip())
        for r in range(header_index + 1, len(matrix)):
            if c >= len(matrix[r]) or matrix[r][c] is None:
                continue
            label = matrix[r][label_col]
            if not isinstance(label, str) or not label.strip():
                continue
            number = _numeric(matrix[r][c])
            if number is not None:
                values[(label.strip(), str(title).strip())] = number
    return values, headers


def _load_central_org_map(central_path: Path) -> dict[str, str]:
    """读参照表：社会信用代码 → 报表项目（大集中行标签）。"""
    book = load_workbook(central_path, read_only=True, data_only=True)
    mapping: dict[str, str] = {}
    try:
        if "参照表" not in book.sheetnames:
            return mapping
        rows = [list(row) for row in book["参照表"].iter_rows(values_only=True)]
    finally:
        book.close()
    if not rows:
        return mapping
    header = [str(v or "").strip() for v in rows[0]]
    try:
        code_col = header.index("统一社会信用代码")
        item_col = header.index("报表项目")
    except ValueError:
        return mapping
    for row in rows[1:]:
        if len(row) <= max(code_col, item_col):
            continue
        code = str(row[code_col] or "").strip()
        item = str(row[item_col] or "").strip()
        if code and item:
            mapping[code] = item
    return mapping


def compare_central(
    current: dict[tuple[str, str], IndicatorValue],
    central_path: Path,
    config: PeriodConfig,
    *,
    tolerance: float = CENTRAL_DIFF_TOLERANCE,
) -> list[dict[str, Any]]:
    matrix, _headers = _load_central_matrix(central_path)
    org_map = _load_central_org_map(central_path)
    rows: list[dict[str, Any]] = []
    for (org_name, code), record in sorted(current.items()):
        indicator = config.indicators.get(code)
        if indicator is None or not indicator.check_central or not indicator.central_name:
            continue
        # VBA 口径：只核对当期有值的指标（零值/空行不读）。
        if not _vba_has_value(record, indicator):
            continue
        org = config.orgs.get(org_name)
        org_code = org.code if org and org.code else record.org_code
        row_label = org_map.get(org_code) or (org.report_item if org else "")
        base_value = _apply_unit(record.value, indicator)
        base_num = _numeric(base_value)
        # 集中系统数据以亿元报送，统一换算成万元再比较（输出列同样存万元）。
        central_value = matrix.get((row_label, indicator.central_name)) if row_label else None
        if central_value is not None:
            central_value = central_value * CENTRAL_UNIT_FACTOR
        if base_num is None:
            continue
        if central_value is None:
            difference, ratio, note = None, None, "谨慎核实"
        else:
            difference = abs(base_num - central_value)
            ratio = (difference / abs(base_num) * 100.0) if base_num else None
            note = "差异超过100元" if difference > tolerance else ""
        rows.append({
            "地区": org.region if org else "",
            "数据属性": indicator.data_type,
            "机构名称": org_name,
            "指标编码": code,
            "指标名称": record.name,
            "基础数据值": base_value,
            "大集中值": central_value,
            "差异绝对值": difference,
            "差异幅度": ratio,
            "是否说明": note,
        })
    return rows


# ---------------------------------------------------------------------------
# 特殊指标与复杂校验规则
# ---------------------------------------------------------------------------

def _period_month(date_text: str) -> tuple[int, int] | None:
    """从文件名日期（2026-06-30）解析 (年, 月)；解析失败返回 None。"""
    match = re.search(r"(\d{4})[-./]?(\d{2})", str(date_text))
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _split_code_list(content: str) -> list[str]:
    """指标代码清单 → 代码列表；支持逗号/顿号/分号/空白分隔。"""
    return [item for item in re.split(r"[,，、;；\s]+", content.strip()) if item]


def apply_special_rules(
    rows: list[dict[str, Any]],
    current: dict[tuple[str, str], IndicatorValue],
    previous: dict[tuple[str, str], IndicatorValue],
    config: PeriodConfig,
    *,
    on_step: Callable[[str], None] | None = None,
) -> None:
    """执行“累计不降”类规则，命中行写“是否说明+级别”。

    当期值 < 上期值（均换算为万元后）即命中。“累计不降(当年)”仅在同一年度
    内比较（跨年累计不可比，跳过并写计算过程说明）；“累计不降(历史)”跨年
    也可比。旧“特殊指标-自定义”表的规则经 ``effective_rules`` 转换后在此统一执行。
    """
    step = on_step or (lambda _text: None)
    rules = [
        r for r in config.effective_rules()
        if not r.disabled and r.type in (RULE_TYPE_ACC_YEAR, RULE_TYPE_ACC_HISTORY)
    ]
    if not rules:
        return
    step(f"累计不降校验：执行 {len(rules)} 条规则")
    cur_years = {_period_month(record.date) for record in current.values()}
    pre_years = {_period_month(record.date) for record in previous.values()}
    cur_year = next(iter({y[0] for y in cur_years if y}), None)
    pre_year = next(iter({y[0] for y in pre_years if y}), None)
    row_map = {(row["机构名称"], row["指标编码"]): row for row in rows}
    hits = 0
    for rule in rules:
        for code in _split_code_list(rule.content):
            indicator = config.indicators.get(code)
            for org_name in {name for name, _code in row_map}:
                row = row_map.get((org_name, code))
                if row is None:
                    continue
                cur = current.get((org_name, code))
                pre = previous.get((org_name, code))
                cur_num = _numeric(_apply_unit(cur.value, indicator)) if cur else None
                pre_num = _numeric(_apply_unit(pre.value, indicator)) if pre else None
                if cur_num is None or pre_num is None:
                    continue
                same_year = cur_year is not None and cur_year == pre_year
                if rule.type == RULE_TYPE_ACC_YEAR and not same_year:
                    row["计算过程"] = (
                        f"跨年累计不比较（当期 {cur.date} / 上期 {pre.date}）"
                        if row.get("计算过程") in (None, "")
                        else row["计算过程"]
                    )
                    continue
                if cur_num < pre_num:
                    row["是否说明"] = rule.desc
                    row["级别"] = rule.level
                    hits += 1
    step(f"累计不降校验命中 {hits} 行")


_PLACEHOLDER = re.compile(r"(\[|\{)([^\]}]*)(\]|\})")


class _V(float):
    """表达式求值的包装数值：相等比较带容差，消除二进制浮点误报。

    值已按万元两位舍入（百元精度），``==``/``!=`` 采用 0.005 万元
    （50 元）容差——勾稽关系两侧十进制相同时不应判“不平衡”。
    ``<``/``>`` 等严格比较保持浮点原语义。
    """

    __slots__ = ()

    def _tolerant_eq(self, other) -> bool:
        try:
            return abs(float(self) - float(other)) <= 0.005
        except (TypeError, ValueError):
            return False

    def __eq__(self, other):
        return self._tolerant_eq(other)

    def __ne__(self, other):
        return not self._tolerant_eq(other)

    def __hash__(self):
        return hash(float(self))

    def __add__(self, other):
        return _V(float(self) + float(other))

    def __radd__(self, other):
        return _V(float(other) + float(self))

    def __sub__(self, other):
        return _V(float(self) - float(other))

    def __rsub__(self, other):
        return _V(float(other) - float(self))

    def __mul__(self, other):
        return _V(float(self) * float(other))

    def __rmul__(self, other):
        return _V(float(other) * float(self))

    def __truediv__(self, other):
        return _V(float(self) / float(other))

    def __rtruediv__(self, other):
        return _V(float(other) / float(self))


def _AND(*values) -> bool:
    return all(bool(v) for v in values)


def _OR(*values) -> bool:
    return any(bool(v) for v in values)


def _NOT(value) -> bool:
    return not bool(value)


def _placeholder_code(inner: str, resolver: Callable[[str], str]) -> tuple[str, str]:
    """占位符内文 → (指标代码, 错误说明)。

    兼容三种写法：旧 8 段 `[机构,地区,类别,代码,…]` 取下标 3；新短写
    `[代码]`；指标名称 `[一级资本净额]`（经指标参照解析，重名/未知即报错）。
    """
    if "," in inner:
        parts = [p.strip() for p in inner.split(",")]
        return (parts[3] if len(parts) > 3 else ""), ""
    code = inner.strip()
    if not code:
        return "", "存在空占位符 []"
    if code.isdigit():
        return code, ""
    try:
        return resolver(code), ""
    except KeyError:
        return "", f"引用的指标名称“{code}”在指标参照中不存在（或重名）"


def _expression_codes(
    expression: str, resolver: Callable[[str], str]
) -> tuple[set[str], list[str]]:
    """提取表达式引用的全部指标代码，附带解析错误列表。"""
    codes: set[str] = set()
    errors: list[str] = []
    for match in _PLACEHOLDER.finditer(expression):
        code, error = _placeholder_code(match.group(2), resolver)
        if code:
            codes.add(code)
        if error:
            errors.append(error)
    return codes, errors


def _render_expression(
    expression: str,
    cur_lookup: Callable[[str], float],
    pre_lookup: Callable[[str], float],
    resolver: Callable[[str], str],
) -> tuple[str, list[str]]:
    """把占位符替换为 ``_V(数值)``，返回（代入后的表达式, 错误列表）。

    缺省指标记 0，值为 0 时按 VBA 口径替换为 0.01（避免除零）；
    数值四舍五入两位（万元，VBA RoundRule）。
    """
    errors: list[str] = []

    def replace(match: re.Match) -> str:
        code, error = _placeholder_code(match.group(2), resolver)
        if error:
            errors.append(error)
            code = ""
        lookup = cur_lookup if match.group(1) == "[" else pre_lookup
        value = lookup(code) if code else 0.0
        return f"_V({0.01 if value == 0 else round(value, 2)!r})"

    return _PLACEHOLDER.sub(replace, expression), errors


def _translate_for_eval(filled: str) -> str:
    """Excel 风格运算符/函数 → Python：`<>`、单独 `=`、大写 AND/OR/NOT。"""
    text = filled.replace("<>", "!=")                          # Excel 不等号
    text = re.sub(r"(?<![<>!])=(?!=)", "==", text)             # 单独 = 为相等比较
    text = re.sub(r"\b(AND|OR|NOT)\b", lambda m: "_" + m.group(1), text)
    return text


def _translate_function_calls(text: str) -> str:
    """二次尝试：小写 and/or/not 的函数式写法（如 ``or(a,b)``）转为 _OR(a,b)。

    仅当首次求值出现语法错误时使用；裸关键字 ``a and b`` 是合法 Python，
    不会走到这里。
    """
    return re.sub(
        r"\b(and|or|not)\s*\(",
        lambda m: "_" + m.group(1).upper() + "(",
        text,
        flags=re.IGNORECASE,
    )


def _compile_translated(filled: str) -> str:
    """试编译代入后的表达式，返回错误说明（可执行返回空串）。

    与运行时求值保持同一套两步尝试：先按主转换，语法不过再试函数式
    小写转换，避免把 ``or(a,b)`` 这类 Excel 习惯写法误报为语法错误。
    """
    primary = _translate_for_eval(filled)
    message = ""
    for eval_text in (primary, _translate_function_calls(primary)):
        try:
            compile(eval_text, "<校验规则>", "eval")
            return ""
        except SyntaxError as exc:
            message = exc.msg or "表达式语法无法识别"
    return message


_EVAL_NAMES = {"_AND": _AND, "_OR": _OR, "_NOT": _NOT, "_V": _V}


def _eval_expression(
    expression: str,
    cur_lookup: Callable[[str], float],
    pre_lookup: Callable[[str], float],
    resolver: Callable[[str], str],
) -> tuple[bool, str, str]:
    """求值一条校验表达式，返回 (是否命中, 代入后的表达式, 错误说明)。

    表达式来自用户自己的配置簿，在空 builtins 沙箱中受限求值；
    语法/引用错误不再静默吞掉，而是随错误说明写入运行日志。
    """
    filled, errors = _render_expression(expression, cur_lookup, pre_lookup, resolver)
    if errors:
        return False, filled, errors[0]
    primary = _translate_for_eval(filled)
    for eval_text in (primary, _translate_function_calls(primary)):
        try:
            result = eval(eval_text, {"__builtins__": {}}, dict(_EVAL_NAMES))  # noqa: S307 - 用户配置簿内的受控表达式
            return bool(result), filled, ""
        except SyntaxError:
            continue
        except Exception as exc:
            return False, filled, f"表达式无法求值：{exc}"
    return False, filled, "表达式语法无法识别"


def apply_complex_rules(
    rows: list[dict[str, Any]],
    current: dict[tuple[str, str], IndicatorValue],
    previous: dict[tuple[str, str], IndicatorValue],
    config: PeriodConfig,
    *,
    on_step: Callable[[str], None] | None = None,
) -> None:
    """执行“表达式”类校验规则：逐机构代入指标值，命中行写“是否说明+级别+计算过程”。

    旧“复杂校验-自定义”表经 ``effective_rules`` 转换后在此统一执行；
    表达式无法求值的规则记入运行日志，不再静默跳过。
    """
    step = on_step or (lambda _text: None)
    rules = [
        r for r in config.effective_rules()
        if not r.disabled and r.type == RULE_TYPE_EXPRESSION
    ]
    if not rules:
        return
    step(f"表达式校验：执行 {len(rules)} 条")

    lookup, _duplicated = _indicator_name_lookup(config)

    def resolver(name: str) -> str:
        return lookup[name]

    def values_by_org(values: dict[tuple[str, str], IndicatorValue]) -> dict[str, dict[str, float]]:
        result: dict[str, dict[str, float]] = {}
        for (org, code), record in values.items():
            number = _numeric(_apply_unit(record.value, config.indicators.get(code)))
            if number is None:
                continue
            result.setdefault(org, {})[code] = number
        return result

    cur_by_org = values_by_org(current)
    pre_by_org = values_by_org(previous)
    row_map = {(row["机构名称"], row["指标编码"]): row for row in rows}
    hit_rules = 0
    rule_errors: dict[str, str] = {}
    for rule in rules:
        label = rule.desc
        involved_codes, code_errors = _expression_codes(rule.content, resolver)
        for error in code_errors:
            rule_errors.setdefault(label, error)
        rule_hit_any_org = False
        for org in sorted(set(cur_by_org) | set(pre_by_org)):
            cur_vals = cur_by_org.get(org, {})
            pre_vals = pre_by_org.get(org, {})
            hit, filled, error = _eval_expression(
                rule.content,
                lambda code, _v=cur_vals: _v.get(code, 0.0),
                lambda code, _v=pre_vals: _v.get(code, 0.0),
                resolver,
            )
            if error:
                rule_errors.setdefault(label, error)
            if not hit:
                continue
            rule_hit_any_org = True
            reference = f"校验规则 {rule.rule_id} 命中" if rule.rule_id else "复杂校验命中"
            for code in involved_codes:
                row = row_map.get((org, code))
                if row is not None and not row.get("是否说明"):
                    row["是否说明"] = rule.desc
                    row["级别"] = rule.level
                    row["计算过程"] = f"{reference}：{filled}"
        if rule_hit_any_org:
            hit_rules += 1
    for label, error in rule_errors.items():
        step(f"表达式校验：规则“{label}”无法执行，已跳过（{error}）")
    step(f"表达式校验命中 {hit_rules} 条规则")




def write_output(
    output_path: Path,
    period_rows: list[dict[str, Any]],
    central_rows: list[dict[str, Any]],
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "两期对比"
    sheet.append(PERIOD_SHEET_HEADERS)
    for row in period_rows:
        sheet.append([row.get(header, "") for header in PERIOD_SHEET_HEADERS])
    central_sheet = book.create_sheet("大集中对比")
    central_sheet.append(CENTRAL_SHEET_HEADERS)
    for row in central_rows:
        central_sheet.append([row.get(header, "") for header in CENTRAL_SHEET_HEADERS])
    for column in range(1, len(PERIOD_SHEET_HEADERS) + 1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = 14
    central_sheet.column_dimensions["A"].width = 14
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)


def run_period_compare(
    *,
    current_dir: Path,
    previous_dir: Path,
    central_path: Path | None,
    output_dir: Path,
    config_path: Path,
    central_tolerance_yuan: float | None = None,
    on_step: Callable[[str], None] | None = None,
) -> Path:
    """执行跨期比较并输出工作簿；返回输出文件路径。

    ``central_tolerance_yuan`` 为大集中核对差异容差（元，来自设置中心）；
    缺省用内置口径 100 元（0.01 万元）。
    """
    step = on_step or (lambda _text: None)
    config = load_period_config(config_path)
    step(f"已加载配置：指标 {len(config.indicators)} 项，机构 {len(config.orgs)} 家，警戒区间 {len(config.alerts)} 档")
    for warning in config.warnings:
        step(f"配置提醒：{warning}")
    tolerance = CENTRAL_DIFF_TOLERANCE
    if central_tolerance_yuan is not None:
        try:
            tolerance = max(float(central_tolerance_yuan), 0.0) / 10000.0
        except (TypeError, ValueError):
            step(f"配置提醒：大集中核对容差“{central_tolerance_yuan}”无效，按默认 100 元处理")
    # 目录路径或已加载的指标字典（便于调用方复用加载结果）均可。
    current = current_dir if isinstance(current_dir, dict) else load_period_directory(current_dir, label="当期", on_step=step)
    previous = previous_dir if isinstance(previous_dir, dict) else load_period_directory(previous_dir, label="上期", on_step=step)
    step(f"当期指标值 {len(current)} 条，上期指标值 {len(previous)} 条")
    period_rows = compare_periods(current, previous, config)
    step(f"两期比较完成：{len(period_rows)} 行")
    apply_special_rules(period_rows, current, previous, config, on_step=step)
    apply_complex_rules(period_rows, current, previous, config, on_step=step)
    central_rows: list[dict[str, Any]] = []
    if central_path is not None:
        central_rows = compare_central(current, central_path, config, tolerance=tolerance)
        flagged = sum(1 for row in central_rows if row["是否说明"])
        step(f"大集中核对完成：{len(central_rows)} 行，其中 {flagged} 行需要说明")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"报表采集系统_比较结果_{timestamp}.xlsx"
    write_output(output_path, period_rows, central_rows)
    return output_path
