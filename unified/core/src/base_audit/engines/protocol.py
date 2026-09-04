"""计算引擎统一协议。

业务层（service 等）只依赖本协议；平台差异全部封装在
``engines.excel_wps``（Microsoft Excel / WPS COM，仅 Windows）与
``engines.libreoffice_adapter``（LibreOffice Calc，UOS/麒麟/Linux）中。
协议定义自 flet/uos ``calculation.py`` 的 ``Calculator`` 抽取精简。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Protocol


@dataclass(frozen=True)
class CalculationResult:
    workbook_path: Path
    engine_name: str
    elapsed_seconds: float
    # 计算后缓存中的公式错误值（如 #REF!、#NAME?），(工作表, 单元格, 错误文本)。
    error_cells: tuple[tuple[str, str, str], ...] = ()
    # 可选说明：引擎是否可用、版本、降级等信息。
    diagnostic: str = ""


class Calculator(Protocol):
    """一个可重算 .xlsx 工作簿并保留公式缓存的计算引擎。"""

    engine_name: str

    def recalculate(self, workbook_path: Path, *, formula_ranges: Iterable[object] = ()) -> CalculationResult: ...


_FORMULA_ERROR_PREFIXES = ("#N/A", "#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#NUM!", "#NULL!", "#SPILL!")


def scan_formula_errors(workbook_path: Path, *, formula_ranges: Iterable[object] = ()) -> tuple[tuple[str, str, str], ...]:
    """用 openpyxl 只读扫描计算后工作簿中的公式错误缓存值。

    只报告错误值，不把普通业务提示文字当成公式错误；未限定区域时扫描全表。
    """
    from openpyxl import load_workbook
    from openpyxl.utils import range_boundaries

    scoped = list(formula_ranges)
    values_book = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    errors: list[tuple[str, str, str]] = []
    try:
        targets: list[tuple[str, tuple[int, int, int, int] | None]] = []
        if scoped:
            for item in scoped:
                sheet_name = getattr(item, "sheet_name", None)
                address = getattr(item, "address", None)
                if sheet_name not in values_book.sheetnames or address is None:
                    continue
                try:
                    targets.append((sheet_name, range_boundaries(address)))
                except Exception:
                    continue
        else:
            targets = [(sheet.title, None) for sheet in values_book.worksheets]
        for sheet_name, bounds in targets:
            sheet = values_book[sheet_name]
            rows = (
                sheet.iter_rows(
                    min_row=bounds[1], max_row=bounds[3],
                    min_col=bounds[0], max_col=bounds[2],
                )
                if bounds
                else sheet.iter_rows()
            )
            for row in rows:
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("#"):
                        matched = next((prefix for prefix in _FORMULA_ERROR_PREFIXES if value.startswith(prefix)), None)
                        if matched:
                            errors.append((sheet_name, cell.coordinate, value))
    finally:
        values_book.close()
    return tuple(errors)
