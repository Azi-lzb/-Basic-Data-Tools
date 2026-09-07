from __future__ import annotations

import os
import threading
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any

from .discovery import (
    classify_source_files,
    detect_period,
    explanation_files,
    recommend_template,
)
from .history import HISTORY_AUDIT_SHEET, HISTORY_HEADERS
from .name_config import (
    HISTORY_WORKBOOK_NAME,
    initialize_config,
    load_config_editor_data,
    reset_default_configuration,
    save_config_editor_draft,
    validate_config_editor_draft,
)
from .engines import available_engines, engine_hint, valid_engine_values
from .service import AuditService
from .settings import SettingsStore, hide_application_data_directory


class WebApi:
    def __init__(self, project_root: Path) -> None:
        self.project_root = project_root
        hide_application_data_directory(project_root / "data")
        self.settings_store = SettingsStore(project_root / "data" / "用户设置.json")
        self.settings = self.settings_store.load()
        # 历史审核配置：优先用上次选择的路径，否则用程序根目录默认文件。
        self.history_path = (
            Path(self.settings.history_config)
            if self.settings.history_config
            else project_root / HISTORY_WORKBOOK_NAME
        )
        bundled_templates = project_root / "templates"
        formal_templates = project_root / "2026-07-31" / "模板文件"
        saved_input = self.settings.last_input_dir
        saved_output = self.settings.last_output_dir
        self.state: dict[str, Any] = {
            "busy": False,
            "status": "就绪",
            "log": [],
            "input": saved_input,
            "templateDir": self.settings.last_template_dir or str(
                bundled_templates if bundled_templates.is_dir() else formal_templates
            ),
            "template": "",
            "templateManual": False,
            "external": self.settings.last_external_file,
            "historyConfig": str(self.history_path),
            "output": saved_output,
            "outputAuto": not self.settings.output_pinned,
            "outputPinned": self.settings.output_pinned,
            "recursiveDepth": self.settings.recursive_depth,
            "extraFiles": list(self.settings.extra_files),
            "writeFlowLogs": self.settings.write_flow_logs,
            "confirmBeforeRun": self.settings.confirm_before_run,
            "calculationEngine": (
                self.settings.calculation_engine
                if self.settings.calculation_engine in valid_engine_values()
                else "自动"
            ),
            "engines": available_engines(),
            "engineHint": engine_hint(),
            "sourceFiles": [],
            "selectedFiles": [],
            "mixedTemplates": [],
            "explanationFiles": [],
            # 报表采集系统路径：独立保存，不和逐笔统计系统的输入/输出混用。
            "pcCurDir": self.settings.period_current_dir,
            "pcPreDir": self.settings.period_previous_dir,
            "pcCentral": self.settings.period_central_file,
            "pcConfig": self.settings.period_config_file or str(project_root / "跨期比较配置.xlsx"),
            "pcOutput": self.settings.period_output_dir,
            "pcOutputAuto": self.settings.period_output_auto,
            "periodPairs": [],
        }
        # This also performs the one-time non-destructive filename migration.
        initialize_config(self.history_path)
        if self.state["pcCurDir"] or self.state["pcPreDir"]:
            self._refresh_period_pairs()

    def get_state(self) -> dict[str, Any]:
        return self.state

    def initialize_config(self) -> dict[str, Any]:
        """Migrate workflow settings to JSON and retain Excel history only."""
        path = initialize_config(self.history_path)
        self._log(f"已初始化流程配置：{path}（仅保留历史核查表审核）")
        self.state["status"] = "配置已初始化"
        return self.state

    def reset_config(self) -> dict[str, Any]:
        """Reset workflow display settings without touching Excel history."""
        path = reset_default_configuration(self.history_path)
        self._log(f"已重置自定义流程和主界面显示设置：{path}（历史核查表审核未改动）")
        self.state["status"] = "配置已重置"
        return self.state

    def get_config_editor_data(self) -> dict[str, object]:
        """Configuration-centre payload; deliberately excludes history records."""
        return load_config_editor_data(self.history_path)

    def validate_config_editor_draft(self, draft: dict[str, Any]) -> dict[str, object]:
        """Validate a pending visual-editor change without writing history."""
        errors = validate_config_editor_draft(
            self.history_path, draft
        )
        return {"valid": not errors, "errors": errors}

    def save_config_editor_draft(self, draft: dict[str, Any]) -> dict[str, object]:
        """Persist functions, workflows and display settings; Excel history stays intact."""
        if self.state["busy"]:
            return {"ok": False, "errors": ["任务正在运行，暂不能修改配置"]}
        try:
            path = save_config_editor_draft(
                self.history_path, draft
            )
        except Exception as exc:
            self._log(f"保存设置中心配置失败：{exc}")
            return {"ok": False, "errors": [str(exc)]}
        self._log(f"已保存功能、流程和主界面显示设置：{path}")
        return {"ok": True, "errors": [], "data": load_config_editor_data(path)}

    def open_config(self) -> dict[str, Any]:
        """Backward-compatible alias for :meth:`open_history_explanation`."""
        return self.open_history_explanation()

    def open_history_explanation(self) -> dict[str, Any]:
        """Open the Excel workbook that stores historical audit explanations."""
        import os
        path = self.history_path
        if not path.is_file():
            self.state["status"] = "历史审核配置不存在"
            self._log(f"未找到历史审核配置：{path}")
            return self.state
        try:
            os.startfile(path)
        except Exception as exc:
            self.state["status"] = "无法打开历史审核配置"
            self._log(f"打开历史审核配置失败：{exc}")
            return self.state
        self._log(f"已打开历史审核配置：{path}")
        return self.state

    def get_history_page(self, query: dict[str, Any] | None = None) -> dict[str, object]:
        """Return one read-only, filtered page of historical audit explanations.

        The worksheet is streamed with openpyxl: large history files never
        become a giant WebView payload.  This endpoint deliberately provides
        no editing operation; formal maintenance remains in Excel.
        """
        from openpyxl import load_workbook

        query = query if isinstance(query, dict) else {}

        def number(name: str, default: int, minimum: int, maximum: int) -> int:
            try:
                value = int(query.get(name, default))
            except (TypeError, ValueError):
                value = default
            return max(minimum, min(maximum, value))

        page = number("page", 1, 1, 1000000)
        page_size = number("pageSize", 50, 10, 100)
        keyword = str(query.get("keyword") or "").strip().casefold()
        error_type = str(query.get("errorType") or "").strip()
        opinion = str(query.get("opinion") or "").strip()
        path = self.history_path
        if not path.is_file():
            return {
                "columns": list(HISTORY_HEADERS), "items": [], "total": 0,
                "page": page, "pageSize": page_size, "totalPages": 0,
                "errorTypes": [], "opinions": [], "error": "历史审核配置不存在",
            }
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if HISTORY_AUDIT_SHEET not in workbook.sheetnames:
                raise ValueError(f"历史审核配置中缺少“{HISTORY_AUDIT_SHEET}”工作表")
            rows = workbook[HISTORY_AUDIT_SHEET].iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            missing = [name for name in HISTORY_HEADERS if name not in headers]
            if missing:
                raise ValueError("历史审核配置缺少列：" + "、".join(missing))
            positions = {name: headers.index(name) for name in HISTORY_HEADERS}
            total = 0
            error_types: set[str] = set()
            opinions: set[str] = set()
            items: list[dict[str, str]] = []
            start = (page - 1) * page_size
            end = start + page_size
            for raw in rows:
                item = {
                    name: str(raw[index] if index < len(raw) and raw[index] is not None else "").strip()
                    for name, index in positions.items()
                }
                if item["错误类型"]:
                    error_types.add(item["错误类型"])
                if item["审核意见"]:
                    opinions.add(item["审核意见"])
                haystack = " ".join(item.values()).casefold()
                if keyword and keyword not in haystack:
                    continue
                if error_type and item["错误类型"] != error_type:
                    continue
                if opinion and item["审核意见"] != opinion:
                    continue
                if start <= total < end:
                    items.append(item)
                total += 1
            total_pages = (total + page_size - 1) // page_size
            # A filter can make the requested page out of range.  Reply with
            # the real last page marker; the front end can request it once.
            return {
                "columns": list(HISTORY_HEADERS), "items": items, "total": total,
                "page": page, "pageSize": page_size, "totalPages": total_pages,
                "errorTypes": sorted(error_types), "opinions": sorted(opinions),
            }
        finally:
            workbook.close()

    def open_user_guide(self) -> dict[str, Any]:
        """Open the adjacent Word user guide with the default application."""
        import os
        path = self.project_root / "基础数据审核工具使用说明.docx"
        if not path.is_file():
            self.state["status"] = "使用说明不存在"
            self._log(f"未找到使用说明：{path}")
            return self.state
        try:
            os.startfile(path)
        except Exception as exc:
            self.state["status"] = "无法打开使用说明"
            self._log(f"打开使用说明失败：{exc}")
            return self.state
        self._log(f"已打开使用说明：{path}")
        return self.state

    def open_path(self, path: str) -> dict[str, Any]:
        """Open a local directory or file with the system default application."""
        import os
        target = Path(path)
        if not target.exists():
            self.state["status"] = "路径不存在"
            self._log(f"未找到路径：{path}")
            return self.state
        try:
            os.startfile(target)
        except Exception as exc:
            self.state["status"] = "无法打开路径"
            self._log(f"打开路径失败：{exc}")
            return self.state
        self._log(f"已打开：{path}")
        return self.state

    def _load_custom_features(self) -> list[dict[str, Any]]:
        """Read custom flows configured to appear as main-workbench entries."""
        try:
            editor = load_config_editor_data(self.history_path)
        except Exception as exc:
            self._log(f"读取流程显示配置失败：{exc}")
            return []
        display = editor.get("flowDisplay", {})
        if not isinstance(display, dict):
            return []
        features: list[dict[str, Any]] = []
        seen: set[str] = set()
        # 常用流程已有固定主按钮；这里只显示用户在流程编排中新建并勾选的流程。
        for row in editor.get("customFlows", []):
            feature_name = str(row.get("流程名") or "").strip()
            if not feature_name or feature_name in seen:
                continue
            seen.add(feature_name)
            features.append({
                "id": feature_name,
                "name": feature_name,
                "flow": feature_name,
                "shown": bool(display.get(feature_name)),
                "remark": "",
            })
        return features

    def get_custom_features(self) -> list[dict[str, Any]]:
        """Visible flat custom features for the workbench."""
        return [
            {"id": item["id"], "name": item["name"], "flow": item["flow"], "remark": item["remark"]}
            for item in self._load_custom_features()
            if item["shown"]
        ]

    def run_custom(self, feature_name: str) -> dict[str, Any]:
        """Run one user-defined flow displayed on the workbench."""
        if self.state["busy"]:
            return self.state
        feature = next(
            (item for item in self._load_custom_features() if item["id"] == feature_name),
            None,
        )
        if feature is None:
            self._log(f"未找到显示流程：{feature_name}")
            return self.state
        if not feature["flow"]:
            self._log(f"显示流程“{feature['name']}”缺少流程名称。")
            self.state["status"] = "缺少流程名称"
            return self.state
        # 主界面自定义流程走宽松校验（strict=False）：组合前置缺失时只警告不报错。
        self.start_flow(feature["flow"], strict=False)
        return self.state

    def start_flow(self, flow_name: str, strict: bool = True) -> bool:
        """Start one configured execution flow by its displayed name."""
        name = str(flow_name).strip()
        if not name:
            self._log("执行流程名称不能为空")
            return False
        return self.start("flow:" + name, strict=strict)

    def start_with_state(
        self,
        action: str,
        values: dict[str, Any] | None = None,
        selected_files: list[str] | None = None,
        strict: bool = True,
    ) -> bool:
        """Synchronize form values and launch an action in one WebView call.

        This is intentionally a single API boundary for the Python 3.7 /
        pywebview 5 compatibility package.  Older WebView bridges may stall
        when the front end sends ``update`` → ``set_selected_files`` →
        ``start`` as three immediate calls.
        """
        if values:
            self.update(values)
        if selected_files is not None:
            self.set_selected_files(selected_files)
        return self.start(action, strict=strict)

    def start_flow_with_state(
        self,
        flow_name: str,
        values: dict[str, Any] | None = None,
        selected_files: list[str] | None = None,
        strict: bool = True,
    ) -> bool:
        """Single-call variant of :meth:`start_flow` for legacy WebView."""
        name = str(flow_name).strip()
        if not name:
            self._log("执行流程名称不能为空")
            return False
        return self.start_with_state("flow:" + name, values, selected_files, strict=strict)

    def choose_folder(self, field: str) -> str:
        import webview
        current = self.state.get(field) or self.project_root
        # pywebview 5 (used by the Python 3.7 Win7 package) exposes integer
        # dialog constants; pywebview 6 uses the FileDialog enum.
        file_dialog = getattr(webview, "FileDialog", None)
        folder_dialog = (
            getattr(file_dialog, "FOLDER", None)
            if file_dialog is not None
            else getattr(webview, "FOLDER_DIALOG")
        )
        selected = webview.windows[0].create_file_dialog(
            folder_dialog, directory=str(current)
        )
        value = str(selected[0]) if selected else ""
        if value:
            previous = self.state.get(field, "")
            self.state[field] = value
            if field == "templateDir" and value != previous:
                self.state["templateManual"] = False
            if field in {"input", "templateDir"}:
                if field == "input" and not self.state["outputPinned"]:
                    self.state["output"] = str(Path(value) / "执行结果")
                    self.state["outputAuto"] = True
                # 仅通过选择源数据目录触发一次模板推荐；选择模板目录只刷新文件清单。
                # 用户手动选定模板后，仍以手动选择为准。
                self._recognize(allow_template_auto=(field == "input"))
            elif field == "output":
                self.state["outputAuto"] = False
                self.state["outputPinned"] = True
        return value

    def choose_file(self, field: str) -> str:
        import webview
        current = self.state.get(field) or self.project_root
        file_dialog = getattr(webview, "FileDialog", None)
        open_dialog = (
            getattr(file_dialog, "OPEN", None)
            if file_dialog is not None
            else getattr(webview, "OPEN_DIALOG")
        )
        selected = webview.windows[0].create_file_dialog(
            open_dialog,
            directory=str(Path(current).parent if Path(current).is_file() else current),
            file_types=("Excel 文件 (*.xlsx;*.xls)",),
        )
        value = str(selected[0]) if selected else ""
        if value:
            self.state[field] = value
            if field == "template":
                self.state["templateManual"] = True
        return value

    def choose_history_config(self) -> str:
        """选择历史审核配置.xlsx 文件并记住路径（历史数据由人工维护）。"""
        import webview
        if self.state["busy"]:
            return str(self.history_path)
        file_dialog = getattr(webview, "FileDialog", None)
        open_dialog = (
            getattr(file_dialog, "OPEN", None)
            if file_dialog is not None
            else getattr(webview, "OPEN_DIALOG")
        )
        current = self.history_path if self.history_path.is_file() else self.project_root
        selected = webview.windows[0].create_file_dialog(
            open_dialog,
            directory=str(Path(current).parent if Path(current).is_file() else current),
            file_types=("Excel 文件 (*.xlsx;*.xls)",),
        )
        value = str(selected[0]) if selected else ""
        if value:
            self.history_path = Path(value)
            self.state["historyConfig"] = value
            self.settings.history_config = value
            self.settings_store.save(self.settings)
            self._log(f"已选择历史审核配置：{value}")
        return value

    def create_template_merge(self) -> bool:
        """Interactive, manual-only creation of one combined audit template.

        It deliberately bypasses the workbench form and template recommendation:
        the user selects both the base and every input workbook in dialogs.
        """
        if self.state["busy"]:
            return False
        import webview
        file_dialog = getattr(webview, "FileDialog", None)
        open_dialog = (
            getattr(file_dialog, "OPEN", None)
            if file_dialog is not None
            else getattr(webview, "OPEN_DIALOG")
        )
        base_dir = self.state.get("templateDir") or self.project_root
        self.state["status"] = "第 1 步：请选择基准模板"
        self._log(
            "第 1 步/2：请选择基准模板（它会另存为联合模板；"
            "应包含集中系统数据、参照表等公共依赖工作表）"
        )
        selected_base = webview.windows[0].create_file_dialog(
            open_dialog,
            directory=str(base_dir),
            file_types=("Excel 文件 (*.xlsx;*.xlsm;*.xls)",),
        )
        if not selected_base:
            self._log("已取消制作联合模板：未选择底稿模板")
            return False
        base_template = Path(selected_base[0]).resolve()
        self.state["status"] = "第 2 步：请选择要并入的模板"
        self._log(
            f"已选择基准模板：{base_template.name}。第 2 步/2："
            "请批量选择要并入的其他模板（不要重复选择基准模板）"
        )
        selected_sources = webview.windows[0].create_file_dialog(
            open_dialog,
            directory=str(base_template.parent),
            allow_multiple=True,
            file_types=("Excel 文件 (*.xlsx;*.xlsm;*.xls)",),
        )
        if not selected_sources:
            self._log("已取消制作联合模板：未选择待复制工作簿")
            return False
        source_templates = [Path(path).resolve() for path in selected_sources]
        if not [path for path in source_templates if path != base_template]:
            self._log("已取消制作联合模板：待复制工作簿不能只有底稿模板本身")
            return False
        self.state["busy"] = True
        self.state["status"] = "正在制作联合模板，请勿关闭窗口……"
        self._log(f"开始制作联合模板：底稿“{base_template.name}”")
        self._log("提示：请将含外部依赖工作表的模板选作底稿；原始文件不会修改")
        threading.Thread(
            target=self._template_merge_worker,
            args=(base_template, source_templates),
            daemon=True,
        ).start()
        return True

    def _template_merge_worker(
        self, base_template: Path, source_templates: list[Path]
    ) -> None:
        started = time.monotonic()
        try:
            service = AuditService(
                config_path=self.history_path,
                engine_preference=str(self.state["calculationEngine"]),
            )
            result = service.merge_template_files(
                base_template=base_template,
                source_templates=source_templates,
                on_step=self._log_detail,
            )
            self.state["status"] = result.summary_text().splitlines()[0]
            self._log(result.summary_text())
        except Exception as exc:
            self.state["status"] = "制作联合模板失败"
            self._log("制作联合模板失败：" + str(exc))
        finally:
            elapsed = time.monotonic() - started
            self._log(f"任务结束，本次耗时：{elapsed:.1f} 秒")
            self.state["busy"] = False

    def append_extra_files(self) -> list[str]:
        """待处理清单：弹窗批量追加文件（源数据目录之外也可）。"""
        if self.state["busy"]:
            return list(self.state.get("extraFiles", []))
        import webview
        file_dialog = getattr(webview, "FileDialog", None)
        open_dialog = (
            getattr(file_dialog, "OPEN", None)
            if file_dialog is not None
            else getattr(webview, "OPEN_DIALOG")
        )
        current = self.state.get("input") or str(self.project_root)
        selected = webview.windows[0].create_file_dialog(
            open_dialog,
            directory=current,
            allow_multiple=True,
            file_types=("Excel 文件 (*.xlsx;*.xlsm;*.xls)",),
        )
        if not selected:
            return list(self.state.get("extraFiles", []))
        added = []
        existing = set(self.state.get("extraFiles", []))
        for item in selected:
            path = str(Path(item).resolve())
            if path not in existing and Path(path).is_file():
                existing.add(path)
                added.append(path)
        if added:
            self.state["extraFiles"] = sorted(existing)
            self.settings.extra_files = list(self.state["extraFiles"])
            self.settings_store.save(self.settings)
            self._log(f"已追加 {len(added)} 个待处理文件")
            self._recognize(allow_template_auto=False)
        return list(self.state.get("extraFiles", []))

    def remove_extra_file(self, path: str) -> list[str]:
        """从追加清单移除一个手动追加的文件（不影响源数据目录扫描结果）。"""
        files = [item for item in self.state.get("extraFiles", []) if item != path]
        if len(files) != len(self.state.get("extraFiles", [])):
            self.state["extraFiles"] = files
            self.settings.extra_files = files
            self.settings_store.save(self.settings)
            self._log(f"已移除追加文件：{Path(path).name}")
            self._recognize(allow_template_auto=False)
            selected = [item for item in (self.state.get("selectedFiles") or []) if item != path]
            self.state["selectedFiles"] = selected
        return list(self.state.get("extraFiles", []))

    def _refresh_period_pairs(self) -> None:
        """按机构+表单配对两期已选目录中的文件，供报表采集页展示。"""
        from .period_compare import list_period_pairs

        cur = self.state.get("pcCurDir")
        pre = self.state.get("pcPreDir")
        if not cur and not pre:
            self.state["periodPairs"] = []
            return
        try:
            self.state["periodPairs"] = list_period_pairs(
                Path(cur) if cur else None, Path(pre) if pre else None
            )
        except Exception as exc:
            self.state["periodPairs"] = []
            self._log(f"两期文件配对失败：{exc}")

    def choose_period_compare(self, kind: str) -> str:
        """报表采集系统页：选择跨期比较的输入或输出路径。

        只弹窗并把结果记入 state，供页面回显；不触发执行。
        """
        if self.state["busy"] or kind not in {"pcCurDir", "pcPreDir", "pcCentral", "pcConfig", "pcOutput"}:
            return ""
        import webview
        file_dialog = getattr(webview, "FileDialog", None)
        folder_dialog = (
            getattr(file_dialog, "FOLDER", None)
            if file_dialog is not None
            else getattr(webview, "FOLDER_DIALOG")
        )
        open_dialog = (
            getattr(file_dialog, "OPEN", None)
            if file_dialog is not None
            else getattr(webview, "OPEN_DIALOG")
        )
        titles = {
            "pcCurDir": "选择当期（本期）数据目录",
            "pcPreDir": "选择上期数据目录",
            "pcCentral": "选择大集中数据文件",
            "pcConfig": "选择跨期比较配置.xlsx",
            "pcOutput": "选择输出目录",
        }
        current = self.state.get(kind) or self.state.get("input") or str(self.project_root)
        if kind == "pcCentral" or kind == "pcConfig":
            selected = webview.windows[0].create_file_dialog(
                open_dialog,
                directory=str(Path(current).parent) if Path(current).is_file() else current,
                file_types=("Excel 文件 (*.xlsx;*.xlsm;*.xls)",),
            )
        else:
            selected = webview.windows[0].create_file_dialog(
                folder_dialog, directory=str(current)
            )
        value = str(selected[0]) if selected else ""
        if value:
            self.state[kind] = str(Path(value).resolve())
            if kind == "pcCurDir" and self.state.get("pcOutputAuto", True):
                self.state["pcOutput"] = str(Path(value).resolve() / "执行结果")
            elif kind == "pcOutput":
                self.state["pcOutputAuto"] = False
            label = {
                "pcCurDir": "当期目录", "pcPreDir": "上期目录",
                "pcCentral": "大集中数据", "pcConfig": "跨期比较配置", "pcOutput": "输出目录",
            }[kind]
            self._log(f"跨期比较：{label}已选择")
            self._save_settings()
        self._refresh_period_pairs()
        return self.state.get(kind, "")

    def start_period_compare(self) -> bool:
        """报表采集系统页：用页面已选择的三路径启动跨期比较。"""
        if self.state["busy"]:
            return False
        cur = self.state.get("pcCurDir")
        pre = self.state.get("pcPreDir")
        if not cur or not pre:
            self._log("跨期比较失败：请先选择本期待处理与上期待处理目录")
            return False
        current_dir = Path(cur)
        previous_dir = Path(pre)
        central_raw = self.state.get("pcCentral")
        central_path = Path(central_raw) if central_raw else None
        if central_path is not None and not central_path.is_file():
            self._log(f"跨期比较失败：大集中数据文件不存在：{central_path}")
            return False
        output_dir = Path(self.state.get("pcOutput") or (current_dir / "执行结果"))
        self.state["busy"] = True
        self.state["status"] = "正在执行跨期比较，请勿关闭窗口……"
        self._log(f"开始跨期比较：当期“{current_dir.name}” vs 上期“{previous_dir.name}”")
        threading.Thread(
            target=self._period_compare_worker,
            args=(current_dir, previous_dir, central_path, output_dir),
            daemon=True,
        ).start()
        return True

    def _period_compare_worker(
        self,
        current_dir: Path,
        previous_dir: Path,
        central_path: Path | None,
        output_dir: Path,
    ) -> None:
        from .period_compare import (
            PeriodCompareError,
            ensure_default_config,
            run_period_compare,
        )

        started = time.monotonic()
        config_path = Path(self.state.get("pcConfig") or self.project_root / "跨期比较配置.xlsx")
        try:
            if ensure_default_config(config_path):
                self._log(f"首次使用：已生成默认配置模板 {config_path.name}，可按需维护指标/机构/警戒区间")
            output_path = run_period_compare(
                current_dir=current_dir,
                previous_dir=previous_dir,
                central_path=central_path,
                output_dir=output_dir,
                config_path=config_path,
                on_step=self._log_detail,
            )
            self.state["status"] = "跨期比较完成"
            self._log(f"跨期比较完成：{output_path}")
        except PeriodCompareError as exc:
            self.state["status"] = "跨期比较失败"
            self._log("跨期比较失败：" + str(exc))
        except Exception as exc:
            self.state["status"] = "跨期比较失败"
            self._log("跨期比较失败：" + str(exc) + "\n" + traceback.format_exc())
        finally:
            elapsed = time.monotonic() - started
            self._log(f"任务结束，本次耗时：{elapsed:.1f} 秒")
            self.state["busy"] = False

    def update(self, values: dict[str, Any]) -> dict[str, Any]:
        for key in ("input", "templateDir", "template", "external", "output"):
            if key in values:
                new_value = str(values[key]).strip()
                previous = self.state.get(key, "")
                if key == "output" and new_value != self.state.get("output", ""):
                    self.state["outputAuto"] = False
                    self.state["outputPinned"] = True
                self.state[key] = new_value
                if key == "template" and new_value and new_value != previous:
                    self.state["templateManual"] = True
                elif key == "templateDir" and new_value != previous:
                    self.state["templateManual"] = False
        if "outputPinned" in values:
            value = values["outputPinned"]
            pinned = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self.state["outputPinned"] = pinned
            self.state["outputAuto"] = not pinned
            if not pinned and self.state["input"]:
                self.state["output"] = str(Path(self.state["input"]) / "执行结果")
        if "recursiveDepth" in values:
            raw = values["recursiveDepth"]
            try:
                depth = int(raw)
            except (TypeError, ValueError):
                depth = -1 if str(raw).strip() in {"最深处", "true", "1", "是"} else 0
            if depth != self.state["recursiveDepth"]:
                self.state["recursiveDepth"] = depth
        if "recursive" in values:
            value = values["recursive"]
            recursive = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            depth = -1 if recursive else 0
            changed = depth != self.state["recursiveDepth"]
            self.state["recursiveDepth"] = depth
            if changed:
                # 复选框直接决定待审核清单的扫描范围；不触发模板自动改写。
                self._recognize(allow_template_auto=False)
        if "writeFlowLogs" in values:
            value = values["writeFlowLogs"]
            self.state["writeFlowLogs"] = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self._save_settings()
        if "confirmBeforeRun" in values:
            value = values["confirmBeforeRun"]
            self.state["confirmBeforeRun"] = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self._save_settings()
        if "calculationEngine" in values:
            candidate = str(values["calculationEngine"]).strip()
            if candidate not in valid_engine_values():
                self._log("计算引擎只能选择：" + "、".join(sorted(valid_engine_values())))
            else:
                self.state["calculationEngine"] = candidate
                self._save_settings()
        return self.state

    def recognize(self) -> dict[str, Any]:
        self._recognize(force_template=True)
        return self.state

    def set_selected_files(self, paths: list[str]) -> dict[str, Any]:
        allowed = {item["path"] for item in self.state.get("sourceFiles", [])}
        self.state["selectedFiles"] = [path for path in paths if path in allowed]
        return self.state

    def start(self, action: str, strict: bool = True) -> bool:
        if self.state["busy"]:
            return False
        self.state["busy"] = True
        self.state["status"] = "正在处理，请勿关闭窗口……"
        # Python 3.7 is used by the Win7 package; str.removeprefix arrived in
        # Python 3.9, so use slicing here instead.
        display = action[len("flow:"):] if action.startswith("flow:") else action
        action_name = {"check": "审核前检查", "audit": "汇总核查表校验", "summary": "汇总校验结果说明"}.get(display, display)
        self._log(f"开始执行：{action_name}")
        threading.Thread(target=self._worker, args=(action, strict), daemon=True).start()
        return True

    def _recognize(
        self,
        *,
        force_template: bool = False,
        allow_template_auto: bool = False,
    ) -> None:
        input_dir = Path(self.state["input"]) if self.state["input"] else None
        template_dir = Path(self.state["templateDir"]) if self.state["templateDir"] else None
        if input_dir and input_dir.is_dir():
            period = detect_period(input_dir, recursive=self.state["recursiveDepth"])
            self.state["detectedPeriod"] = period.period
            self.state["explanationFiles"] = [
                {"path": str(path), "name": path.name}
                for path in explanation_files(
                    input_dir, recursive=self.state["recursiveDepth"]
                )
            ]
            if not self.state["outputPinned"]:
                self.state["output"] = str(input_dir / "执行结果")
                self.state["outputAuto"] = True
        elif self.state.get("explanationFiles"):
            self.state["explanationFiles"] = []
        if input_dir and template_dir and input_dir.is_dir() and template_dir.is_dir():
            files = classify_source_files(
                template_dir, input_dir, self.project_root / "data" / "模板索引.json",
                recursive=self.state["recursiveDepth"],
            )
            self.state["sourceFiles"] = files
            current = set(self.state.get("selectedFiles") or [])
            available = {item["path"] for item in files}
            self.state["selectedFiles"] = (
                [item["path"] for item in files]
                if not current else [path for path in current if path in available]
            )
            templates = sorted({item["template"] for item in files if item["template"] != "未识别"})
            self.state["mixedTemplates"] = templates
            # 自动匹配只能由两处触发：用户点击“自动识别模板”，或选择源数据目录。
            # 执行流程、刷新目录、选择模板目录等场景只更新清单，不得改写模板选择。
            should_recommend = force_template or (
                allow_template_auto and not self.state.get("templateManual")
            )
            if should_recommend:
                result = recommend_template(
                    template_dir, input_dir, self.project_root / "data" / "模板索引.json",
                    recursive=self.state["recursiveDepth"],
                )
                if result.template_path:
                    self.state["template"] = str(result.template_path)
                    self.state["templateManual"] = False
                self._log(result.details)
            if len(templates) > 1:
                self._log("发现多种报表：" + "、".join(templates) + "。建议取消勾选不属于本次审核类型的文件。")
            # 用户手动追加的文件（源数据目录之外）并入清单，同样参与模板匹配。
            extra = [Path(path) for path in self.state.get("extraFiles", []) if Path(path).is_file()]
            known = {item["path"] for item in files}
            fresh = [path for path in extra if str(path.resolve()) not in known and path.parent != input_dir]
            for path in fresh:
                rows = classify_source_files(
                    template_dir, path.parent, self.project_root / "data" / "模板索引.json",
                    recursive=False,
                )
                for row in rows:
                    if Path(row["path"]).resolve() == path.resolve() and row["path"] not in known:
                        files.append(row)
                        known.add(row["path"])

    def _worker(self, action: str, strict: bool = True) -> None:
        started = time.monotonic()
        try:
            # 执行前仅刷新待处理文件、数据期等状态，不重新匹配或改写模板。
            self._log("正在刷新待审核文件清单……")
            self._recognize(allow_template_auto=False)
            self._log(
                f"待处理文件：{len(self.state.get('selectedFiles', []))} 个；"
                "正在读取执行流程……"
            )
            self._save_settings()
            service = AuditService(
                config_path=self.history_path,
                engine_preference=str(self.state["calculationEngine"]),
            )
            output = Path(self.state["output"])
            selected = [Path(path) for path in self.state.get("selectedFiles", [])]
            if action in {"check", "audit"} and not selected:
                raise ValueError("请至少勾选一个待审核文件")
            if action == "check":
                result = service.preflight(template_path=Path(self.state["template"]), input_dir=Path(self.state["input"]), output_dir=output, selected_files=selected, external_path=Path(self.state["external"]) if self.state["external"] else None, extra_files=[Path(p) for p in self.state.get("extraFiles", [])], recursive=self.state["recursiveDepth"], on_step=self._log_detail)
            else:
                if action == "audit":
                    action = "flow:汇总核查表校验"
                if action.startswith("flow:"):
                    flow_name = action[len("flow:"):]
                else:
                    flow_name = "汇总校验结果说明"
                standalone_combine_flow = service.is_standalone_combine_flow(flow_name)
                if standalone_combine_flow:
                    self._log(f"当前流程：{flow_name}；仅使用源数据目录")
                else:
                    self._log(f"当前流程：{flow_name}；模板：{Path(self.state['template']).name}")
                self._log("正在启动流程处理引擎……", detail=True)
                result = service.run_flow(
                    flow_name=flow_name,
                    template_path=None if standalone_combine_flow else Path(self.state["template"]),
                    input_dir=Path(self.state["input"]), output_dir=output,
                    period="" if standalone_combine_flow else (self.state.get("detectedPeriod") or Path(self.state["input"]).name),
                    history_path=self.history_path, selected_files=selected,
                    external_path=None if standalone_combine_flow else (Path(self.state["external"]) if self.state["external"] else None),
                    extra_files=[] if standalone_combine_flow else [Path(p) for p in self.state.get("extraFiles", [])],
                    recursive=self.state["recursiveDepth"], on_step=self._log_detail, strict=strict,
                    write_flow_logs=bool(self.state["writeFlowLogs"]),
                )
            self.state["status"] = result.summary_text().splitlines()[0]
            self._log(result.summary_text())
        except Exception as exc:
            self.state["status"] = "执行失败"
            self._log("执行失败：" + str(exc))
        finally:
            elapsed = time.monotonic() - started
            self._log(f"任务结束，本次耗时：{elapsed:.1f} 秒")
            self.state["busy"] = False

    def win_minimize(self) -> None:
        """自绘标题栏的窗口控制：最小化。"""
        import webview
        if webview.windows:
            webview.windows[0].minimize()

    def win_maximize(self, restore: bool = False) -> None:
        """自绘标题栏的窗口控制：最大化 / 还原。restore=True 时还原。"""
        import webview
        if not webview.windows:
            return
        window = webview.windows[0]
        if restore:
            window.restore()
        else:
            window.maximize()

    def win_close(self) -> None:
        """自绘标题栏的窗口控制：关闭窗口。"""
        import webview
        if webview.windows:
            webview.windows[0].destroy()

    def _log(self, text: str, detail: bool = False) -> None:
        """记录一条运行日志。detail=False 为面向用户的简明结果，detail=True
        为逐功能/逐文件的调试细节，前端“简单”模式会过滤掉 detail 条目。"""
        stamp = datetime.now().strftime("%H:%M:%S")
        entry = {"text": f"[{stamp}] {text}", "detail": bool(detail)}
        self.state["log"] = (self.state["log"] + [entry])[-100:]

    def _log_detail(self, text: str) -> None:
        self._log(text, detail=True)

    def _save_settings(self) -> None:
        self.settings.last_input_dir = self.state["input"]
        self.settings.last_template_dir = self.state["templateDir"]
        self.settings.last_external_file = self.state["external"]
        self.settings.last_output_dir = self.state["output"]
        self.settings.output_pinned = bool(self.state["outputPinned"])
        self.settings.recursive_depth = int(self.state["recursiveDepth"])
        self.settings.recursive_folders = self.settings.recursive_depth != 0
        self.settings.write_flow_logs = bool(self.state["writeFlowLogs"])
        self.settings.confirm_before_run = bool(self.state["confirmBeforeRun"])
        self.settings.calculation_engine = str(self.state["calculationEngine"])
        self.settings.period_current_dir = str(self.state.get("pcCurDir") or "")
        self.settings.period_previous_dir = str(self.state.get("pcPreDir") or "")
        self.settings.period_central_file = str(self.state.get("pcCentral") or "")
        self.settings.period_config_file = str(self.state.get("pcConfig") or "")
        self.settings.period_output_dir = str(self.state.get("pcOutput") or "")
        self.settings.period_output_auto = bool(self.state.get("pcOutputAuto", True))
        self.settings_store.save(self.settings)


def launch_web(project_root: Path) -> None:
    bundle_root = Path(getattr(sys, "_MEIPASS", project_root))
    # Python.NET 3 normally discovers the interpreter DLL from a regular
    # Python installation.  In a PyInstaller one-file build (in particular
    # the Python 3.7 Win7 package) that DLL lives in the temporary _MEI
    # directory instead.  Tell the CLR bridge its exact location before
    # pywebview imports ``clr``; otherwise pywebview hides the real loader
    # error behind the misleading "pythonnet is not installed" message.
    if getattr(sys, "frozen", False) and sys.platform == "win32":
        python_dll = bundle_root / ("python{}{}.dll".format(sys.version_info[0], sys.version_info[1]))
        if python_dll.is_file():
            os.environ.setdefault("PYTHONNET_PYDLL", str(python_dll))
        os.environ.setdefault("PYTHONNET_RUNTIME", "netfx")

        # pywebview turns every WinForms import failure into the generic
        # "pythonnet is not installed" message.  Verify the bridge first and
        # leave a diagnosable report beside the EXE when an old Windows system
        # cannot load one of its native/CLR dependencies.
        try:
            import clr

            clr.AddReference("System.Windows.Forms")
        except Exception as exc:
            report = project_root / "Win7界面启动诊断.txt"
            details = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__))
            try:
                report.write_text(
                    "基础数据审核工具界面启动诊断\n"
                    "Python.NET / WinForms 初始化失败。\n\n"
                    + details,
                    encoding="utf-8",
                )
            except OSError:
                pass
            raise RuntimeError(
                "Windows 界面组件初始化失败。请将 EXE 同目录的“Win7界面启动诊断.txt”发给维护人员。"
            ) from exc

    import webview

    initialize_config(project_root / HISTORY_WORKBOOK_NAME)
    hide_application_data_directory(project_root / "data")
    html = bundle_root / "web" / "index.html" if getattr(sys, "frozen", False) else project_root / "frontend" / "web" / "index.html"
    if not html.is_file():
        raise RuntimeError("本地界面文件缺失")
    page = _bridge_alias_page(html)
    # 使用操作系统原生标题栏（最小化/最大化/关闭由系统提供），界面内不再自绘。
    webview.create_window("基础数据审核工具", page.as_uri(), js_api=WebApi(project_root), width=1180, height=820, min_size=(900, 650), frameless=False)
    webview.start(gui="edgechromium")


def _bridge_alias_page(html: Path) -> Path:
    """为真 pywebview 窗口生成带别名垫片的临时页面。

    前端统一调用 ``bridge.api.方法名(...)`` 并等待 ``bridgeready``；pywebview
    原生只提供 ``window.pywebview`` 和 ``pywebviewready``。这里在 <head> 注入
    别名脚本把两者桥接起来，共用页面文件本身保持与 shell-flask 完全一致。
    """
    import tempfile

    source = html.read_text(encoding="utf-8")
    if "</head>" not in source:
        raise RuntimeError("本地界面文件缺失 <head>，无法注入桥接别名")
    alias = (
        "<script>window.addEventListener('pywebviewready',function(){"
        "window.bridge=window.pywebview;"
        "window.dispatchEvent(new Event('bridgeready'));"
        "});</script>"
    )
    handle, name = tempfile.mkstemp(prefix="audit_bridge_", suffix=".html")
    import os as _os

    with _os.fdopen(handle, "w", encoding="utf-8") as stream:
        stream.write(source.replace("</head>", alias + "</head>", 1))
    return Path(name)
