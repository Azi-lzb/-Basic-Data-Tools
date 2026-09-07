"""审核前检查的原生（openpyxl）实现。

与 COM 版同语义：命名区域调查（缺失也记录到运行日志）、模板体检（结构化
审核规则表校验在 read_template 内完成）、外部文件规划核对、逐文件表结构
比对并可选写检查报告。全程不启动 LibreOffice/Excel。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

from openpyxl import load_workbook

from ..external import make_external_sheet_plan
from ..feature_log import FeatureLog
from ..models import PreflightItem, PreflightRunResult
from ..name_config import FeatureMapping
from ..preflight_xlsx import (
    validate_source_xlsx,
    write_preflight_report_xlsx,
    write_structure_report_xlsx,
)
from ..template import TemplateError
from .openpyxl_workbook import (
    named_ranges,
    read_template,
    template_formulas,
    template_structure_values,
)


def _external_plan_items(
    template_path: Path, external_path: Path, definition
) -> list[PreflightItem]:
    """外部文件规划核对：按模板公式识别应复制的工作表并核对存在性。"""
    external = load_workbook(external_path, read_only=True, data_only=False, keep_links=False)
    try:
        plan = make_external_sheet_plan(
            formulas=template_formulas(template_path),
            available_sheets=external.sheetnames,
        )
        missing = [name for name in plan.sheet_names if name not in external.sheetnames]
        status = "不通过" if missing else "通过"
        detail = "将复制 {} 个工作表（{}）：{}".format(
            len(plan.sheet_names), plan.source, "、".join(plan.sheet_names)
        )
        if missing:
            detail += "；缺失：" + "、".join(missing)
        return [PreflightItem("外部文件检查", "错误" if missing else "提示", status, "", "", detail)]
    finally:
        external.close()


def run_native_preflight(
    *,
    template_path: Path,
    input_dir: Path,
    output_dir: Path,
    config_path: Path | None = None,
    selected_files: list[Path] | None = None,
    external_path: Path | None = None,
    recursive: bool = False,
    write_report: bool = True,
    on_step: Optional[Callable[[str], None]] = None,
    named_range_features: tuple[tuple[int, FeatureMapping], ...] = (),
    feature_log: Optional[FeatureLog] = None,
) -> PreflightRunResult:
    say = on_step or (lambda _text: None)
    template_path = template_path.resolve()
    from ..discovery import source_workbooks

    sources = source_workbooks(input_dir, recursive=recursive)
    if selected_files is not None:
        allowed = {path.resolve() for path in sources}
        sources = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not sources:
        raise ValueError("源数据目录中没有可检查的 .xlsx 文件")
    output_dir.mkdir(parents=True, exist_ok=True)
    batch_id = datetime.now().strftime("%Y%m%d%H%M%S")

    template_items: list[PreflightItem] = []
    definition = None
    try:
        mappings = [mapping for _order, mapping in named_range_features]
        # 先逐映射调查（缺失也记录），再读模板做强校验。
        if mappings:
            book = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
            try:
                for mapping in mappings:
                    areas = named_ranges(book, [mapping])
                    rows = [(area.sheet_name, area.address, "通过") for area in areas]
                    if feature_log is not None:
                        feature_log.add_sheet(mapping.name, ("工作表", "命名区域名", "覆盖区域", "结果"),
                                              rows or [("", "、".join(mapping.range_names), "缺失")])
                    if areas:
                        per_sheet: dict[str, int] = {}
                        for area in areas:
                            per_sheet[area.sheet_name] = per_sheet.get(area.sheet_name, 0) + 1
                        detail = "、".join("{} {} 处".format(sheet, count) for sheet, count in per_sheet.items())
                        say("{}：{}".format(mapping.name, detail))
                        template_items.append(PreflightItem(
                            "命名区域检查", "提示", "通过", "", "",
                            "模块“{}”已找到 {} 个区域".format(mapping.name, len(areas)),
                        ))
                    else:
                        message = "模块“{}”缺少命名区域：{}".format(mapping.name, "、".join(mapping.range_names))
                        say(message)
                        template_items.append(PreflightItem("命名区域检查", "错误", "不通过", "", "", message))
            finally:
                book.close()

        definition = read_template(template_path, formula_mappings=mappings)
        template_items.append(PreflightItem(
            "模板体检", "提示", "通过", "", "",
            "规则 {} 条（{}）；结构区域 {} 个".format(
                len(definition.rules), "结构化审核规则表" if definition.structured else "命名区域模板",
                len(definition.structure_ranges),
            ),
        ))
        if definition.structure_ranges:
            expected = template_structure_values(template_path, definition)
            template_items.append(PreflightItem(
                "表结构区域", "提示", "通过", "", "",
                "已读取 {} 个结构区域用于文件比对".format(len(expected)),
            ))
        if external_path is not None:
            template_items.extend(_external_plan_items(template_path, external_path, definition))
    except TemplateError as exc:
        template_items = [PreflightItem("总体", "错误", "不通过", "", "", str(exc))]
        report_path = output_dir / "审核前检查_{}.xlsx".format(batch_id)
        if feature_log is not None:
            feature_log.add_template_health(template_path, template_items)
        if write_report:
            write_preflight_report_xlsx(report_path, template_path, template_items)
        raise TemplateError(
            "模板体检不通过：{}".format(exc)
            + (f"。体检报告：{report_path}" if write_report else "")
        ) from exc

    if feature_log is not None:
        feature_log.add_template_health(template_path, template_items)

    # 逐文件表结构比对（纯 openpyxl；旧 .xls 记为跳过）。
    source_matches = []
    expected_structure = template_structure_values(template_path, definition) if definition and definition.structure_ranges else []
    for source in sources:
        if source.suffix.casefold() == ".xls":
            from ..models import SourceMatch
            source_matches.append((source, SourceMatch(
                matched=False, score=0.0, matched_labels=0, checked_labels=0,
                missing_sheets=(), details="旧版 .xls 仅支持汇总读取；请先转换为 .xlsx",
            )))
            continue
        if expected_structure:
            match = validate_source_xlsx(source, definition, expected_structure, external_sheet_names=())
        else:
            # 模板没有“表结构区域”：没有需要比对的固定标签，视作通过（与 COM 版一致）。
            from ..models import SourceMatch
            match = SourceMatch(matched=True, score=1.0, matched_labels=0, checked_labels=0,
                                missing_sheets=(), details="模板未配置表结构区域，未做结构比对")
        source_matches.append((source, match))
        say("{}：{}".format(source.name, "通过" if match.matched else "不通过"))

    structure_step_done = bool(definition and definition.structure_ranges)
    if structure_step_done:
        say("完成：表结构比对")

    report_path = output_dir / "审核前检查_{}.xlsx".format(batch_id)
    if write_report:
        write_preflight_report_xlsx(report_path, template_path, template_items)
        if structure_step_done:
            write_structure_report_xlsx(
                output_dir / "表结构比对_{}.xlsx".format(batch_id), source_matches
            )

    return PreflightRunResult(
        report_path=report_path if write_report else None,
        total_files=len(source_matches),
        matched_files=sum(1 for _, item in source_matches if item.matched),
        template_warnings=sum(1 for item in template_items if item.level in {"提示", "警告"}),
        source_matches=tuple(source_matches),
        batch_id=batch_id,
    )
