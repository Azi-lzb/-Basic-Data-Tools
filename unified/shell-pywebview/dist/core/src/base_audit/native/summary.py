"""原生汇总模块（UOS/麒麟）：任意行/固定行汇总、汇总表合并、组合工作表。

任意行/固定行汇总对齐 Windows COM 版输出语义：来源文件列写“去日期”的
机构名（``normalize_template_name``），含全部身份字段的表末尾追加本地校验
历史三列（历史触发条数/历史分类说明/审核结果，复合键多行用 ``|`` 拼接）。
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import re
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

from ..discovery import source_workbooks
from ..history import (
    LOCAL_VALIDATION_HISTORY_FIELDS,
    LOCAL_VALIDATION_IDENTITY_FIELDS,
    join_local_validation_history,
)
from ..name_config import (
    FIXED_ROW_SUMMARY_FUNCTION,
    USED_RANGE_SUMMARY_FUNCTION,
    FeatureMapping,
)
from ..template import normalize_template_name
from .openpyxl_workbook import named_ranges, named_single_cells_with_prefix


def _bounds(address: str) -> tuple[int, int, int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(address)
    return min_row, min_col, max_row, max_col


def _header_mapping(feature: FeatureMapping) -> FeatureMapping:
    return FeatureMapping(feature.name, feature.feature_type, ("表头区域",), False, "", "")


def _field_key(value: object) -> str:
    return str(value or "").replace(" ", "").replace("\n", "").casefold()


def _sheet_name(base: str, used: set[str]) -> str:
    text = re.sub(r"[\\/*?:\[\]]", "_", str(base))[:31] or "汇总"
    result = text
    number = 2
    while result.casefold() in used:
        result = "{}_{}".format(text[:28], number)
        number += 1
    used.add(result.casefold())
    return result


def _joins_local_validation_history(base_headers: list[str]) -> bool:
    """历史三列只对含全部身份字段的表追加；“业务说明”等普通汇总表保持原样。"""
    keys = {_field_key(value) for value in base_headers}
    return all(_field_key(name) in keys for name in LOCAL_VALIDATION_IDENTITY_FIELDS)


def _read_header(template_book, data_range, headers) -> tuple[list[object], int, int]:
    """选与汇总区域同表且列范围覆盖最合适的表头区域，返回表头值与列范围。"""
    matching = [area for area in headers if area.sheet_name == data_range.sheet_name]
    if not matching:
        raise ValueError("工作表“{}”缺少命名区域“表头区域”".format(data_range.sheet_name))
    top, left, bottom, right = _bounds(data_range.address)
    best = None
    for item in matching:
        h_top, h_left, h_bottom, h_right = _bounds(item.address)
        if h_bottom < h_top:
            continue
        lo, hi = max(left, h_left), min(right, h_right)
        if lo > hi:
            continue
        if best is None or (hi - lo) > (best[2] - best[1]):
            best = (h_top, lo, hi, h_bottom)
    if best is None:
        raise ValueError("汇总区域与表头区域没有共同列")
    h_top, lo, hi, h_bottom = best
    sheet = template_book[data_range.sheet_name]
    # 多行表头自上而下用下划线拼接（与 Windows 版 _combine_header_rows 一致）。
    header_values: list[object] = []
    columns = range(lo, hi + 1)
    merged: dict[int, list[str]] = {}
    for row in sheet.iter_rows(min_row=h_top, max_row=h_bottom, min_col=lo, max_col=hi, values_only=True):
        for offset, value in enumerate(row):
            text = str(value or "").strip()
            if text:
                merged.setdefault(offset, []).append(text)
    for column in columns:
        offset = column - lo
        header_values.append("_".join(merged.get(offset, [])) or "")
    return header_values, lo, hi


def run_region_summaries(
    *, template_path: Path, input_dir: Path, output_dir: Path,
    features: Iterable[FeatureMapping], recursive: bool = True,
    selected_files: Iterable[Path] | None = None,
    on_step: Callable[[str], None] | None = None,
    output_name: str = "区域汇总",
    history_config_path: Path | None = None,
) -> Path:
    """运行配置的原生行汇总并返回输出工作簿路径。"""
    say = on_step or (lambda _text: None)
    features = [item for item in features if item.feature_type in {USED_RANGE_SUMMARY_FUNCTION, FIXED_ROW_SUMMARY_FUNCTION}]
    if not features:
        raise ValueError("没有启用“汇总_任意行汇总”或“汇总_固定行汇总”模块")
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files:
        allowed = {path.resolve() for path in files}
        files = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not files:
        raise FileNotFoundError("没有可汇总的源工作簿")
    template = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
    history_lookup: dict[tuple[str, ...], list[str]] = {}
    try:
        global_fields = named_single_cells_with_prefix(template, "全局单元格区域.")
        local_fields = named_single_cells_with_prefix(template, "单元格区域.")
        groups: dict[str, dict[str, object]] = {}
        for feature in features:
            say("正在汇总：{}".format(feature.name))
            data_ranges = named_ranges(template, (feature,))
            headers = named_ranges(template, (_header_mapping(feature),))
            if not data_ranges:
                raise ValueError("汇总功能“{}”未找到命名区域：{}".format(feature.name, "、".join(feature.range_names)))
            if not headers:
                raise ValueError("汇总功能“{}”缺少命名区域“表头区域”".format(feature.name))
            for data_range in data_ranges:
                header_values, left, right = _read_header(template, data_range, headers)
                local = [(name, area) for name, area in local_fields if area.sheet_name == data_range.sheet_name]
                extras = [*global_fields, *local]
                fields = [name for name, _ in extras]
                key = data_range.sheet_name.casefold()
                base_headers = ["来源文件", "来源工作表", *fields, *header_values]
                join_history = _joins_local_validation_history(base_headers)
                out_headers = [*base_headers, *LOCAL_VALIDATION_HISTORY_FIELDS] if join_history else base_headers
                group = groups.setdefault(key, {
                    "title": data_range.sheet_name, "headers": out_headers, "rows": [],
                })
                top, _left, bottom, _right = _bounds(data_range.address)
                for file_path in files:
                    if file_path.suffix.casefold() == ".xls":
                        continue
                    book = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
                    try:
                        if data_range.sheet_name not in book.sheetnames:
                            continue
                        sheet = book[data_range.sheet_name]
                        last_row = bottom if feature.feature_type == FIXED_ROW_SUMMARY_FUNCTION else max(top, sheet.max_row)
                        values = list(sheet.iter_rows(min_row=top, max_row=last_row, min_col=left, max_col=right, values_only=True))
                        metadata = []
                        for field, area in extras:
                            if area.sheet_name not in book.sheetnames:
                                raise ValueError("源文件“{}”缺少单元格区域字段“{}”所在工作表“{}”".format(file_path.name, field, area.sheet_name))
                            row, col = _bounds(area.address)[:2]
                            metadata.append(book[area.sheet_name].cell(row, col).value)
                        for row in values:
                            if any(value not in (None, "") for value in row):
                                base_row = [normalize_template_name(file_path.stem), data_range.sheet_name, *metadata, *row]
                                if join_history:
                                    from ..history import _composite_key
                                    identity = _composite_key(base_headers, base_row, LOCAL_VALIDATION_IDENTITY_FIELDS)
                                    if not history_lookup:
                                        history_lookup = _load_local_validation_history(history_config_path)
                                    joined = history_lookup.get(identity)
                                    extra = joined if joined is not None else [""] * len(LOCAL_VALIDATION_HISTORY_FIELDS)
                                    group["rows"].append([*base_row, *extra])
                                else:
                                    group["rows"].append(base_row)
                    finally:
                        book.close()
            say("完成：{}".format(feature.name))
    finally:
        template.close()
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "{}_{}.xlsx".format(output_name, datetime.now().strftime("%Y%m%d_%H%M%S"))
    output, used_names = Workbook(), set()
    try:
        first = True
        for group in groups.values():
            sheet = output.active if first else output.create_sheet()
            first = False
            sheet.title = _sheet_name(str(group["title"]), used_names)
            sheet.append(group["headers"])
            for row in group["rows"]:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(40, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        if first:
            output.active.title = "汇总"
        output.save(path)
    finally:
        output.close()
    return path


def _load_local_validation_history(config_path: Path | None) -> dict[tuple[str, ...], list[str]]:
    """读取「历史本地校验结果审核」表并按复合键聚合历史三列。"""
    from ..history import LOCAL_VALIDATION_HISTORY_SHEET
    from ..history import join_local_validation_history as _join

    if config_path is None or not Path(config_path).is_file():
        return {}
    book = load_workbook(config_path, read_only=True, data_only=True, keep_links=False)
    try:
        if LOCAL_VALIDATION_HISTORY_SHEET not in book.sheetnames:
            return {}
        rows = book[LOCAL_VALIDATION_HISTORY_SHEET].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        return _join(headers, [list(row) for row in rows])
    finally:
        book.close()


_INVALID = re.compile(r"[\[\]:*?/\\]")


def _copy_sheet(source, target, name, style_cache=None):
    """跨工作簿复制工作表；样式经缓存去重注册（与 Windows 实现同思路）。"""
    from copy import copy as _copy

    target_sheet = target.create_sheet(name)
    target_sheet.sheet_state = source.sheet_state
    target_sheet.freeze_panes = source.freeze_panes

    def cached_style(cell):
        if style_cache is None:
            return _copy(cell._style)
        key = cell.style_id if hasattr(cell, "style_id") else str(cell._style)
        if key not in style_cache:
            style_cache[key] = _copy(cell._style)
        return style_cache[key]

    for row in source.iter_rows():
        for cell in row:
            destination = target_sheet[cell.coordinate]
            destination.value = cell.value
            if cell.has_style:
                destination._style = cached_style(cell)
            if cell.comment:
                destination.comment = _copy(cell.comment)
            if cell.hyperlink:
                destination._hyperlink = _copy(cell.hyperlink)
    for key, dimension in source.row_dimensions.items():
        target_sheet.row_dimensions[key].height = dimension.height
        target_sheet.row_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.column_dimensions.items():
        target_sheet.column_dimensions[key].width = dimension.width
        target_sheet.column_dimensions[key].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        target_sheet.merge_cells(str(merged))


def _unique_sheet_name(value: str, used: set[str]) -> str:
    base = _INVALID.sub("_", value)[:31] or "工作表"
    result = base
    number = 2
    while result.casefold() in used:
        tail = "_" + str(number)
        result = base[:31 - len(tail)] + tail
        number += 1
    used.add(result.casefold())
    return result


def merge_by_first_filename_part(
    *, input_dir: Path, output_dir: Path, recursive: bool = True, selected_files=None,
    output_name: str = "组合联合核查表", on_step=None,
):
    """`修改_组合联合核查表`：按文件名第一段合并同一机构的工作簿。"""
    say = on_step or (lambda _text: None)
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files:
        allowed = {path.resolve() for path in files}
        files = [Path(p) for p in selected_files if Path(p).resolve() in allowed]
    groups: dict[str, list[Path]] = defaultdict(list)
    for path in files:
        groups[path.stem.split("_", 1)[0] or path.stem].append(path)
    target_dir = output_dir / (str(output_name or "组合联合核查表") + "_" + datetime.now().strftime("%Y%m%d%H%M%S"))
    target_dir.mkdir(parents=True, exist_ok=True)
    records = []
    for group, paths in sorted(groups.items()):
        out = target_dir / (group + "_合并_" + (datetime.now().strftime("%Y%m") or "本期") + ".xlsx")
        book = Workbook()
        del book[book.sheetnames[0]]
        used: set[str] = set()
        copied: list[str] = []
        try:
            for path in paths:
                source = load_workbook(path, data_only=False, keep_links=False)
                try:
                    style_cache: dict[str, object] = {}
                    for sheet in source.worksheets:
                        name = _unique_sheet_name(sheet.title, used)
                        _copy_sheet(sheet, book, name, style_cache)
                        copied.append(path.name + "｜" + sheet.title)
                finally:
                    source.close()
            if not book.sheetnames:
                book.create_sheet("空工作表")
            book.save(out)
            records.append((group, out.name, "；".join(copied), "成功"))
            say("已组合：{}（合并源文件 {} 个）".format(group, len(paths)))
        except Exception as exc:
            records.append((group, "", "", "失败：" + str(exc)))
        finally:
            book.close()
    return target_dir, records


def merge_workbook_tables(
    *, input_dir: Path, output_dir: Path, recursive: bool = True,
    output_name: str = "汇总表合并", selected_files: Iterable[Path] | None = None,
    on_step=None,
) -> Path:
    """`汇总_汇总表合并`：仅首行表头完全一致的工作表合并为一张表。"""
    say = on_step or (lambda _text: None)
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files is not None:
        allowed = {path.resolve() for path in files}
        files = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not files:
        raise FileNotFoundError("没有可合并的源工作簿")
    groups: dict[tuple[str, ...], dict[str, object]] = {}
    for path in files:
        if path.suffix.casefold() == ".xls":
            continue  # 旧版 .xls 由 Windows 版处理；原生管线聚焦 .xlsx
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = tuple(next(rows, ()))
                if not header or not any(value not in (None, "") for value in header):
                    continue
                signature = tuple(str(value or "").strip() for value in header)
                group = groups.setdefault(signature, {"headers": ["来源文件", "来源工作表", *header], "rows": [], "sheets": []})
                if sheet.title not in group["sheets"]:
                    group["sheets"].append(sheet.title)
                for row in rows:
                    if any(value not in (None, "") for value in row):
                        group["rows"].append([path.name, sheet.title, *row])
        finally:
            book.close()
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "{}_{}.xlsx".format(output_name, datetime.now().strftime("%Y%m%d_%H%M%S"))
    output, used = Workbook(), set()
    try:
        first = True
        for group in groups.values():
            sheet = output.active if first else output.create_sheet()
            first = False
            sheet.title = _sheet_name("_".join(group["sheets"][:2]) or "汇总", used)
            sheet.append(group["headers"])
            for row in group["rows"]:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        if first:
            output.active.title = "汇总"
        output.save(path)
        say("汇总表合并完成：按表头分为 {} 张表".format(len(groups)))
    finally:
        output.close()
    return path
