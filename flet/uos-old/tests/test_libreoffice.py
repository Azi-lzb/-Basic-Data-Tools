from __future__ import annotations

import shutil
import sys
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import unquote, urlparse
from urllib.request import url2pathname

from openpyxl import Workbook, load_workbook

from uos_audit import libreoffice


def test_direct_engine_respects_configured_executable(monkeypatch) -> None:
    monkeypatch.setenv("BASE_AUDIT_ENGINE", "direct")
    monkeypatch.setenv("BASE_AUDIT_SOFFICE", sys.executable)
    engine = libreoffice.find_calc_engine()
    assert engine is not None
    assert engine.source == "环境变量 BASE_AUDIT_SOFFICE"
    assert engine.command_prefix == (sys.executable,)


def test_linglong_engine_uses_verified_command_prefix(monkeypatch) -> None:
    monkeypatch.setenv("BASE_AUDIT_ENGINE", "ll-cli")
    monkeypatch.setattr(libreoffice.shutil, "which", lambda name: "/usr/bin/ll-cli" if name == "ll-cli" else None)
    monkeypatch.setattr(
        libreoffice.subprocess, "run",
        lambda *args, **kwargs: SimpleNamespace(stdout="org.libreoffice.libreoffice 25.8", stderr="", returncode=0),
    )
    engine = libreoffice.find_calc_engine()
    assert engine is not None
    assert engine.source == "玲珑商店版"
    assert engine.command_prefix == ("/usr/bin/ll-cli", "run", "org.libreoffice.libreoffice", "--", "soffice")


def test_recalculate_stages_and_replaces_only_after_success(tmp_path: Path, monkeypatch) -> None:
    workbook_path = tmp_path / "审核版.xlsx"
    book = Workbook()
    book.active["A1"] = "=SUM(2,3)"
    book.save(workbook_path)
    book.close()
    commands: list[list[str]] = []

    def fake_run(command, **kwargs):
        commands.append(command)
        output_dir = Path(command[command.index("--outdir") + 1])
        staged_arg = command[-1]
        staged = Path(url2pathname(unquote(urlparse(staged_arg).path))) if staged_arg.startswith("file:") else Path(staged_arg)
        shutil.copy2(staged, output_dir / staged.name)
        return SimpleNamespace(stdout="已设置 GTK_PATH=/mock", stderr="GTK warning", returncode=0)

    monkeypatch.setattr(libreoffice.subprocess, "run", fake_run)
    result = libreoffice.LibreOfficeCalculator(Path(sys.executable)).recalculate(workbook_path)

    assert result.workbook_path == workbook_path
    assert commands[0][1:5] == ["--headless", "--nologo", "--nodefault", "--nolockcheck"]
    assert "-env:UserInstallation=file:" in " ".join(commands[0])
    restored = load_workbook(workbook_path, data_only=False)
    assert restored.active["A1"].value == "=SUM(2,3)"
    restored.close()
