from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from uos_audit.name_config import FeatureMapping, USED_RANGE_SUMMARY_FUNCTION
from uos_audit.region_summary import merge_workbook_tables, run_region_summaries


def test_used_range_summary_keeps_metadata_and_two_line_header(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "明细"
    sheet["A1"], sheet["B1"], sheet["A2"], sheet["B2"] = "贷款", "贷款", "余额", "利率"
    sheet["D1"] = "2026-08"
    template.defined_names.add(DefinedName("任意行汇总区域", attr_text="'明细'!$A$3:$B$9"))
    template.defined_names.add(DefinedName("表头区域", attr_text="'明细'!$A$1:$B$2"))
    template.defined_names.add(DefinedName("全局单元格区域.数据日期", attr_text="'明细'!$D$1"))
    template.save(template_path)
    template.close()
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source = Workbook()
    data = source.active
    data.title = "明细"
    data.append(["贷款", "贷款"])
    data.append(["余额", "利率"])
    data.append([10, 0.03])
    data["D1"] = "2026-08"
    source.save(source_dir / "机构A.xlsx")
    source.close()
    feature = FeatureMapping("任意行汇总", USED_RANGE_SUMMARY_FUNCTION, ("任意行汇总区域",), False, "", "")
    output = run_region_summaries(template_path=template_path, input_dir=source_dir, output_dir=tmp_path / "out", features=(feature,))
    book = load_workbook(output, data_only=True)
    sheet = book["明细"]
    assert [cell.value for cell in sheet[1]] == ["来源文件", "来源工作表", "数据日期", "贷款_余额", "贷款_利率"]
    assert [cell.value for cell in sheet[2]] == ["机构A.xlsx", "明细", "2026-08", 10, 0.03]
    book.close()


def test_workbook_table_merge_respects_selected_files(tmp_path: Path) -> None:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    for name, value in (("机构A.xlsx", "A"), ("机构B.xlsx", "B")):
        book = Workbook()
        sheet = book.active
        sheet.title = "数据"
        sheet.append(["字段"])
        sheet.append([value])
        book.save(source_dir / name)
        book.close()

    output = merge_workbook_tables(
        input_dir=source_dir, output_dir=tmp_path / "out",
        selected_files=[source_dir / "机构B.xlsx"], output_name="选定文件合并",
    )
    result = load_workbook(output, data_only=True)
    rows = list(result.active.iter_rows(values_only=True))
    result.close()

    assert rows[0] == ("来源文件", "来源工作表", "字段")
    assert rows[1] == ("机构B.xlsx", "数据", "B")
    assert len(rows) == 2
