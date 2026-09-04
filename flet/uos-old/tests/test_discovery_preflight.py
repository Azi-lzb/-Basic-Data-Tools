from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from uos_audit.audit import run_native_audit
from uos_audit.discovery import recommend_template, source_workbooks
from uos_audit.name_config import FeatureMapping, STRUCTURE_COMPARE_FUNCTION
from uos_audit.name_config import initialize_config, load_config_editor_data, save_config_editor_draft
from uos_audit.workflow import run_flow


def _workbook(path: Path, sheets: dict[str, object]) -> None:
    book = Workbook()
    first = True
    for name, value in sheets.items():
        sheet = book.active if first else book.create_sheet()
        sheet.title = name
        sheet["A1"] = value
        first = False
    book.save(path)
    book.close()


def test_template_recommendation_uses_sheet_set_and_skips_output_directories(tmp_path: Path) -> None:
    templates = tmp_path / "模板文件"
    templates.mkdir()
    _workbook(templates / "单位贷款模板.xlsx", {"单位表": "单位", "极端表": "极端"})
    _workbook(templates / "个人贷款模板.xlsx", {"个人表": "个人"})

    sources = tmp_path / "源数据"
    sources.mkdir()
    _workbook(sources / "机构A_报送.xlsx", {"单位表": 1, "极端表": 2})
    skipped = sources / "执行结果" / "历史"
    skipped.mkdir(parents=True)
    _workbook(skipped / "不应参与识别.xlsx", {"个人表": 1})

    suggestion = recommend_template(
        templates, sources, tmp_path / "data" / "模板索引.json", recursive=True
    )

    assert suggestion.matched
    assert suggestion.template_path is not None
    assert suggestion.template_path.name == "单位贷款模板.xlsx"
    assert [path.name for path in source_workbooks(sources, recursive=True)] == ["机构A_报送.xlsx"]


def test_template_recommendation_ignores_unmatched_summary_workbook(tmp_path: Path) -> None:
    templates = tmp_path / "模板文件"
    templates.mkdir()
    _workbook(templates / "单位贷款模板.xlsx", {"单位表": "单位", "极端表": "极端"})
    _workbook(templates / "个人贷款模板.xlsx", {"个人表": "个人"})
    sources = tmp_path / "源数据"
    sources.mkdir()
    # This name sorts before the institutions, recreating the former first-file
    # failure while the two actual raw reports still agree on one template.
    _workbook(sources / "00_汇总信息.xlsx", {"汇总": "不是报送表"})
    _workbook(sources / "机构A_报送.xlsx", {"单位表": 1, "极端表": 2})
    _workbook(sources / "机构B_报送.xlsx", {"单位表": 3, "极端表": 4})

    suggestion = recommend_template(
        templates, sources, tmp_path / "data" / "模板索引.json"
    )

    assert suggestion.matched
    assert suggestion.template_path is not None
    assert suggestion.template_path.name == "单位贷款模板.xlsx"
    assert "已忽略未能识别" in suggestion.details


def test_structure_mismatch_is_skipped_before_creating_audit_copy(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["A1"] = "模板固定表头"
    template.defined_names.add(DefinedName("表结构区域", attr_text="'报表'!$A$1"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    _workbook(source_dir / "机构A.xlsx", {"报表": "错误表头"})
    mapping = FeatureMapping(
        "表结构比对", STRUCTURE_COMPARE_FUNCTION, ("表结构区域",), False, "", "", "", ""
    )

    result = run_native_audit(
        input_dir=source_dir,
        template_path=template_path,
        output_dir=tmp_path / "执行结果",
        structure_mappings=[mapping],
        formula_mappings=[],
        extraction_mappings=[],
        include_external=False,
        include_formula_copy=False,
        include_formula_extraction=False,
        write_summary=False,
    )

    assert len(result.files) == 1
    assert "报送文件与模板不匹配" in result.files[0].error
    assert result.files[0].audit_path is None
    assert not (tmp_path / "执行结果" / "机构A_审核版.xlsx").exists()


def test_structure_step_can_explicitly_keep_its_own_report(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["A1"] = "固定表头"
    template.defined_names.add(DefinedName("表结构区域", attr_text="'报表'!$A$1"))
    template.save(template_path)
    template.close()
    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    _workbook(source_dir / "机构A.xlsx", {"报表": "固定表头"})

    history_path = tmp_path / "历史审核说明.xlsx"
    initialize_config(history_path)
    editor = load_config_editor_data(history_path)
    draft = {
        "customModules": [],
        "customFlows": [{
            "流程名": "结构报告", "顺序": "10", "功能名": "表结构比对",
            "启用": "是", "失败后处理": "跳过", "是否输出结果": "是",
            "输出文件名": "结构核对结果", "输入": "源数据目录", "输出": "运行日志", "备注": "",
        }],
        "flowDisplay": {},
        "combineSheetsPlans": editor["combineSheetsPlans"],
        "activeCombineSheetsPlanId": editor["activeCombineSheetsPlanId"],
    }
    save_config_editor_draft(history_path, draft)

    outcome = run_flow(
        flow_name="结构报告", template_path=template_path, input_dir=source_dir,
        output_dir=tmp_path / "执行结果", history_path=history_path,
        config_path=history_path, write_flow_logs=False,
    )

    report = Path(outcome["output"])
    assert report.is_file()
    assert report.name.startswith("结构核对结果_")
