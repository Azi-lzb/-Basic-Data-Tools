from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

from .models import Issue


HISTORY_HEADERS = (
    "工作簿名",
    "工作表名",
    "定位单元格",
    "错误类型",
    "校验指标",
    "描述",
    "当前值",
    "对比值",
    "差值",
    "规则编号",
    "历史校验说明",
    "审核意见",
)


def read_history_xlsx(path: Path, *, sheet_name: str = "历史审核记录") -> list[Issue]:
    """Read the concise business history sheet without relying on COM."""
    if not path.is_file():
        return []
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = book[sheet_name] if sheet_name in book.sheetnames else (
            book["历史审核结果"] if "历史审核结果" in book.sheetnames else None
        )
        if sheet is None:
            return []
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return []
        headers = {str(value or "").strip(): index for index, value in enumerate(rows[0])}

        def get(row: tuple[object, ...], name: str) -> object:
            index = headers.get(name)
            return row[index] if index is not None and index < len(row) else ""

        result: list[Issue] = []
        for row in rows[1:]:
            source_file, sheet_value, target_cell = (str(get(row, key) or "") for key in ("工作簿名", "工作表名", "定位单元格"))
            rule_id = str(get(row, "规则编号") or "")
            if not any((source_file, sheet_value, target_cell, rule_id)):
                continue
            indicator = str(get(row, "校验指标") or "")
            identity = rule_id or "｜".join((source_file, sheet_value, target_cell, indicator or "校验指标未识别"))
            result.append(Issue(
                issue_id=identity, period=str(get(row, "数据期") or ""), batch_id="", audit_time="",
                triggered=True, status="", first_seen_period="", previous_seen_period="", consecutive_count=1,
                org_code="", org_name="", report_code="", sheet_name=sheet_value, rule_id=identity,
                severity=str(get(row, "错误类型") or "错误"), formula_cell=target_cell, target_cell=target_cell,
                target_value=get(row, "当前值"), formula_result="", message=str(get(row, "描述") or ""),
                source_file=source_file, audit_file="", check_field=indicator, comparison_value=get(row, "对比值"),
                difference_value=get(row, "差值"), detail=str(get(row, "描述") or ""),
                institution_feedback=str(get(row, "历史校验说明") or ""), auditor_opinion=str(get(row, "审核意见") or ""),
            ))
        return result
    finally:
        book.close()


def write_history_xlsx(path: Path, issues: Iterable[Issue], *, sheet_name: str = "历史审核记录") -> None:
    """Replace only the program-managed history table, preserving other sheets."""
    path.parent.mkdir(parents=True, exist_ok=True)
    book = load_workbook(path) if path.is_file() else Workbook()
    try:
        if sheet_name in book.sheetnames:
            del book[sheet_name]
        sheet = book.create_sheet(sheet_name, 0)
        sheet.append(HISTORY_HEADERS)
        for item in issues:
            sheet.append((
                Path(item.source_file).name if item.source_file else "", item.sheet_name, item.target_cell,
                item.severity, item.check_field, item.detail or item.message, item.target_value,
                item.comparison_value, item.difference_value, item.issue_id,
                item.institution_feedback, item.auditor_opinion,
            ))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, len(HISTORY_HEADERS) + 1):
            sheet.column_dimensions[chr(64 + column)].width = 22
        if "Sheet" in book.sheetnames and len(book.sheetnames) > 1:
            del book["Sheet"]
        book.save(path)
    finally:
        book.close()


def _period_key(value: str) -> tuple[int, int, str]:
    text = value.strip()
    for fmt in ("%Y-%m", "%Y/%m", "%Y%m", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(text, fmt)
            return dt.year, dt.month, text
        except ValueError:
            pass
    if len(text) >= 6 and text[:4].isdigit() and "Q" in text.upper():
        try:
            quarter = int(text.upper().split("Q", 1)[1][:1])
            return int(text[:4]), quarter * 3, text
        except (ValueError, IndexError):
            pass
    return 0, 0, text


def classify_current_issues(
    current_period: str,
    current: list[Issue],
    history: Iterable[Issue],
    *,
    audited_org_codes: set[str] | None = None,
    batch_id: str = "",
    audit_time: str = "",
) -> tuple[list[Issue], list[Issue]]:
    old_rows = [row for row in history if row.period != current_period]
    old_triggered = [row for row in old_rows if row.triggered]
    periods = sorted({row.period for row in old_rows}, key=_period_key)
    previous_period = periods[-1] if periods else ""

    by_id: dict[str, list[Issue]] = defaultdict(list)
    for row in old_triggered:
        by_id[row.issue_id].append(row)
    for values in by_id.values():
        values.sort(key=lambda item: _period_key(item.period))

    current_ids = {item.issue_id for item in current}
    classified: list[Issue] = []
    for item in current:
        prior = by_id.get(item.issue_id, [])
        if not prior:
            item.status = "新增"
            item.first_seen_period = current_period
            item.previous_seen_period = ""
            item.consecutive_count = 1
        else:
            last = prior[-1]
            item.first_seen_period = prior[0].first_seen_period or prior[0].period
            item.previous_seen_period = last.period
            if previous_period and last.period == previous_period:
                item.status = "连续出现"
                item.consecutive_count = max(1, last.consecutive_count) + 1
            else:
                item.status = "再次出现"
                item.consecutive_count = 1
            item.institution_feedback = last.institution_feedback
            item.auditor_opinion = last.auditor_opinion
        classified.append(item)

    resolved: list[Issue] = []
    if previous_period:
        latest_by_id: dict[str, Issue] = {}
        for row in old_triggered:
            if row.period == previous_period:
                latest_by_id[row.issue_id] = row
        for issue_id, old in latest_by_id.items():
            in_scope = audited_org_codes is None or old.org_code in audited_org_codes
            if in_scope and issue_id not in current_ids:
                resolved.append(
                    old.clone(
                        period=current_period,
                        batch_id=batch_id or (current[0].batch_id if current else ""),
                        audit_time=audit_time or (current[0].audit_time if current else ""),
                        triggered=False,
                        status="已整改",
                        previous_seen_period=old.period,
                        consecutive_count=0,
                        target_value="",
                        comparison_value="",
                        reference_value="",
                        difference_value="",
                        formula_result="",
                    )
                )
    return classified, resolved


def merge_history(existing: Iterable[Issue], additions: Iterable[Issue]) -> list[Issue]:
    """Carry all manually written notes for the same readable rule number."""
    grouped: dict[str, list[Issue]] = defaultdict(list)
    for item in existing:
        grouped[item.issue_id].append(item)

    def combine(rows: Iterable[Issue], field: str) -> str:
        values: list[str] = []
        for row in rows:
            value = str(getattr(row, field, "") or "").strip()
            if value and value not in values:
                values.append(value)
        return "\n".join(values)

    for item in additions:
        previous = grouped.get(item.issue_id, [])
        if previous:
            item.institution_feedback = combine(previous, "institution_feedback")
            item.auditor_opinion = combine(previous, "auditor_opinion")
    merged: dict[str, Issue] = {item.issue_id: item for item in existing}
    for item in additions:
        merged[item.issue_id] = item
    return sorted(
        merged.values(),
        key=lambda item: (Path(item.source_file).name, item.sheet_name, item.formula_cell, item.issue_id),
    )
