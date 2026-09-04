from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

from uos_audit.name_config import FeatureMapping
from uos_audit.history import read_history_xlsx, write_history_xlsx
from uos_audit.openpyxl_workbook import (
    add_external_sheets,
    copy_formula_ranges,
    extract_issues,
    named_ranges,
    read_template,
)


def _mapping(name: str) -> FeatureMapping:
    return FeatureMapping("公式校验复制", "修改_公式校验复制", (name,), False, "", "")


def _template(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "报表"
    sheet["A1"] = "字段"
    sheet["B2"] = '=IF(A2>0,"错误|余额|余额不应大于0|"&A2&"||"&A2,"")'
    sheet["B2"].comment = Comment("余额校验", "test")
    book.defined_names.add(DefinedName("校验区域_001", attr_text="'报表'!$B$2"))
    book.defined_names.add(DefinedName("_xlfn.SUMIFS", attr_text="SUMIFS(A:A,B:B,1)"))
    book.save(path)
    book.close()


def _source(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "报表"
    sheet["A2"] = 4
    book.save(path)
    book.close()


def test_named_range_formula_copy_and_cached_extraction(tmp_path: Path) -> None:
    template_path, audit_path = tmp_path / "模板.xlsx", tmp_path / "审核版.xlsx"
    _template(template_path)
    _source(audit_path)
    mapping = _mapping("校验区域")
    definition = read_template(template_path, formula_mappings=(mapping,))
    assert [(item.sheet_name, item.address) for item in definition.copy_ranges] == [("报表", "B2")]
    copy_formula_ranges(template_path, audit_path, definition)
    formulas = load_workbook(audit_path, data_only=False)
    assert formulas["报表"]["B2"].value.startswith("=IF(")
    formulas.close()

    # A LibreOffice-saved cache is represented by a value-only copy in this
    # unit test; extraction itself must not depend on COM or a live office app.
    cached = load_workbook(audit_path, data_only=False)
    cached["报表"]["B2"] = "错误|余额|余额不应大于0|4||4"
    cached.save(audit_path)
    cached.close()
    issues = extract_issues(audit_path, definition.rules, period="2026-08", batch_id="test")
    assert len(issues) == 1
    assert issues[0].check_field == "余额"
    assert issues[0].target_value == "4"
    assert issues[0].rule_id == issues[0].issue_id
    assert issues[0].rule_id == "审核版｜报表｜B2｜余额"


def test_legacy_formula_comment_rule_id_matches_windows_discovery(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    _template(template_path)
    book = load_workbook(template_path)
    book["报表"]["B2"].comment = Comment("规则编号：余额规则-001\n余额校验", "test")
    book.save(template_path)
    book.close()

    definition = read_template(template_path, formula_mappings=(_mapping("校验区域"),))

    assert definition.rules[0].rule_id == "余额规则-001"
    assert definition.rules[0].message == "余额校验"


def test_extract_issues_accepts_live_com_value_cache_without_reopening_xlsx(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    _template(template_path)
    definition = read_template(template_path, formula_mappings=(_mapping("校验区域"),))

    issues = extract_issues(
        tmp_path / "无需存在的审核版.xlsx", definition.rules, period="2026-08", batch_id="test",
        cached_values={
            ("报表", "B2"): "错误|余额|余额不应大于0|4||4",
            ("报表", "A2"): 4,
        },
    )

    assert len(issues) == 1
    assert issues[0].target_value == "4"


def test_named_range_ignores_non_cell_formula_names(tmp_path: Path) -> None:
    path = tmp_path / "names.xlsx"
    _template(path)
    book = load_workbook(path)
    assert [(item.sheet_name, item.address) for item in named_ranges(book, (_mapping("校验区域"),))] == [("报表", "B2")]
    assert named_ranges(book, (_mapping("_xlfn"),)) == []
    book.close()


def test_named_ranges_accept_worksheet_scoped_name() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "报表"
    sheet.defined_names.add(DefinedName("任意行汇总区域", attr_text="'报表'!$A$3:$B$9"))

    assert [(item.sheet_name, item.address) for item in named_ranges(
        book, (_mapping("任意行汇总区域"),)
    )] == [("报表", "A3:B9")]
    book.close()


def test_external_sheet_copy_preserves_hidden_sheet(tmp_path: Path) -> None:
    audit_path, external_path = tmp_path / "audit.xlsx", tmp_path / "external.xlsx"
    _source(audit_path)
    external = Workbook()
    sheet = external.active
    sheet.title = "集中系统数据"
    sheet.sheet_state = "hidden"
    sheet["A1"] = "辅助数据"
    external.create_sheet("说明")
    external.save(external_path)
    external.close()
    add_external_sheets(audit_path, external_path, ("集中系统数据",))
    book = load_workbook(audit_path)
    assert book["集中系统数据"]["A1"].value == "辅助数据"
    assert book["集中系统数据"].sheet_state == "hidden"
    book.close()


def test_history_keeps_manual_notes(tmp_path: Path) -> None:
    path = tmp_path / "历史审核说明.xlsx"
    from uos_audit.models import Issue
    record = Issue("来源｜报表｜B2｜余额", "2026-08", "1", "", True, "新增", "", "", 1, "", "", "", "报表", "来源｜报表｜B2｜余额", "错误", "B2", "B2", "4", "错误", "说明", "来源.xlsx", "", check_field="余额", institution_feedback="已核实", auditor_opinion="通过")
    write_history_xlsx(path, [record])
    restored = read_history_xlsx(path)
    assert restored[0].institution_feedback == "已核实"
    assert restored[0].auditor_opinion == "通过"
