"""Manual WPS COM probe for rendered conditional-format colours.

Run with the Windows runtime Python. It creates a temporary xlsx whose A1 is
yellow when its value is greater than zero, then prints base and displayed COM
colours. This stays outside normal unit tests because it starts WPS.
"""
from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill


def main() -> None:
    path = Path(tempfile.mkdtemp(prefix="base_audit_wps_probe_")) / "条件格式探测.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "条件格式"
    sheet["A1"] = 1
    sheet.conditional_formatting.add(
        "A1", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFFF00"))
    )
    book.save(path)
    book.close()

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = workbook = None
    try:
        app = win32com.client.DispatchEx("ket.Application")
        app.Visible = False
        app.DisplayAlerts = False
        workbook = app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True)
        sheet = workbook.Worksheets(1)
        sheet.Activate()
        try:
            app.CalculateFullRebuild()
        except Exception:
            app.Calculate()
        cell = sheet.Cells(1, 1)
        print("WPS version:", getattr(app, "Version", ""))
        print("Rule count:", sheet.UsedRange.FormatConditions.Count)
        print("Base Interior.Color:", int(cell.Interior.Color))
        try:
            print("DisplayFormat.Interior.Color:", int(cell.DisplayFormat.Interior.Color))
        except Exception as exc:
            print("DisplayFormat error:", repr(exc))
        # Confirm the Windows WPS fallback reads the same rule from OOXML.
        sys_path = str(Path(__file__).resolve().parents[1] / "src")
        import sys
        sys.path.insert(0, sys_path)
        from uos_audit.calculation import WpsComCalculator
        from uos_audit.models import CopyRange
        issues = WpsComCalculator()._extract_rule_based_conditional_formats(
            workbook_path=path, ranges=[CopyRange("条件格式", "A1")], structure_ranges=[],
            period="2026-01-01", batch_id="probe", source_file=path,
        )
        print("OOXML fallback issues:", len(issues), issues[0].detail if issues else "")
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if app is not None:
            app.Quit()
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    main()
