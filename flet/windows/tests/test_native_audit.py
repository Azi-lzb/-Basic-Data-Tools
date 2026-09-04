from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

from uos_audit.audit import run_native_audit
from uos_audit.conditional_format import extract_simple_conditional_format_issues
from uos_audit.models import CopyRange
from uos_audit.name_config import CONDITIONAL_FORMAT_EXTRACT_FUNCTION, FeatureMapping


def test_conditional_only_module_reads_source_without_audit_copy(tmp_path: Path, monkeypatch) -> None:
    """The source already owns its conditional formatting; no formula stage is needed."""
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = "条件格式说明"
    sheet["A1"] = "指标"
    template.defined_names.add(DefinedName("条件格式区域", attr_text="'报表'!$B$2"))
    template.defined_names.add(DefinedName("表结构区域", attr_text="'报表'!$A$1"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source_path = source_dir / "机构A.xlsx"
    source = Workbook()
    source.active.title = "报表"
    source.save(source_path)
    source.close()

    observed: list[Path] = []
    monkeypatch.setattr(
        "uos_audit.audit.extract_simple_conditional_format_issues",
        lambda **kwargs: (observed.append(Path(kwargs["workbook_path"])) or [], 0),
    )
    mapping = FeatureMapping(
        "条件格式结果提取", CONDITIONAL_FORMAT_EXTRACT_FUNCTION, ("条件格式区域",),
        False, "", "", "审核结果集", "源数据目录",
    )
    result = run_native_audit(
        input_dir=source_dir, template_path=template_path, output_dir=tmp_path / "执行结果",
        conditional_mappings=[mapping], include_external=False, include_formula_copy=False,
        include_formula_extraction=False, write_summary=False,
    )

    assert observed == [source_path], [item.error for item in result.files]
    assert result.summary_path is None
    assert not (tmp_path / "执行结果" / "机构A_审核版.xlsx").exists()


def test_simple_true_false_conditional_rule_is_extracted_without_libreoffice(tmp_path: Path) -> None:
    path = tmp_path / "条件格式.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "报表"; sheet["A2"] = 1
    sheet.conditional_formatting.add("A2", FormulaRule(formula=["A2>0"]))
    book.save(path); book.close()
    issues, unsupported = extract_simple_conditional_format_issues(
        workbook_path=path, ranges=[CopyRange("报表", "A2")],
        structure_ranges=[], period="2026-08", batch_id="test", source_file=path,
    )
    assert len(issues) == 1
    assert unsupported == 0
