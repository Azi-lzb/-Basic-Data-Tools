from pathlib import Path

from openpyxl import Workbook

from uos_audit.calculation import (
    ExcelComCalculator,
    inspect_formula_capability,
    scan_formula_errors,
)
from uos_audit.models import CopyRange


class _FakeFormatCondition:
    Formula1 = "=AND($C8<>\"\",$C8>0)"
    Formula2 = ""


class _FakeCellValueFormatCondition:
    Type = 1
    Operator = 5
    Formula1 = "=0"
    Formula2 = ""


def test_rendered_conditional_format_uses_destination_relative_rule_text() -> None:
    text = ExcelComCalculator._conditional_rule_text(
        _FakeFormatCondition(), origin="C8", destination="C10"
    )
    assert text == '条件格式规则：AND(C10<>"",C10>0)'


def test_rendered_cell_value_condition_includes_triggered_cell_and_operator() -> None:
    text = ExcelComCalculator._conditional_rule_text(
        _FakeCellValueFormatCondition(), origin="D2", destination="D26"
    )
    assert text == "条件格式规则：D26>0"


def _book(path: Path, formula: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, formula
    book.save(path)
    book.close()


def test_capability_scan_is_limited_to_copied_formula_ranges(tmp_path: Path) -> None:
    """Unrelated helper formulas must not block the current audit step."""
    path = tmp_path / "审核版.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, "=SUM(A1:A2)"
    sheet["Z1"] = '=INDIRECT("A1")'
    book.save(path)
    book.close()

    full = inspect_formula_capability(path)
    scoped = inspect_formula_capability(path, formula_ranges=[CopyRange("数据", "A3")])

    assert not full.python_supported
    assert scoped.python_supported
    assert scoped.formula_count == 1


def test_capability_recognises_supported_coordinate_and_info_functions(tmp_path: Path) -> None:
    path = tmp_path / "坐标函数.xlsx"
    _book(path, '=IF(ISNUMBER(A1),ADDRESS(ROW(),COLUMN()),"")')

    capability = inspect_formula_capability(path)

    assert capability.python_supported
    assert set(capability.functions) == {"ADDRESS", "COLUMN", "IF", "ISNUMBER", "ROW"}


def test_capability_reports_error_formula_locations(tmp_path: Path) -> None:
    path = tmp_path / "错误公式.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["B2"] = "=SUM(#REF!)"
    sheet["C3"] = "=#NAME?"
    book.save(path)
    book.close()

    capability = inspect_formula_capability(path)

    assert capability.formula_errors == (("数据", "B2", "#REF!"), ("数据", "C3", "#NAME?"))


def test_formula_error_scan_is_limited_to_validation_range(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"] = "#REF!"
    sheet["Z1"] = "#VALUE!"
    book.save(path)
    book.close()

    assert scan_formula_errors(path, formula_ranges=[CopyRange("数据", "A1")]) == (("数据", "A1", "#REF!"),)
