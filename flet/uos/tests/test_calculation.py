from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from uos_audit.calculation import (
    ENGINE_COMBINED,
    ENGINE_OPTIONS,
    ENGINE_IRONCALC,
    ENGINE_PYTHON,
    FormulaEngineError,
    get_calculator,
    inspect_formula_capability,
    scan_formula_errors,
)
from uos_audit.models import CopyRange


def _book(path: Path, formula: str) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, formula
    book.save(path)
    book.close()


def test_python_formula_engine_writes_cache(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    _book(path, "=SUM(A1:A2)")
    result = get_calculator(ENGINE_PYTHON).recalculate(path)
    assert result.engine_name == ENGINE_PYTHON
    book = load_workbook(path, data_only=True)
    assert book["数据"]["A3"].value == 5
    book.close()


def test_formulas_attempts_indirect_despite_compatibility_warning(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    _book(path, '=INDIRECT("A1")')
    capability = inspect_formula_capability(path)
    assert not capability.python_supported
    assert "INDIRECT" in capability.functions
    result = get_calculator(ENGINE_PYTHON).recalculate(path)
    assert result.workbook_path == path


def test_legacy_combined_setting_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    _book(path, "=SUM(A1:A2)")
    assert ENGINE_COMBINED not in ENGINE_OPTIONS
    try:
        get_calculator(ENGINE_COMBINED)
    except FormulaEngineError:
        pass
    else:
        raise AssertionError("已停用的组合引擎不应自动兜底")


def test_ironcalc_engine_writes_cache(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    _book(path, "=SUM(A1:A2)")
    result = get_calculator(ENGINE_IRONCALC).recalculate(path)
    assert result.engine_name == ENGINE_IRONCALC
    book = load_workbook(path, data_only=True)
    assert book["数据"]["A3"].value == 5
    book.close()


def test_capability_scan_is_limited_to_copied_formula_ranges(tmp_path: Path) -> None:
    """Unrelated helper formulas must not block the current audit step."""
    path = tmp_path / "审核版.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, "=SUM(A1:A2)"
    sheet["Z1"] = '=INDIRECT("A1")'
    book.save(path)
    book.close()

    full = inspect_formula_capability(path)
    scoped = inspect_formula_capability(path, formula_ranges=[CopyRange("数据", "A3")])

    assert not full.python_supported
    assert scoped.python_supported
    assert scoped.formula_count == 1


def test_capability_recognises_supported_coordinate_and_info_functions(tmp_path: Path) -> None:
    path = tmp_path / "坐标函数.xlsx"
    _book(path, '=IF(ISNUMBER(A1),ADDRESS(ROW(),COLUMN()),"")')

    capability = inspect_formula_capability(path)

    assert capability.python_supported
    assert set(capability.functions) == {"ADDRESS", "COLUMN", "IF", "ISNUMBER", "ROW"}


def test_capability_reports_error_formula_locations(tmp_path: Path) -> None:
    path = tmp_path / "错误公式.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["B2"] = "=SUM(#REF!)"
    sheet["C3"] = "=#NAME?"
    book.save(path)
    book.close()

    capability = inspect_formula_capability(path)

    assert capability.formula_errors == (("数据", "B2", "#REF!"), ("数据", "C3", "#NAME?"))


def test_python_formula_engine_preserves_base_style_sheet_and_named_range(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"], sheet["A2"], sheet["B1"] = 2, 3, "标题"
    sheet["A3"] = "=SUM(A1:A2)"
    sheet["A3"].font = Font(bold=True, color="FFFFFF")
    sheet["A3"].fill = PatternFill("solid", fgColor="1F4E78")
    external = book.create_sheet("参照表")
    external["A1"] = "保留"
    book.defined_names.add(DefinedName("校验区域", attr_text="'数据'!$A$3"))
    book.save(path)
    book.close()

    get_calculator(ENGINE_PYTHON).recalculate(path)

    calculated = load_workbook(path, data_only=True)
    assert calculated["数据"]["A3"].value == 5
    assert "参照表" in calculated.sheetnames
    assert calculated["参照表"]["A1"].value == "保留"
    calculated.close()
    formulas = load_workbook(path, data_only=False)
    assert formulas["数据"]["A3"].value == "=SUM(A1:A2)"
    assert formulas["数据"]["A3"].font.bold
    assert formulas["数据"]["A3"].fill.fgColor.rgb.endswith("1F4E78")
    assert "校验区域" in {item.name for item in formulas.defined_names.values()}
    formulas.close()


def test_formula_error_scan_is_limited_to_validation_range(tmp_path: Path) -> None:
    path = tmp_path / "审核版.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet["A1"] = "#REF!"
    sheet["Z1"] = "#VALUE!"
    book.save(path)
    book.close()

    assert scan_formula_errors(path, formula_ranges=[CopyRange("数据", "A1")]) == (("数据", "A1", "#REF!"),)
