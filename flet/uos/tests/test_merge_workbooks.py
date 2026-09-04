from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from uos_audit.merge_workbooks import merge_by_first_filename_part, create_combined_template, combine_by_plan

def _book(path, sheet, value):
    b=Workbook(); b.active.title=sheet; b.active["A1"]=value; b.save(path); b.close()

def test_group_and_template_combine(tmp_path: Path):
    src=tmp_path/"src"; src.mkdir(); _book(src/"机构A_表一_B00_2026-08-01.xlsx","表一",1); _book(src/"机构A_表二_B00_2026-08-01.xlsx","表二",2)
    folder,records=merge_by_first_filename_part(input_dir=src,output_dir=tmp_path/"out",output_name="选定联合")
    assert folder.name.startswith("选定联合_")
    result=load_workbook(folder/records[0][1]); assert result.sheetnames==["表一","表二"]; result.close()
    base=src/"base.xlsx"; other=src/"other.xlsx"; _book(base,"底稿",1); _book(other,"新增",2)
    workbook = load_workbook(base)
    workbook.defined_names.add(DefinedName("表结构区域_001", attr_text="'底稿'!$A$1"))
    workbook.save(base)
    workbook.close()
    workbook = load_workbook(other)
    workbook.defined_names.add(DefinedName("表结构区域_001", attr_text="'新增'!$A$1"))
    workbook.save(other)
    workbook.close()
    output,copied,skipped=create_combined_template(base_template=base,source_templates=[other])
    joined=load_workbook(output); assert "新增" in joined.sheetnames
    names = {item.name for item in joined.defined_names.values()}
    assert "表结构区域_001_底稿" in names
    assert "表结构区域_001_新增" in names
    joined.close(); assert copied and not skipped
    combined, rows=combine_by_plan(input_dir=src,output_dir=tmp_path/"comb",plan={"mode":"name","groups":[{"name":"机构A组合","keywords":["机构A"]}]},output_name="选定组合")
    assert combined.name.startswith("选定组合_")
    assert rows and (combined/rows[0][1]).is_file()
