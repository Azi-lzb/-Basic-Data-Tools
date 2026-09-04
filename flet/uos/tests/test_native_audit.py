from pathlib import Path
from contextlib import nullcontext

import pytest

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from uos_audit.audit import run_native_audit
from uos_audit.name_config import CONDITIONAL_FORMAT_EXTRACT_FUNCTION, FeatureMapping, initialize_config
from uos_audit.workflow import run_flow


def test_native_audit_uses_python_engine_for_simple_template(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = '=IF(A2>0,"错误|余额|余额不应大于0|"&A2&"||"&A2,"")'
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source = Workbook()
    source.active.title = "报表"
    source.active["A2"] = 4
    source.save(source_dir / "机构A.xlsx")
    source.close()

    history_path = tmp_path / "历史审核说明.xlsx"
    initialize_config(history_path)
    result = run_native_audit(
        input_dir=source_dir, template_path=template_path, output_dir=tmp_path / "执行结果",
        config_path=history_path, history_path=history_path, calculation_engine="formulas",
    )

    assert result.files[0].error == ""
    assert result.calculation_engine == "formulas"
    assert len(result.current_issues) == 1
    assert result.current_issues[0].check_field == "余额"


def test_default_flow_runs_structure_copy_extract_and_history(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["A1"] = "固定表头"
    sheet["B2"] = '=IF(A2>0,"错误|余额|余额不应大于0|"&A2&"||"&A2,"")'
    template.defined_names.add(DefinedName("表结构区域", attr_text="'报表'!$A$1"))
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    # 条件格式步骤依赖 LibreOffice UNO；此纯公式流程只验证其余链路。
    template.defined_names.add(DefinedName("条件格式区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source = Workbook()
    source.active.title = "报表"
    source.active["A1"] = "固定表头"
    source.active["A2"] = 4
    source.save(source_dir / "机构A.xlsx")
    source.close()

    history_path = tmp_path / "历史审核说明.xlsx"
    initialize_config(history_path)
    outcome = run_flow(
        flow_name="汇总核查表校验", template_path=template_path, input_dir=source_dir,
        output_dir=tmp_path / "执行结果", history_path=history_path, config_path=history_path,
        calculation_engine="formulas", write_flow_logs=True,
    )
    assert Path(outcome["output"]).is_file()
    assert Path(outcome["log"]).is_file()
    assert history_path.is_file()


def test_conditional_only_module_reads_source_without_audit_copy(tmp_path: Path, monkeypatch) -> None:
    """The source already owns its conditional formatting; no formula stage is needed."""
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = "条件格式说明"
    sheet["A1"] = "指标"
    template.defined_names.add(DefinedName("条件格式区域", attr_text="'报表'!$B$2"))
    template.defined_names.add(DefinedName("表结构区域", attr_text="'报表'!$A$1"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source_path = source_dir / "机构A.xlsx"
    source = Workbook()
    source.active.title = "报表"
    source.save(source_path)
    source.close()

    observed: list[Path] = []
    monkeypatch.setattr("uos_audit.audit.libreoffice_conditional_batch", lambda: nullcontext())
    monkeypatch.setattr(
        "uos_audit.audit.extract_conditional_format_issues",
        lambda **kwargs: observed.append(Path(kwargs["workbook_path"])) or [],
    )
    mapping = FeatureMapping(
        "条件格式结果提取", CONDITIONAL_FORMAT_EXTRACT_FUNCTION, ("条件格式区域",),
        False, "", "", "审核结果集", "源数据目录",
    )
    result = run_native_audit(
        input_dir=source_dir, template_path=template_path, output_dir=tmp_path / "执行结果",
        conditional_mappings=[mapping], include_external=False, include_formula_copy=False,
        include_formula_extraction=False, write_summary=False,
    )

    assert observed == [source_path], [item.error for item in result.files]
    assert result.summary_path is None
    assert not (tmp_path / "执行结果" / "机构A_审核版.xlsx").exists()
def test_external_sheet_is_added_before_formula_calculation(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = '=IF(参照表!A1>0,"错误|外部值|外部数据异常|||"," ")'
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source = Workbook()
    source.active.title = "报表"
    source.save(source_dir / "机构A.xlsx")
    source.close()

    external_path = tmp_path / "外部.xlsx"
    external = Workbook()
    external.active.title = "参照表"
    external.active["A1"] = 1
    external.save(external_path)
    external.close()

    result = run_native_audit(
        input_dir=source_dir, template_path=template_path, output_dir=tmp_path / "执行结果",
        external_path=external_path, calculation_engine="formulas",
    )

    assert result.files[0].error == ""
    assert len(result.current_issues) == 1
    audit = __import__("openpyxl").load_workbook(result.files[0].audit_path, data_only=False)
    assert "参照表" in audit.sheetnames
    audit.close()


def test_formula_copy_can_be_used_as_temporary_step_without_retaining_audit_copy(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = '=IF(A2>0,"错误|余额|说明|"&A2&"|||","")'
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source = Workbook()
    source.active.title = "报表"
    source.active["A2"] = 1
    source.save(source_dir / "机构A.xlsx")
    source.close()

    result = run_native_audit(
        input_dir=source_dir, template_path=template_path,
        output_dir=tmp_path / "执行结果", calculation_engine="formulas",
        keep_audit_copies=False,
    )

    assert len(result.current_issues) == 1
    assert result.files[0].audit_path is None
    assert not (tmp_path / "执行结果" / "机构A_审核版.xlsx").exists()


def test_stop_on_file_failure_honours_flow_stop_policy(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = "=1"
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    _book = Workbook()
    _book.active.title = "其他报表"
    _book.save(source_dir / "机构A.xlsx")
    _book.close()

    with pytest.raises(RuntimeError, match="处理“机构A.xlsx”失败"):
        run_native_audit(
            input_dir=source_dir, template_path=template_path,
            output_dir=tmp_path / "执行结果", calculation_engine="formulas",
            stop_on_file_failure=True,
        )


def test_copy_only_engine_rejects_formula_result_extraction(tmp_path: Path) -> None:
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["B2"] = "=1"
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()
    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source = Workbook()
    source.active.title = "报表"
    source.save(source_dir / "机构A.xlsx")
    source.close()

    with pytest.raises(RuntimeError, match="不支持的计算引擎"):
        run_native_audit(
            input_dir=source_dir, template_path=template_path,
            output_dir=tmp_path / "执行结果", calculation_engine="仅复制不计算",
        )


def test_default_flow_end_to_end_with_structure_external_formula_history_and_output(tmp_path: Path) -> None:
    """The main Flet audit action must execute the whole configured pipeline."""
    template_path = tmp_path / "模板.xlsx"
    template = Workbook()
    sheet = template.active
    sheet.title = "报表"
    sheet["A1"] = "固定表头"
    sheet["B2"] = '=IF(参照表!A1=1,"错误|外部校验|外部值异常|1|1|0","")'
    template.defined_names.add(DefinedName("表结构区域", attr_text="'报表'!$A$1"))
    template.defined_names.add(DefinedName("校验区域", attr_text="'报表'!$B$2"))
    template.save(template_path)
    template.close()

    source_dir = tmp_path / "源数据"
    source_dir.mkdir()
    source = Workbook()
    source.active.title = "报表"
    source.active["A1"] = "固定表头"
    source.save(source_dir / "机构A.xlsx")
    source.close()
    external_path = tmp_path / "外部.xlsx"
    external = Workbook()
    external.active.title = "参照表"
    external.active["A1"] = 1
    external.save(external_path)
    external.close()

    history_path = tmp_path / "历史审核说明.xlsx"
    initialize_config(history_path)
    outcome = run_flow(
        flow_name="汇总核查表校验", template_path=template_path,
        input_dir=source_dir, output_dir=tmp_path / "执行结果",
        history_path=history_path, config_path=history_path,
        external_path=external_path, calculation_engine="formulas",
        write_flow_logs=True,
    )

    result_path = Path(outcome["output"])
    log_path = Path(outcome["log"])
    assert result_path.is_file()
    assert log_path.is_file()
    audit_copy = next((tmp_path / "执行结果").glob("机构审核副本_*/*_审核版.xlsx"))
    assert audit_copy.is_file()
    audit_book = __import__("openpyxl").load_workbook(audit_copy, data_only=False)
    assert "参照表" in audit_book.sheetnames
    audit_book.close()
    result_book = __import__("openpyxl").load_workbook(result_path, data_only=True)
    result_rows = list(result_book["本期审核结果"].iter_rows(min_row=2, values_only=True))
    result_book.close()
    assert result_rows[0][3:9] == ("错误", "外部校验", "1", "1", "0", "外部值异常")
    result_book = __import__("openpyxl").load_workbook(result_path, data_only=False)
    result_sheet = result_book["本期审核结果"]
    assert result_sheet["A2"].hyperlink is None
    assert result_sheet["C2"].hyperlink is not None
    assert result_sheet["C2"].hyperlink.location == "'报表'!B2"
    result_book.close()
    assert history_path.is_file()
