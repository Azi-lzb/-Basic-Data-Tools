"""Openpyxl implementation of the fixed-row and used-range summary modules."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

from .discovery import source_workbooks
from .name_config import FIXED_ROW_SUMMARY_FUNCTION, USED_RANGE_SUMMARY_FUNCTION, FeatureMapping
from .openpyxl_workbook import named_ranges, named_single_cells_with_prefix


def _bounds(address: str) -> tuple[int, int, int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(address)
    return min_row, min_col, max_row, max_col


def _header_mapping(feature: FeatureMapping) -> FeatureMapping:
    return FeatureMapping(feature.name, feature.feature_type, ("表头区域",), False, "", "")


def _field_key(value: object) -> str:
    return str(value or "").replace(" ", "").replace("\n", "").casefold()


def _sheet_name(base: str, used: set[str]) -> str:
    raw = "".join("_" if char in r'[]:*?/\\' else char for char in base)[:31] or "汇总"
    candidate, number = raw, 2
    while candidate.casefold() in used:
        suffix = "_{}".format(number)
        candidate = raw[: 31 - len(suffix)] + suffix
        number += 1
    used.add(candidate.casefold())
    return candidate


def _xls_sheet_values(path: Path, sheet_name: str):
    """Read a legacy xls sheet as a zero-based pandas table using xlrd."""
    import pandas as pd

    return pd.read_excel(path, sheet_name=sheet_name, header=None, engine="xlrd")


def _xls_value(table, row: int, column: int):
    if row < 1 or column < 1 or row > len(table.index) or column > len(table.columns):
        return None
    value = table.iat[row - 1, column - 1]
    return None if pd_is_na(value) else value


def pd_is_na(value: object) -> bool:
    """Avoid importing pandas for the normal xlsx hot path."""
    try:
        import pandas as pd
        return bool(pd.isna(value))
    except Exception:
        return False


def _read_header(template_book, data_range, headers) -> tuple[list[object], int, int]:
    matching = [header for header in headers if header.sheet_name == data_range.sheet_name]
    if not matching:
        raise ValueError("工作表“{}”缺少命名区域“表头区域”".format(data_range.sheet_name))
    data_top, data_left, _data_bottom, data_right = _bounds(data_range.address)
    candidates = []
    for header in matching:
        header_top, header_left, header_bottom, header_right = _bounds(header.address)
        left, right = max(data_left, header_left), min(data_right, header_right)
        if left <= right:
            candidates.append((right - left, header, left, right, header_top, header_bottom))
    if not candidates:
        raise ValueError("工作表“{}”的表头区域与汇总区域没有共同列".format(data_range.sheet_name))
    _, header, left, right, top, bottom = max(candidates, key=lambda item: item[0])
    sheet = template_book[data_range.sheet_name]
    # 多行表头按照从上到下的非空值拼接，与 Windows 汇总逻辑一致。
    values = []
    for column in range(left, right + 1):
        pieces = []
        for row in range(top, bottom + 1):
            value = sheet.cell(row, column).value
            if value not in (None, ""):
                text = str(value).strip()
                if text and text not in pieces:
                    pieces.append(text)
        values.append("_".join(pieces) if pieces else "列{}".format(column))
    return values, left, right


def run_region_summaries(
    *, template_path: Path, input_dir: Path, output_dir: Path,
    features: Iterable[FeatureMapping], recursive: bool = True,
    selected_files: Iterable[Path] | None = None,
    on_step: Callable[[str], None] | None = None,
    output_name: str = "区域汇总",
) -> Path:
    """Run configured native row summaries and return the output workbook path."""
    say = on_step or (lambda _text: None)
    features = [item for item in features if item.feature_type in {USED_RANGE_SUMMARY_FUNCTION, FIXED_ROW_SUMMARY_FUNCTION}]
    if not features:
        raise ValueError("没有启用“汇总_任意行汇总”或“汇总_固定行汇总”模块")
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files:
        allowed = {path.resolve() for path in files}
        files = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not files:
        raise FileNotFoundError("没有可汇总的源工作簿")
    template = load_workbook(template_path, read_only=True, data_only=False, keep_links=False)
    try:
        global_fields = named_single_cells_with_prefix(template, "全局单元格区域.")
        local_fields = named_single_cells_with_prefix(template, "单元格区域.")
        groups: dict[str, dict[str, object]] = {}
        for feature in features:
            say("正在汇总：{}".format(feature.name))
            data_ranges, headers = named_ranges(template, (feature,)), named_ranges(template, (_header_mapping(feature),))
            if not data_ranges:
                raise ValueError("汇总功能“{}”未找到命名区域：{}".format(feature.name, "、".join(feature.range_names)))
            if not headers:
                raise ValueError("汇总功能“{}”缺少命名区域“表头区域”".format(feature.name))
            for data_range in data_ranges:
                header_values, left, right = _read_header(template, data_range, headers)
                local = [(name, area) for name, area in local_fields if area.sheet_name == data_range.sheet_name]
                extras = [*global_fields, *local]
                fields = [name for name, _ in extras]
                key = data_range.sheet_name.casefold()
                group = groups.setdefault(key, {
                    "title": data_range.sheet_name, "headers": ["来源文件", "来源工作表", *fields, *header_values], "rows": [],
                })
                top, _left, bottom, _right = _bounds(data_range.address)
                for file_path in files:
                    if file_path.suffix.casefold() == ".xls":
                        try:
                            table = _xls_sheet_values(file_path, data_range.sheet_name)
                        except ValueError:
                            continue
                        last_row = bottom if feature.feature_type == FIXED_ROW_SUMMARY_FUNCTION else max(top, len(table.index))
                        metadata = []
                        for field, area in extras:
                            if area.sheet_name != data_range.sheet_name:
                                raise ValueError("旧版 .xls 汇总暂不支持跨工作表单元格字段“{}”".format(field))
                            row, col = _bounds(area.address)[:2]
                            metadata.append(_xls_value(table, row, col))
                        for row_number in range(top, last_row + 1):
                            row = [_xls_value(table, row_number, column) for column in range(left, right + 1)]
                            if any(value not in (None, "") for value in row):
                                group["rows"].append([file_path.name, data_range.sheet_name, *metadata, *row])
                        continue
                    book = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
                    try:
                        if data_range.sheet_name not in book.sheetnames:
                            continue
                        sheet = book[data_range.sheet_name]
                        last_row = bottom if feature.feature_type == FIXED_ROW_SUMMARY_FUNCTION else max(top, sheet.max_row)
                        values = list(sheet.iter_rows(min_row=top, max_row=last_row, min_col=left, max_col=right, values_only=True))
                        metadata = []
                        for field, area in extras:
                            if area.sheet_name not in book.sheetnames:
                                raise ValueError("源文件“{}”缺少单元格区域字段“{}”所在工作表“{}”".format(file_path.name, field, area.sheet_name))
                            row, col = _bounds(area.address)[:2]
                            metadata.append(book[area.sheet_name].cell(row, col).value)
                        for row in values:
                            if any(value not in (None, "") for value in row):
                                group["rows"].append([file_path.name, data_range.sheet_name, *metadata, *row])
                    finally:
                        book.close()
            say("完成：{}".format(feature.name))
    finally:
        template.close()
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "{}_{}.xlsx".format(output_name, datetime.now().strftime("%Y%m%d_%H%M%S"))
    output, used_names = Workbook(), set()
    try:
        first = True
        for group in groups.values():
            sheet = output.active if first else output.create_sheet()
            first = False
            sheet.title = _sheet_name(str(group["title"]), used_names)
            sheet.append(group["headers"])
            for row in group["rows"]:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(40, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        if first:
            output.active.title = "汇总"
        output.save(path)
    finally:
        output.close()
    return path


def merge_workbook_tables(
    *, input_dir: Path, output_dir: Path, recursive: bool = True,
    output_name: str = "汇总表合并", selected_files: Iterable[Path] | None = None,
    on_step=None,
) -> Path:
    """`汇总_汇总表合并`: only sheets with identical first-row headers merge."""
    say = on_step or (lambda _text: None)
    files = source_workbooks(input_dir, recursive=recursive)
    if selected_files is not None:
        allowed = {path.resolve() for path in files}
        files = [Path(path).resolve() for path in selected_files if Path(path).resolve() in allowed]
    if not files:
        raise FileNotFoundError("没有可合并的源工作簿")
    groups: dict[tuple[str, ...], dict[str, object]] = {}
    for path in files:
        if path.suffix.casefold() == ".xls":
            # xlrd only supports legacy xls and pandas turns each sheet into
            # the same row iterator shape used by the xlsx branch below.
            import pandas as pd

            with pd.ExcelFile(path, engine="xlrd") as legacy:
                for sheet_name in legacy.sheet_names:
                    frame = pd.read_excel(legacy, sheet_name=sheet_name, header=None)
                    values = frame.where(pd.notna(frame), None).values.tolist()
                    if not values:
                        continue
                    header = tuple(values[0])
                    if not any(value not in (None, "") for value in header):
                        continue
                    signature = tuple(str(value or "").strip() for value in header)
                    group = groups.setdefault(signature, {"headers": ["来源文件", "来源工作表", *header], "rows": [], "sheets": []})
                    if sheet_name not in group["sheets"]:
                        group["sheets"].append(sheet_name)
                    for row in values[1:]:
                        if any(value not in (None, "") for value in row):
                            group["rows"].append([path.name, sheet_name, *row])
            continue
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = tuple(next(rows, ()))
                if not header or not any(value not in (None, "") for value in header):
                    continue
                signature = tuple(str(value or "").strip() for value in header)
                group = groups.setdefault(signature, {"headers": ["来源文件", "来源工作表", *header], "rows": [], "sheets": []})
                if sheet.title not in group["sheets"]:
                    group["sheets"].append(sheet.title)
                for row in rows:
                    if any(value not in (None, "") for value in row):
                        group["rows"].append([path.name, sheet.title, *row])
        finally:
            book.close()
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "{}_{}.xlsx".format(output_name, datetime.now().strftime("%Y%m%d_%H%M%S"))
    output, used = Workbook(), set()
    try:
        first = True
        for group in groups.values():
            sheet = output.active if first else output.create_sheet(); first = False
            sheet.title = _sheet_name("_".join(group["sheets"][:2]) or "汇总", used)
            sheet.append(group["headers"])
            for row in group["rows"]: sheet.append(row)
            sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        if first: output.active.title = "汇总"
        output.save(path); say("汇总表合并完成：按表头分为 {} 张表".format(len(groups)))
    finally: output.close()
    return path
