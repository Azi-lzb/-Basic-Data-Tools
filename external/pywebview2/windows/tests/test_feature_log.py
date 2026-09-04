from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from src.base_audit.excel_com import ExcelSession
from src.base_audit.feature_log import FeatureLog, write_feature_log_xlsx
from src.base_audit.name_config import FeatureMapping, NAMED_RANGE_CHECK_FUNCTION


class _Area:
    def __init__(self, sheet: str, address: str):
        self.Worksheet = type("W", (), {"Name": sheet})()
        self.Address = address
        self.Rows = type("R", (), {"Count": 1})()
        self.Columns = type("C", (), {"Count": 1})()


class _Name:
    def __init__(self, full_name: str, *areas):
        self.Name = full_name
        self.RefersToRange = type("R", (), {"Areas": areas})()


class _Names:
    def __init__(self, items):
        self._items = items
        self.Count = len(items)

    def Item(self, index):
        return self._items[index - 1]


class _FakeWorkbook:
    def __init__(self, names):
        self.Names = _Names(names)


class FeatureLogTests(unittest.TestCase):
    def test_write_produces_one_sheet_per_feature(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            log = FeatureLog("汇总核查表校验", root)
            log.add_sheet(
                "检查校验区域",
                ("工作表", "命名区域名", "覆盖区域", "结果"),
                [("贷款明细", "校验区域", "A1:D100", "通过")],
            )
            log.add_sheet(
                "校验结果提取",
                ("报送文件", "工作表", "定位单元格", "级别"),
                [("机构甲.xlsx", "贷款明细", "D5", "错误")],
            )
            path = log.write()
            self.assertIsNotNone(path)
            self.assertTrue(path.name.startswith("汇总核查表校验_运行日志_"))
            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(["检查校验区域", "校验结果提取"], workbook.sheetnames)
            self.assertEqual("贷款明细", workbook["检查校验区域"]["A2"].value)
            self.assertEqual("A1:D100", workbook["检查校验区域"]["C2"].value)
            workbook.close()

    def test_write_returns_none_when_no_sheets(self):
        with TemporaryDirectory() as folder:
            log = FeatureLog("汇总核查表校验", Path(folder))
            self.assertIsNone(log.write())

    def test_survey_named_ranges_lists_sheets_coverage_and_missing(self):
        workbook = _FakeWorkbook(
            [
                _Name("'贷款明细'!校验区域", _Area("贷款明细", "A1:D100")),
                _Name("'抵押明细'!校验区域", _Area("抵押明细", "B3:H50")),
            ]
        )
        excel = ExcelSession()  # 不进入 COM 上下文，仅测试枚举逻辑
        mapping = FeatureMapping(
            "检查校验区域", NAMED_RANGE_CHECK_FUNCTION, ("校验区域",), False, "", ""
        )
        self.assertEqual(
            [
                ("贷款明细", "校验区域", "A1:D100", "通过"),
                ("抵押明细", "校验区域", "B3:H50", "通过"),
            ],
            excel.survey_named_ranges(workbook, mapping),
        )

    def test_survey_named_ranges_marks_missing_name(self):
        workbook = _FakeWorkbook([])
        excel = ExcelSession()
        mapping = FeatureMapping(
            "检查表结构区域", NAMED_RANGE_CHECK_FUNCTION, ("表结构区域",), False, "", ""
        )
        self.assertEqual([("—", "表结构区域", "—", "缺失")], excel.survey_named_ranges(workbook, mapping))

    def test_write_feature_log_xlsx_handles_header_only(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "运行日志.xlsx"
            write_feature_log_xlsx(
                path, [("模板体检", ("类别", "级别", "状态"), [])]
            )
            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual("类别", workbook["模板体检"]["A1"].value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
