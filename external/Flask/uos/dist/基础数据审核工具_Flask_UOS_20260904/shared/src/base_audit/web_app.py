from __future__ import annotations

import os
import threading
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from .discovery import (
    classify_source_files,
    detect_period,
    explanation_files,
    recommend_template,
)
from .history import HISTORY_AUDIT_SHEET, HISTORY_HEADERS
from .name_config import (
    HISTORY_WORKBOOK_NAME,
    initialize_config,
    load_config_editor_data,
    reset_default_configuration,
    save_config_editor_draft,
    validate_config_editor_draft,
)
from .service import AuditService
from .settings import SettingsStore, hide_application_data_directory


def available_engines() -> list[dict[str, str]]:
    """当前环境真正可用的计算引擎选项（启动时按操作系统识别）。

    Windows 使用 Microsoft Excel / WPS 表格 COM；统信 UOS / 麒麟使用
    LibreOffice Calc/UNO。界面只显示当前环境可用的引擎。
    """
    if sys.platform == "win32":
        return [
            {"value": "自动", "label": "自动（推荐）"},
            {"value": "Microsoft Excel", "label": "Microsoft Excel"},
            {"value": "WPS 表格", "label": "WPS 表格"},
        ]
    return [
        {"value": "自动", "label": "自动（推荐）"},
        {"value": "LibreOffice Calc", "label": "LibreOffice Calc"},
    ]


def engine_hint() -> str:
    if sys.platform == "win32":
        return (
            "自动模式优先使用 Microsoft Excel；无法启动时尝试 WPS 表格。"
            "指定引擎后，程序不会再自动切换，适合排查公式兼容性问题。"
        )
    return "统信 UOS / 麒麟环境使用系统安装的 LibreOffice Calc 计算公式和渲染条件格式。"


def valid_engine_values() -> set[str]:
    return {item["value"] for item in available_engines()}


class WebApi:
    def __init__(self, project_root: Path) -> None:
        self.project_root = project_root
        hide_application_data_directory(project_root / "data")
        self.settings_store = SettingsStore(project_root / "data" / "用户设置.json")
        self.settings = self.settings_store.load()
        # 历史审核配置：优先用上次选择的路径，否则用程序根目录默认文件。
        self.history_path = (
            Path(self.settings.history_config)
            if self.settings.history_config
            else project_root / HISTORY_WORKBOOK_NAME
        )
        bundled_templates = project_root / "templates"
        formal_templates = project_root / "2026-07-31" / "模板文件"
        saved_input = self.settings.last_input_dir
        saved_output = self.settings.last_output_dir
        self.state: dict[str, Any] = {
            "busy": False,
            "status": "就绪",
            "log": [],
            "input": saved_input,
            "templateDir": self.settings.last_template_dir or str(
                bundled_templates if bundled_templates.is_dir() else formal_templates
            ),
            "template": "",
            "templateManual": False,
            "external": self.settings.last_external_file,
            "historyConfig": str(self.history_path),
            "output": saved_output,
            "outputAuto": not self.settings.output_pinned,
            "outputPinned": self.settings.output_pinned,
            "recursive": self.settings.recursive_folders,
            "writeFlowLogs": self.settings.write_flow_logs,
            "confirmBeforeRun": self.settings.confirm_before_run,
            "calculationEngine": (
                self.settings.calculation_engine
                if self.settings.calculation_engine in valid_engine_values()
                else "自动"
            ),
            "engines": available_engines(),
            "engineHint": engine_hint(),
            "showCustomFeatures": self.settings.show_custom_features,
            "sourceFiles": [],
            "selectedFiles": [],
            "mixedTemplates": [],
            "explanationFiles": [],
        }
        # This also performs the one-time non-destructive filename migration.
        initialize_config(self.history_path)

    def get_state(self) -> dict[str, Any]:
        return self.state

    def initialize_config(self) -> dict[str, Any]:
        """Migrate workflow settings to JSON and retain Excel history only."""
        path = initialize_config(self.history_path)
        self._log(f"已初始化流程配置：{path}（仅保留历史核查表审核）")
        self.state["status"] = "配置已初始化"
        return self.state

    def reset_config(self) -> dict[str, Any]:
        """Reset workflow display settings without touching Excel history."""
        path = reset_default_configuration(self.history_path)
        self._log(f"已重置自定义流程和主界面显示设置：{path}（历史核查表审核未改动）")
        self.state["status"] = "配置已重置"
        return self.state

    def get_config_editor_data(self) -> dict[str, object]:
        """Configuration-centre payload; deliberately excludes history records."""
        return load_config_editor_data(self.history_path)

    def validate_config_editor_draft(self, draft: dict[str, Any]) -> dict[str, object]:
        """Validate a pending visual-editor change without writing history."""
        errors = validate_config_editor_draft(
            self.history_path, draft
        )
        return {"valid": not errors, "errors": errors}

    def save_config_editor_draft(self, draft: dict[str, Any]) -> dict[str, object]:
        """Persist functions, workflows and display settings; Excel history stays intact."""
        if self.state["busy"]:
            return {"ok": False, "errors": ["任务正在运行，暂不能修改配置"]}
        try:
            path = save_config_editor_draft(
                self.history_path, draft
            )
        except Exception as exc:
            self._log(f"保存设置中心配置失败：{exc}")
            return {"ok": False, "errors": [str(exc)]}
        self._log(f"已保存功能、流程和主界面显示设置：{path}")
        return {"ok": True, "errors": [], "data": load_config_editor_data(path)}

    def open_config(self) -> dict[str, Any]:
        """Backward-compatible alias for :meth:`open_history_explanation`."""
        return self.open_history_explanation()

    def open_history_explanation(self) -> dict[str, Any]:
        """Open the Excel workbook that stores historical audit explanations."""
        import os
        path = self.history_path
        if not path.is_file():
            self.state["status"] = "历史审核配置不存在"
            self._log(f"未找到历史审核配置：{path}")
            return self.state
        try:
            os.startfile(path)
        except Exception as exc:
            self.state["status"] = "无法打开历史审核配置"
            self._log(f"打开历史审核配置失败：{exc}")
            return self.state
        self._log(f"已打开历史审核配置：{path}")
        return self.state

    def get_history_page(self, query: dict[str, Any] | None = None) -> dict[str, object]:
        """Return one read-only, filtered page of historical audit explanations.

        The worksheet is streamed with openpyxl: large history files never
        become a giant WebView payload.  This endpoint deliberately provides
        no editing operation; formal maintenance remains in Excel.
        """
        from openpyxl import load_workbook

        query = query if isinstance(query, dict) else {}

        def number(name: str, default: int, minimum: int, maximum: int) -> int:
            try:
                value = int(query.get(name, default))
            except (TypeError, ValueError):
                value = default
            return max(minimum, min(maximum, value))

        page = number("page", 1, 1, 1000000)
        page_size = number("pageSize", 50, 10, 100)
        keyword = str(query.get("keyword") or "").strip().casefold()
        error_type = str(query.get("errorType") or "").strip()
        opinion = str(query.get("opinion") or "").strip()
        path = self.history_path
        if not path.is_file():
            return {
                "columns": list(HISTORY_HEADERS), "items": [], "total": 0,
                "page": page, "pageSize": page_size, "totalPages": 0,
                "errorTypes": [], "opinions": [], "error": "历史审核配置不存在",
            }
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if HISTORY_AUDIT_SHEET not in workbook.sheetnames:
                raise ValueError(f"历史审核配置中缺少“{HISTORY_AUDIT_SHEET}”工作表")
            rows = workbook[HISTORY_AUDIT_SHEET].iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            missing = [name for name in HISTORY_HEADERS if name not in headers]
            if missing:
                raise ValueError("历史审核配置缺少列：" + "、".join(missing))
            positions = {name: headers.index(name) for name in HISTORY_HEADERS}
            total = 0
            error_types: set[str] = set()
            opinions: set[str] = set()
            items: list[dict[str, str]] = []
            start = (page - 1) * page_size
            end = start + page_size
            for raw in rows:
                item = {
                    name: str(raw[index] if index < len(raw) and raw[index] is not None else "").strip()
                    for name, index in positions.items()
                }
                if item["错误类型"]:
                    error_types.add(item["错误类型"])
                if item["审核意见"]:
                    opinions.add(item["审核意见"])
                haystack = " ".join(item.values()).casefold()
                if keyword and keyword not in haystack:
                    continue
                if error_type and item["错误类型"] != error_type:
                    continue
                if opinion and item["审核意见"] != opinion:
                    continue
                if start <= total < end:
                    items.append(item)
                total += 1
            total_pages = (total + page_size - 1) // page_size
            # A filter can make the requested page out of range.  Reply with
            # the real last page marker; the front end can request it once.
            return {
                "columns": list(HISTORY_HEADERS), "items": items, "total": total,
                "page": page, "pageSize": page_size, "totalPages": total_pages,
                "errorTypes": sorted(error_types), "opinions": sorted(opinions),
            }
        finally:
            workbook.close()

    def open_user_guide(self) -> dict[str, Any]:
        """Open the adjacent Word user guide with the default application."""
        import os
        path = self.project_root / "基础数据审核工具使用说明.docx"
        if not path.is_file():
            self.state["status"] = "使用说明不存在"
            self._log(f"未找到使用说明：{path}")
            return self.state
        try:
            os.startfile(path)
        except Exception as exc:
            self.state["status"] = "无法打开使用说明"
            self._log(f"打开使用说明失败：{exc}")
            return self.state
        self._log(f"已打开使用说明：{path}")
        return self.state

    def open_path(self, path: str) -> dict[str, Any]:
        """Open a local directory or file with the system default application."""
        import os
        target = Path(path)
        if not target.exists():
            self.state["status"] = "路径不存在"
            self._log(f"未找到路径：{path}")
            return self.state
        try:
            os.startfile(target)
        except Exception as exc:
            self.state["status"] = "无法打开路径"
            self._log(f"打开路径失败：{exc}")
            return self.state
        self._log(f"已打开：{path}")
        return self.state

    def _load_custom_features(self) -> list[dict[str, Any]]:
        """Read custom flows configured to appear as main-workbench entries."""
        try:
            editor = load_config_editor_data(self.history_path)
        except Exception as exc:
            self._log(f"读取流程显示配置失败：{exc}")
            return []
        display = editor.get("flowDisplay", {})
        if not isinstance(display, dict):
            return []
        features: list[dict[str, Any]] = []
        seen: set[str] = set()
        # 常用流程已有固定主按钮；这里只显示用户在流程编排中新建并勾选的流程。
        for row in editor.get("customFlows", []):
            feature_name = str(row.get("流程名") or "").strip()
            if not feature_name or feature_name in seen:
                continue
            seen.add(feature_name)
            features.append({
                "id": feature_name,
                "name": feature_name,
                "flow": feature_name,
                "shown": bool(display.get(feature_name)),
                "remark": "",
            })
        return features

    def get_custom_features(self) -> list[dict[str, Any]]:
        """Visible flat custom features for the workbench."""
        if not self.state["showCustomFeatures"]:
            return []
        return [
            {"id": item["id"], "name": item["name"], "flow": item["flow"], "remark": item["remark"]}
            for item in self._load_custom_features()
            if item["shown"]
        ]

    def run_custom(self, feature_name: str) -> dict[str, Any]:
        """Run one user-defined flow displayed on the workbench."""
        if self.state["busy"]:
            return self.state
        feature = next(
            (item for item in self._load_custom_features() if item["id"] == feature_name),
            None,
        )
        if feature is None:
            self._log(f"未找到显示流程：{feature_name}")
            return self.state
        if not feature["flow"]:
            self._log(f"显示流程“{feature['name']}”缺少流程名称。")
            self.state["status"] = "缺少流程名称"
            return self.state
        # 主界面自定义流程走宽松校验（strict=False）：组合前置缺失时只警告不报错。
        self.start_flow(feature["flow"], strict=False)
        return self.state

    def start_flow(self, flow_name: str, strict: bool = True) -> bool:
        """Start one configured execution flow by its displayed name."""
        name = str(flow_name).strip()
        if not name:
            self._log("执行流程名称不能为空")
            return False
        return self.start("flow:" + name, strict=strict)

    def start_with_state(
        self,
        action: str,
        values: dict[str, Any] | None = None,
        selected_files: list[str] | None = None,
        strict: bool = True,
    ) -> bool:
        """Synchronize form values and launch an action in one WebView call.

        单请求边界：前端把表单同步、勾选文件和启动合并成一次调用，
        避免桥接层连续多次请求的时序问题。
        """
        if values:
            self.update(values)
        if selected_files is not None:
            self.set_selected_files(selected_files)
        return self.start(action, strict=strict)

    def start_flow_with_state(
        self,
        flow_name: str,
        values: dict[str, Any] | None = None,
        selected_files: list[str] | None = None,
        strict: bool = True,
    ) -> bool:
        """Single-call variant of :meth:`start_flow` for legacy WebView."""
        name = str(flow_name).strip()
        if not name:
            self._log("执行流程名称不能为空")
            return False
        return self.start_with_state("flow:" + name, values, selected_files, strict=strict)

    def _template_merge_worker(
        self, base_template: Path, source_templates: list[Path]
    ) -> None:
        started = time.monotonic()
        try:
            service = AuditService(
                config_path=self.history_path,
                engine_preference=str(self.state["calculationEngine"]),
            )
            result = service.merge_template_files(
                base_template=base_template,
                source_templates=source_templates,
                on_step=self._log_detail,
            )
            self.state["status"] = result.summary_text().splitlines()[0]
            self._log(result.summary_text())
        except Exception as exc:
            self.state["status"] = "制作联合模板失败"
            self._log("制作联合模板失败：" + str(exc))
        finally:
            elapsed = time.monotonic() - started
            self._log(f"任务结束，本次耗时：{elapsed:.1f} 秒")
            self.state["busy"] = False

    def update(self, values: dict[str, Any]) -> dict[str, Any]:
        for key in ("input", "templateDir", "template", "external", "output"):
            if key in values:
                new_value = str(values[key]).strip()
                previous = self.state.get(key, "")
                if key == "output" and new_value != self.state.get("output", ""):
                    self.state["outputAuto"] = False
                    self.state["outputPinned"] = True
                self.state[key] = new_value
                if key == "template" and new_value and new_value != previous:
                    self.state["templateManual"] = True
                elif key == "templateDir" and new_value != previous:
                    self.state["templateManual"] = False
        if "outputPinned" in values:
            value = values["outputPinned"]
            pinned = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self.state["outputPinned"] = pinned
            self.state["outputAuto"] = not pinned
            if not pinned and self.state["input"]:
                self.state["output"] = str(Path(self.state["input"]) / "执行结果")
        if "recursive" in values:
            value = values["recursive"]
            recursive = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            changed = recursive != self.state["recursive"]
            self.state["recursive"] = recursive
            if changed:
                # 复选框直接决定待审核清单的扫描范围；不触发模板自动改写。
                self._recognize(allow_template_auto=False)
        if "writeFlowLogs" in values:
            value = values["writeFlowLogs"]
            self.state["writeFlowLogs"] = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self._save_settings()
        if "confirmBeforeRun" in values:
            value = values["confirmBeforeRun"]
            self.state["confirmBeforeRun"] = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self._save_settings()
        if "calculationEngine" in values:
            candidate = str(values["calculationEngine"]).strip()
            if candidate not in valid_engine_values():
                self._log("计算引擎只能选择：" + "、".join(sorted(valid_engine_values())))
            else:
                self.state["calculationEngine"] = candidate
                self._save_settings()
        if "showCustomFeatures" in values:
            value = values["showCustomFeatures"]
            self.state["showCustomFeatures"] = (
                value if isinstance(value, bool)
                else str(value).strip().casefold() in {"1", "true", "yes", "y", "是"}
            )
            self._save_settings()
        return self.state

    def recognize(self) -> dict[str, Any]:
        self._recognize(force_template=True)
        return self.state

    def set_selected_files(self, paths: list[str]) -> dict[str, Any]:
        allowed = {item["path"] for item in self.state.get("sourceFiles", [])}
        self.state["selectedFiles"] = [path for path in paths if path in allowed]
        return self.state

    def start(self, action: str, strict: bool = True) -> bool:
        if self.state["busy"]:
            return False
        self.state["busy"] = True
        self.state["status"] = "正在处理，请勿关闭窗口……"
        # Python 3.7 is used by the Win7 package; str.removeprefix arrived in
        # Python 3.9, so use slicing here instead.
        display = action[len("flow:"):] if action.startswith("flow:") else action
        action_name = {"check": "审核前检查", "audit": "汇总核查表校验", "summary": "汇总校验结果说明"}.get(display, display)
        self._log(f"开始执行：{action_name}")
        threading.Thread(target=self._worker, args=(action, strict), daemon=True).start()
        return True

    def _recognize(
        self,
        *,
        force_template: bool = False,
        allow_template_auto: bool = False,
    ) -> None:
        input_dir = Path(self.state["input"]) if self.state["input"] else None
        template_dir = Path(self.state["templateDir"]) if self.state["templateDir"] else None
        if input_dir and input_dir.is_dir():
            period = detect_period(input_dir, recursive=bool(self.state["recursive"]))
            self.state["detectedPeriod"] = period.period
            self.state["explanationFiles"] = [
                {"path": str(path), "name": path.name}
                for path in explanation_files(
                    input_dir, recursive=bool(self.state["recursive"])
                )
            ]
            if not self.state["outputPinned"]:
                self.state["output"] = str(input_dir / "执行结果")
                self.state["outputAuto"] = True
        elif self.state.get("explanationFiles"):
            self.state["explanationFiles"] = []
        if input_dir and template_dir and input_dir.is_dir() and template_dir.is_dir():
            files = classify_source_files(
                template_dir, input_dir, self.project_root / "data" / "模板索引.json",
                recursive=bool(self.state["recursive"]),
            )
            self.state["sourceFiles"] = files
            current = set(self.state.get("selectedFiles") or [])
            available = {item["path"] for item in files}
            self.state["selectedFiles"] = (
                [item["path"] for item in files]
                if not current else [path for path in current if path in available]
            )
            templates = sorted({item["template"] for item in files if item["template"] != "未识别"})
            self.state["mixedTemplates"] = templates
            # 自动匹配只能由两处触发：用户点击“自动识别模板”，或选择源数据目录。
            # 执行流程、刷新目录、选择模板目录等场景只更新清单，不得改写模板选择。
            should_recommend = force_template or (
                allow_template_auto and not self.state.get("templateManual")
            )
            if should_recommend:
                result = recommend_template(
                    template_dir, input_dir, self.project_root / "data" / "模板索引.json",
                    recursive=bool(self.state["recursive"]),
                )
                if result.template_path:
                    self.state["template"] = str(result.template_path)
                    self.state["templateManual"] = False
                self._log(result.details)
            if len(templates) > 1:
                self._log("发现多种报表：" + "、".join(templates) + "。建议取消勾选不属于本次审核类型的文件。")

    def _worker(self, action: str, strict: bool = True) -> None:
        started = time.monotonic()
        try:
            # 执行前仅刷新待处理文件、数据期等状态，不重新匹配或改写模板。
            self._log("正在刷新待审核文件清单……")
            self._recognize(allow_template_auto=False)
            self._log(
                f"待处理文件：{len(self.state.get('selectedFiles', []))} 个；"
                "正在读取执行流程……"
            )
            self._save_settings()
            service = AuditService(
                config_path=self.history_path,
                engine_preference=str(self.state["calculationEngine"]),
            )
            output = Path(self.state["output"])
            selected = [Path(path) for path in self.state.get("selectedFiles", [])]
            if action in {"check", "audit"} and not selected:
                raise ValueError("请至少勾选一个待审核文件")
            if action == "check":
                result = service.preflight(template_path=Path(self.state["template"]), input_dir=Path(self.state["input"]), output_dir=output, selected_files=selected, external_path=Path(self.state["external"]) if self.state["external"] else None, recursive=bool(self.state["recursive"]), on_step=self._log_detail)
            else:
                if action == "audit":
                    action = "flow:汇总核查表校验"
                if action.startswith("flow:"):
                    flow_name = action[len("flow:"):]
                else:
                    flow_name = "汇总校验结果说明"
                standalone_combine_flow = service.is_standalone_combine_flow(flow_name)
                if standalone_combine_flow:
                    self._log(f"当前流程：{flow_name}；仅使用源数据目录")
                else:
                    self._log(f"当前流程：{flow_name}；模板：{Path(self.state['template']).name}")
                self._log("正在启动流程处理引擎……", detail=True)
                result = service.run_flow(
                    flow_name=flow_name,
                    template_path=None if standalone_combine_flow else Path(self.state["template"]),
                    input_dir=Path(self.state["input"]), output_dir=output,
                    period="" if standalone_combine_flow else (self.state.get("detectedPeriod") or Path(self.state["input"]).name),
                    history_path=self.history_path, selected_files=selected,
                    external_path=None if standalone_combine_flow else (Path(self.state["external"]) if self.state["external"] else None),
                    recursive=bool(self.state["recursive"]), on_step=self._log_detail, strict=strict,
                    write_flow_logs=bool(self.state["writeFlowLogs"]),
                )
            self.state["status"] = result.summary_text().splitlines()[0]
            self._log(result.summary_text())
        except Exception as exc:
            self.state["status"] = "执行失败"
            self._log("执行失败：" + str(exc))
        finally:
            elapsed = time.monotonic() - started
            self._log(f"任务结束，本次耗时：{elapsed:.1f} 秒")
            self.state["busy"] = False

    def _log(self, text: str, detail: bool = False) -> None:
        """记录一条运行日志。detail=False 为面向用户的简明结果，detail=True
        为逐功能/逐文件的调试细节，前端“简单”模式会过滤掉 detail 条目。"""
        stamp = datetime.now().strftime("%H:%M:%S")
        entry = {"text": f"[{stamp}] {text}", "detail": bool(detail)}
        self.state["log"] = (self.state["log"] + [entry])[-100:]

    def _log_detail(self, text: str) -> None:
        self._log(text, detail=True)

    def _save_settings(self) -> None:
        self.settings.last_input_dir = self.state["input"]
        self.settings.last_template_dir = self.state["templateDir"]
        self.settings.last_external_file = self.state["external"]
        self.settings.last_output_dir = self.state["output"]
        self.settings.output_pinned = bool(self.state["outputPinned"])
        self.settings.recursive_folders = bool(self.state["recursive"])
        self.settings.write_flow_logs = bool(self.state["writeFlowLogs"])
        self.settings.confirm_before_run = bool(self.state["confirmBeforeRun"])
        self.settings.calculation_engine = str(self.state["calculationEngine"])
        self.settings.show_custom_features = bool(self.state["showCustomFeatures"])
        self.settings_store.save(self.settings)
